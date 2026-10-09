# -*- coding: utf-8 -*-
"""
Extrai da planilha-base (layout CONSULTORIA.xlsx) os dados de UM aluno / UM bloco.

Uso:
  python extrair_dados.py planilha.xlsx --saida pasta/ [--bloco N] [--usa-pse]

Gera em --saida:
  dados.json            -> entrada do gerar_pdf.py (prescrição EXATAMENTE como está na planilha)
  relatorio_revisao.md  -> para MIGUEL (nunca vai ao aluno): alertas, inconsistências, dados faltantes

Regras:
  * Nunca inventa dados. Célula vazia = ausente.
  * %1RM exibido = o valor calculado pela própria planilha. Recalculado de forma independente
    só para CONFERIR; divergências vão para o relatório.
  * Nada é alterado silenciosamente: tudo que for estranho vira alerta no relatório.
"""
import argparse
import datetime
import json
import os
import re
import sys
import unicodedata

import openpyxl

# --- estrutura da planilha-base -------------------------------------------------
TREINO_LINHAS = {'01': (6, 12), '02': (18, 24), '03': (30, 36), '04': (42, 48)}   # só fallback (layout antigo)


def detectar_treinos(ws_val, ws_form=None):
    """Descobre onde estão os treinos na aba 'bloco NN' (o espaçamento mudou entre versões da planilha-base:
    antigo = 7 linhas por treino a cada 12; atual = 10 linhas por treino a cada 15). Procura 'TREINO NN' na coluna A;
    as linhas de exercício começam 3 linhas abaixo (após SETS/REPS/... e 'exercícios') e vão até a última linha com a
    fórmula de %1RM (coluna F) ou até a linha anterior ao próximo treino."""
    heads = []
    for r in range(1, 200):
        v = ws_val.cell(r, 1).value
        m = re.fullmatch(r'\s*TREINO\s*(\d+)\s*', str(v or ''), flags=re.I)
        if m:
            heads.append((f'{int(m.group(1)):02d}', r))
    if not heads:
        return dict(TREINO_LINHAS)
    out = {}
    for i, (num, h) in enumerate(heads):
        r0 = h + 3
        limite = (heads[i + 1][1] - 1) if i + 1 < len(heads) else r0 + 14
        # Lê TODAS as linhas até o próximo treino (linhas vazias são ignoradas depois). NÃO depender da fórmula de %1RM:
        # se Miguel apagar a fórmula de uma linha, os exercícios dela não podem sumir.
        r1 = r0
        for r in range(r0, limite + 1):
            if not vazio(ws_val.cell(r, 2).value) or (ws_form is not None and str(ws_form.cell(r, 6).value or '').startswith('=')):
                r1 = r
        out[num] = (r0, r1)
    return out
SEMANA_COL0 = [3, 9, 15, 21]          # C, I, O, U  (SETS, REPS, KG, %1RM, OBS) — layout antigo
COLUNAS_ANTIGO = [(2, c0) for c0 in SEMANA_COL0]


def colunas_semanas(ws_val):
    """[(coluna do exercício, coluna de SETS)] das 4 semanas da aba 'bloco NN' (1 = A).
    Layout antigo: exercício só em B; semanas em C, I, O, U.
    Layout novo (planilhas-base de out/2026, Miguel 7/out/2026): cada semana com a própria coluna de exercício,
    7 colunas por semana: B/C, I/J, P/Q, W/X (exercício, SETS, REPS, KG, %1RM, OBS).
    Lido do cabeçalho do 1º treino ('SETS' na linha abaixo de 'TREINO NN', 'exercícios' na seguinte)."""
    for r in range(1, 200):
        if re.fullmatch(r'\s*TREINO\s*\d+\s*', str(ws_val.cell(r, 1).value or ''), flags=re.I):
            sets = [c for c in range(2, 60) if str(ws_val.cell(r + 1, c).value or '').strip().upper() == 'SETS']
            if len(sets) < 4:
                break
            out = []
            for c0 in sets[:4]:
                rot = sem_acento(ws_val.cell(r + 2, c0 - 1).value or '').strip()
                out.append((c0 - 1 if rot.startswith('exercicio') else 2, c0))
            return out
    return list(COLUNAS_ANTIGO)
MACRO_COL_INICIO = {1: 2, 2: 6, 3: 10}  # B6, F6, J6 = data inicial de cada bloco
PRS_LINHA0 = {1: 6, 2: 18, 3: 30}       # 1ª linha de PRs do bloco (7 linhas)
# Ordem de correspondência idêntica à da fórmula da planilha (REGEXMATCH aninhado)
BASES = [('agachamento', 'Agachamento'), ('supino', 'Supino'), ('terra', 'Terra'),
         ('desenvolvimento', 'Desenvolvimento'), ('remada', 'Remada'),
         ('arranco', 'Arranco'), ('arremesso', 'Arremesso')]   # Arranco/Arremesso: planilha-base nova (out/2026)
OBS_GERAL_PADRAO = ['faca os exercicios na ordem proposta', 'considere kg = 20kg (barra) + peso em anilhas']
PALAVRAS_ESTETICA = ['graxa', 'gordura', 'emagrec', 'hipertrof', 'estetic', 'definicao', 'barriga', 'secar',
                     'perder peso', 'massa magra', 'shape']
DIAS_SEMANA = ['SEGUNDA', 'TERÇA', 'QUARTA', 'QUINTA', 'SEXTA', 'SÁBADO', 'DOMINGO']


def sem_acento(s):
    return ''.join(c for c in unicodedata.normalize('NFD', str(s)) if unicodedata.category(c) != 'Mn').lower()


def brzycki_frac(reps):
    return 1.0278 - 0.0278 * reps


def vazio(v):
    return v is None or (isinstance(v, str) and not v.strip())


def numero(v):
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return v
    try:
        return float(str(v).replace(',', '.'))
    except ValueError:
        return None


def resolver(wbf, wbv, sheet, coord, _prof=0):
    """Valor em cache; se vazio e for referência simples (='aba'!C2), segue a referência."""
    v = wbv[sheet][coord].value
    if not vazio(v):
        return v
    f = wbf[sheet][coord].value
    if isinstance(f, str) and f.startswith('=') and _prof < 3:
        m = re.fullmatch(r"=\s*'?([^'!]+)'?!\$?([A-Z]+)\$?(\d+)", f.strip())
        if m and m.group(1) in wbv.sheetnames:
            return resolver(wbf, wbv, m.group(1), m.group(2) + m.group(3), _prof + 1)
    return None


URL_RE = re.compile(r'https?://\S+')


def separar_link(obs):
    """Tira o link (YouTube etc.) da OBS: devolve (obs_sem_link, link)."""
    if not obs:
        return None, None
    m = URL_RE.search(obs)
    if not m:
        return obs, None
    link = m.group(0)
    resto = (obs[:m.start()] + ' ' + obs[m.end():]).strip()
    resto = re.sub(r'^(\s*-\s*)+|(\s*-\s*)+$', '', resto).strip()
    return (resto or None), link


def chave_ex(ex):
    k = re.sub(r'\s+', ' ', sem_acento(ex)).strip()
    # Decisão de Miguel: variação 'sem step' é o mesmo exercício para a regra Aquecimento/Principal
    return re.sub(r'\s*-\s*sem step$', '', k, flags=re.I).strip()


def aplicar_regra_series(rows, sem, num, alertas, infos):
    """Regra de Miguel para exercícios que se repetem no MESMO treino (mesmo nome):
    1ª ocorrência = Aquecimento; última = Principal; intermediária com carga maior que a anterior = Aquecimento.
    Linha com a MESMA carga da última (ex.: 2×3 + 1×AMRAP em duas linhas) também é Principal (Miguel, 7/out/2026).
    Exercício que NÃO se repete: sem etiqueta Aquecimento/Principal. Outras etiquetas (Gravar, Biset, AMRAP, notas) e links
    são preservados."""
    grupos = {}
    for i, r in enumerate(rows):
        obs, link = separar_link(r['obs'])
        r['link'] = link or r.pop('link_celula', None)
        r.pop('link_celula', None)
        partes = [p.strip() for p in obs.split(' - ')] if obs else []
        r['_orig'] = ' - '.join(p for p in partes if p.lower() in ('aquecimento', 'principal')) or None
        r['_partes'] = [p for p in partes if p.lower() not in ('aquecimento', 'principal')]
        grupos.setdefault(chave_ex(r['ex']), []).append(i)
    for idxs in grupos.values():
        tags = {}
        if len(idxs) > 1:
            kg_ult = rows[idxs[-1]]['kg']
            for pos, i in enumerate(idxs):
                if pos == len(idxs) - 1 or (kg_ult and rows[i]['kg'] == kg_ult):
                    tags[i] = 'Principal'
                elif pos == 0:
                    tags[i] = 'Aquecimento'
                else:
                    ant, atual = rows[idxs[pos - 1]]['kg'], rows[i]['kg']
                    tags[i] = 'Aquecimento'
                    if not (ant and atual and atual > ant):
                        alertas.append(f'S{sem:02d} T{num} "{rows[i]["ex"]}": linha intermediária sem aumento de carga em relação à anterior; rotulada "Aquecimento" por padrão. Confirme.')
        for i in idxs:
            r = rows[i]
            partes = ([tags[i]] if i in tags else []) + r['_partes']
            r['obs'] = ' - '.join(partes) if partes else None
            novo = tags.get(i)
            r['_mudou'] = (r['_orig'] or '').lower() != (novo or '').lower()
            if r['_orig'] and r['_mudou']:
                infos.append(f'S{sem:02d} T{num} {r["ex"]} ({fs(r["sets"])}×{r["reps"]}): etiqueta da planilha "{r["_orig"] or "—"}" → PDF "{novo or "—"}" (regra de aquecimento/principal).')
    n_auto = sum(1 for r in rows if r.pop('_mudou', False))
    infos.append(('_AUTO', n_auto)) if n_auto else None
    for r in rows:
        r.pop('_orig', None)
        r.pop('_partes', None)



def aplicar_gravar(rows, sem, num, adicionados):
    """Regra de Miguel: toda série com %1RM (valor exibido, arredondado) de 80% OU MAIS leva GRAVAR automaticamente (Miguel, 7/out/2026: inclui 80%)."""
    for r in rows:
        partes = [p.strip() for p in r['obs'].split(' - ')] if r['obs'] else []
        if any(p.lower() == 'aquecimento' for p in partes):
            continue          # Miguel (7/out/2026): aquecimento não ganha GRAVAR automático
        if r['pct'] is not None and round(r['pct'] * 100) >= 80:
            if not any(p.lower() == 'gravar' for p in partes):
                r['obs'] = ' - '.join(partes + ['Gravar'])
                adicionados.append(f'S{sem:02d} T{num} {r["ex"]} {fs(r["sets"])}×{r["reps"]} @ {r["kg"]:g} kg ({round(r["pct"]*100)}%)')


def aplicar_amrap(rows, sem, num, adicionados, alertas):
    """Regra de Miguel (4/out/2026): série AMRAP SEMPRE leva GRAVAR, e AMRAP aparece só na coluna REPS
    (a etiqueta AMRAP sai da OBS)."""
    for r in rows:
        partes = [p.strip() for p in r['obs'].split(' - ')] if r['obs'] else []
        if 'amrap' not in (str(r['reps']) + ' ' + ' '.join(partes)).lower():
            continue
        if 'amrap' not in str(r['reps']).lower():
            if not vazio(r['reps']) and r['reps'] != '—':
                alertas.append(f'S{sem:02d} T{num} "{r["ex"]}": REPS = {r["reps"]} e AMRAP na OBS; PDF mostra "{r["reps"]} + AMRAP" nas repetições. Confirme.')
                r['reps'] = f'{r["reps"]} + AMRAP'
            else:
                r['reps'] = 'AMRAP'
        partes = [p for p in partes if p.lower() != 'amrap']
        if not any(p.lower() == 'gravar' for p in partes):
            partes.append('Gravar')
            adicionados.append(f'S{sem:02d} T{num} {r["ex"]} {fs(r["sets"])}×{r["reps"]} (AMRAP)')
        r['obs'] = ' - '.join(partes) or None


FAIXA_RE = re.compile(r'^\s*(\d+)\s*[-–—]\s*(\d+)\s*$')


def aplicar_anotar(rows, sem, num, adicionados):
    """Regra de Miguel (5/out/2026): toda linha com FAIXA de repetições ("3 - 5", "8 - 12"...) ou AMRAP leva ANOTAR
    na OBS, com ou sem carga prescrita: o aluno anota as repetições (e a carga) que fez em cada série."""
    for r in rows:
        amrap = 'amrap' in str(r['reps']).lower()
        if not faixa(r['reps']) and not amrap:
            continue
        partes = [p.strip() for p in r['obs'].split(' - ')] if r['obs'] else []
        if not any(p.lower() == 'anotar' for p in partes):
            r['obs'] = ' - '.join(partes + ['Anotar'])
            adicionados.append(f'S{sem:02d} T{num} {r["ex"]} {fs(r["sets"])}×{r["reps"]}' + (' (AMRAP)' if amrap else '') + ('' if r['kg'] else ' (sem carga)'))


def faixa(v):
    """Regra de Miguel (4/out/2026): faixa sempre com espaço dos dois lados do hífen. '3-4' -> '3 - 4'; não-faixa -> None."""
    m = FAIXA_RE.match(str(v)) if isinstance(v, str) else None
    return f'{m.group(1)} - {m.group(2)}' if m else None


def fs(v):
    """Séries para textos do relatório: 3.0 -> '3'; faixa '3 - 4' fica como está."""
    return f'{v:g}' if isinstance(v, (int, float)) else str(v)


def reps_superior(reps):
    """Para checagens: 12 -> 12 ; '8 - 12' -> 12 ; texto sem número -> None."""
    n = numero(reps)
    if n is not None:
        return n
    nums = re.findall(r'\d+', str(reps))
    return float(nums[-1]) if nums else None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('planilha')
    ap.add_argument('--saida', required=True)
    ap.add_argument('--bloco', type=int, help='nº do bloco (padrão: último bloco com prescrição)')
    ap.add_argument('--numeros', help='exercícios que aparecem em "SEUS NÚMEROS" (máx. 4), ex.: "Supino,Terra,Remada,Arranco". '
                    'Padrão: todos os da aba prs (o PDF aborta se passar de 4).')
    ap.add_argument('--usa-pse', action='store_true', help='aluno usa PSE (inclui PSE na legenda)')
    a = ap.parse_args()
    os.makedirs(a.saida, exist_ok=True)

    wbf = openpyxl.load_workbook(a.planilha)
    wbv = openpyxl.load_workbook(a.planilha, data_only=True)
    nomes = wbv.sheetnames
    blocos_disp = sorted(int(m.group(1)) for n in nomes if (m := re.fullmatch(r'bloco (\d+)', n)))
    if not blocos_disp:
        sys.exit('ERRO: nenhuma aba "bloco NN" encontrada.')

    def tem_prescricao(n):
        ws = wbv[f'bloco {n:02d}']
        for a0, b0 in detectar_treinos(ws, wbf[f'bloco {n:02d}']).values():
            for r in range(a0, b0 + 1):
                if any(not vazio(ws.cell(r, ce).value) and any(not vazio(ws.cell(r, c0 + i).value) and ws.cell(r, c0 + i).value != 0
                                                               for i in (0, 1))
                       for ce, c0 in colunas_semanas(ws)):
                    return True
        return False

    bloco = a.bloco or max((n for n in blocos_disp if tem_prescricao(n)), default=None)
    if bloco is None or f'bloco {bloco:02d}' not in nomes:
        sys.exit('ERRO: nenhum bloco com prescrição encontrado (a planilha parece ser só o layout vazio).')
    bsheet = f'bloco {bloco:02d}'
    asheet = f'aluno {bloco:02d}' if f'aluno {bloco:02d}' in nomes else 'aluno 01'
    ws, wa = wbv[bsheet], wbv[asheet]
    treino_linhas = detectar_treinos(ws, wbf[bsheet])
    colunas = colunas_semanas(ws)

    alertas, infos = [], []   # (texto)
    infos.append('Layout da aba do bloco: ' + ', '.join(f'T{n}: linhas {a0}–{b0}' for n, (a0, b0) in treino_linhas.items()) + '.')
    if colunas != COLUNAS_ANTIGO:
        infos.append('Layout novo: exercício de cada semana na própria coluna (B, I, P, W).')

    # --- aluno ---------------------------------------------------------------
    nome = resolver(wbf, wbv, asheet, 'C2') or resolver(wbf, wbv, 'aluno 01', 'C2')
    objetivo = resolver(wbf, wbv, asheet, 'C3') or resolver(wbf, wbv, 'aluno 01', 'C3')
    if vazio(nome):
        sys.exit('ERRO: nome do aluno (aba aluno, C2) está vazio.')
    nome = str(nome).strip()
    objetivo = None if vazio(objetivo) else str(objetivo).strip()
    if not objetivo:
        alertas.append('Objetivo do aluno (aba aluno, C3) está vazio: a página de perfil mostrará "—".')

    dias, outros_dias = {}, {}
    for c in 'BCDEF':
        d = resolver(wbf, wbv, asheet, c + '6') or resolver(wbf, wbv, 'aluno 01', c + '6')
        t = resolver(wbf, wbv, asheet, c + '5') or resolver(wbf, wbv, 'aluno 01', c + '5')
        if not vazio(d) and not vazio(t):
            if re.fullmatch(r'T\d+', str(d).strip().upper()):
                dias[str(t).strip().upper()] = str(d).strip().upper()   # {'SEGUNDA': 'T01'}
            else:
                outros_dias[str(t).strip().upper()] = str(d).strip()

    # --- data de início ------------------------------------------------------
    inicio = None
    mcol = MACRO_COL_INICIO.get(bloco)
    if mcol:
        v = wbv['macro'].cell(6, mcol).value
        if isinstance(v, datetime.datetime):
            inicio = v
        elif bloco > 1:
            base = wbv['macro'].cell(6, 2).value
            if isinstance(base, datetime.datetime):
                inicio = base + datetime.timedelta(days=28 * (bloco - 1))
                infos.append('Data de início do bloco calculada (início do bloco 01 + 28 dias × blocos): célula da macro sem valor em cache.')
    if inicio is None:
        alertas.append('Data de início do bloco não encontrada na aba macro: a capa ficará sem data.')

    # --- PRs / 1RM estimado (aba prs) ---------------------------------------
    rm, rm_prev = {}, {}
    wp = wbv['prs']

    def ler_prs(n):
        out = {}
        r0 = PRS_LINHA0.get(n)
        if not r0:
            return out
        for r in range(r0, r0 + 7):
            nm, reps, kg = wp.cell(r, 1).value, numero(wp.cell(r, 2).value), numero(wp.cell(r, 3).value)
            if not vazio(nm) and reps and kg:
                est = kg / brzycki_frac(reps)
                cache = numero(wp.cell(r, 4).value)
                if cache is not None and abs(cache - est) > 0.01:
                    alertas.append(f'PRs: 1RM estimado de {nm} na planilha ({cache:.2f}) difere do recálculo ({est:.2f}).')
                out[str(nm).strip()] = {'reps': int(reps) if float(reps).is_integer() else reps, 'kg': kg, 'est': est}
        return out
    rm = ler_prs(bloco)
    if bloco > 1:
        rm_prev = ler_prs(bloco - 1)
    if not rm:
        infos.append('Aba prs sem séries de referência para este bloco: a seção "SEUS NÚMEROS" só aparecerá se você enviar PRs reais com fonte no contexto.')

    def base_de(ex):
        n = sem_acento(ex)
        for k, lab in BASES:
            if k in n:
                return lab
        return None

    rm_pdf = rm
    if a.numeros:
        pedidos = [sem_acento(x).strip() for x in a.numeros.split(',') if x.strip()]
        rm_pdf = {k: v for k, v in rm.items() if sem_acento(k) in pedidos}
        faltam = [x for x in pedidos if x not in {sem_acento(k) for k in rm}]
        if faltam:
            alertas.append('--numeros pediu exercício(s) sem referência na aba prs: ' + ', '.join(faltam) + '.')
        infos.append('SEUS NÚMEROS limitado por Miguel a: ' + (', '.join(rm_pdf) or 'nenhum') + '.')

    # --- hiperlinks de célula (texto visível = título; URL fica atrás) ---------------
    # Miguel: SEMPRE que houver link de vídeo, vira o botão "Ver vídeo". Os links podem estar (a) como URL no texto da OBS,
    # (b) como hiperlink de célula na OBS da aba bloco ou (c) como hiperlink de célula na aba aluno (que espelha o bloco).
    def alvo_link(h):
        t = (h.target or h.location or '') if h is not None else ''
        return t.strip() if str(t).strip().lower().startswith('http') else None
    links_cel = {}
    af = wbf[asheet]
    semana_atual = None
    for rr in range(1, min(af.max_row, 400) + 1):
        a0 = af.cell(rr, 1).value
        m = re.fullmatch(r'\s*SEMANA\s*(\d+)\s*', str(a0 or ''), flags=re.I)
        if m:
            semana_atual = int(m.group(1))
            continue
        u = alvo_link(af.cell(rr, 6).hyperlink)
        mm = re.search(r"'?!\$?[A-Z]+\$?(\d+)", str(a0 or ''))
        if u and semana_atual and mm:
            links_cel[(semana_atual, int(mm.group(1)))] = u

    # --- AQC -----------------------------------------------------------------
    series_texto = []   # séries escritas como texto/faixa ("3-4"): o template só exibe número → parar com aviso claro
    aqc = []
    for r in (13, 14, 15):
        ex = wa.cell(r, 1).value
        if vazio(ex):
            continue
        sets, reps, kg, obs = wa.cell(r, 2).value, wa.cell(r, 3).value, wa.cell(r, 4).value, wa.cell(r, 5).value
        if not vazio(sets) and numero(sets) is None:
            if faixa(sets):
                sets = faixa(sets)
            else:
                series_texto.append(f'AQC linha {r} "{str(ex).strip()}": séries = "{sets}"')
        elif vazio(sets):
            infos.append(f'AQC "{str(ex).strip()}" sem séries: o PDF mostra "—" na coluna SÉRIES.')
        if faixa(reps):
            reps = faixa(reps)
        item = {'ex': str(ex).strip(), 'sets': numero(sets) if numero(sets) is not None else sets,
                'reps': ('' if reps is None else (int(reps) if isinstance(reps, (int, float)) and float(reps).is_integer() else re.sub(r'(?<![\d.,])(\d+)\.0+(?!\d)', r'\1', str(reps)))), 'link': None, 'nota': None}
        if not vazio(obs):
            if str(obs).strip().lower().startswith('http'):
                item['link'] = str(obs).strip()
            else:
                item['nota'] = str(obs).strip()
        u_aqc = alvo_link(wbf[asheet].cell(r, 5).hyperlink)
        if u_aqc:
            item['link'], item['nota'] = u_aqc, None   # texto visível da célula era só o título do vídeo
        if not vazio(kg):
            item['kg'] = numero(kg) if numero(kg) is not None else str(kg).strip()   # Miguel: mostrar SEMPRE a carga do AQC quando existir
        aqc.append(item)

    # --- semanas / treinos ----------------------------------------------------
    gravar_auto, gravar_amrap, anotar_auto = [], [], []
    semanas, semanas_vazias = [], []
    rows_por_semana = {}
    for k, (ce, c0) in enumerate(colunas):
        treinos = {}
        for num, (r0, r1) in treino_linhas.items():
            rows = []
            for r in range(r0, r1 + 1):
                ex = ws.cell(r, ce).value
                sets, reps = ws.cell(r, c0).value, ws.cell(r, c0 + 1).value
                kg = numero(ws.cell(r, c0 + 2).value)
                pct_cache = ws.cell(r, c0 + 3).value
                obs = ws.cell(r, c0 + 4).value
                if vazio(reps) and not vazio(ex) and not vazio(sets) and sets != 0 and 'amrap' in str(obs or '').lower():
                    reps = 'AMRAP'   # Miguel (4/out/2026): AMRAP aparece na coluna REPS, não como etiqueta na OBS
                    alertas.append(f'S{k+1:02d} T{num} "{ex}": AMRAP estava só na OBS (REPS vazia); PDF mostra "AMRAP" nas repetições. Confirme.')
                if vazio(ex) or vazio(sets) or vazio(reps) or sets == 0:
                    if not vazio(ex) and (not vazio(sets) or not vazio(reps)):
                        alertas.append(f'Semana {k+1:02d} · Treino {num} · linha {r}: "{ex}" com séries/reps incompletas (não incluído).')
                    continue
                if numero(sets) is None:
                    if faixa(sets):
                        sets = faixa(sets)
                        infos.append(f'S{k+1:02d} T{num} "{str(ex).strip()}": séries em faixa ({sets}); PDF mostra "{sets}" na coluna SÉRIES.')
                    else:
                        series_texto.append(f'Semana {k+1:02d} · Treino {num} · linha {r} "{str(ex).strip()}": séries = "{sets}"')
                        continue
                if faixa(reps):
                    reps = faixa(reps)
                base = base_de(ex)
                calc = (kg / rm[base]['est']) if (kg and base in rm) else None
                pc = numero(pct_cache)
                if pc is not None and pc > 0:
                    pct = pc
                    if calc is not None and abs(pc - calc) > 0.002:
                        alertas.append(f'S{k+1:02d} T{num} "{ex}": %1RM da planilha ({pc*100:.1f}%) difere do recálculo ({calc*100:.1f}%).')
                else:
                    pct = None
                    if calc is not None:
                        alertas.append(f'S{k+1:02d} T{num} "{ex}": planilha não calculou %1RM, mas há 1RM de {base} (daria {calc*100:.0f}%). PDF mostra "—".')
                rows.append({'ex': str(ex).strip(), 'sets': numero(sets) if numero(sets) is not None else sets,
                             'reps': (int(numero(reps)) if numero(reps) is not None and float(numero(reps)).is_integer() else reps),
                             'kg': kg if kg else None, 'pct': pct,
                             'obs': None if vazio(obs) else str(obs).strip(), 'base': base, 'linha': r,
                             'link_celula': alvo_link(wbf[bsheet].cell(r, c0 + 4).hyperlink) or links_cel.get((k + 1, r))})
            if rows:
                aplicar_regra_series(rows, k + 1, num, alertas, infos)
                aplicar_gravar(rows, k + 1, num, gravar_auto)
                aplicar_amrap(rows, k + 1, num, gravar_amrap, alertas)
                aplicar_anotar(rows, k + 1, num, anotar_auto)
                treinos[num] = rows
        if treinos:
            semanas.append({'numero': k + 1, 'treinos': treinos})
            rows_por_semana[k + 1] = treinos
        else:
            semanas_vazias.append(k + 1)
    if series_texto:
        sys.exit('ERRO: séries escritas como texto, que o PDF não exibe (a coluna SÉRIES aceita número ou faixa "3 - 4"):\n  - '
                 + '\n  - '.join(series_texto) +
                 '\nNÃO escolha um número por conta própria (alteraria a prescrição): pergunte a Miguel qual valor usar '
                 'e peça para corrigir a planilha.')
    if not semanas:
        sys.exit('ERRO: o bloco não tem nenhuma semana prescrita.')
    if semanas_vazias:
        infos.append(f'Semanas sem prescrição (não entram no PDF): {", ".join(f"{s:02d}" for s in semanas_vazias)}.')
    for s in semanas:
        faltam = [n for n in rows_por_semana[semanas[0]['numero']] if n not in s['treinos']]
        if faltam:
            alertas.append(f'Semana {s["numero"]:02d} não tem o(s) Treino(s) {", ".join(faltam)} que existem na semana {semanas[0]["numero"]:02d}.')

    # --- checagens ------------------------------------------------------------
    n_treinos = max(len(s['treinos']) for s in semanas)
    if outros_dias:
        alertas.append('Marcador(es) na aba aluno que NÃO são treino: ' + ', '.join(f'{k} = "{v}"' for k, v in outros_dias.items()) +
                       '. Não contam como dia de treino no PDF (perfil mostra só os dias T01, T02...). Confirme.')
    if dias and len(dias) != n_treinos:
        alertas.append(f'Dias de treino na aba aluno ({len(dias)}: {", ".join(dias)}) ≠ treinos prescritos por semana ({n_treinos}).')
    if not dias:
        alertas.append('Dias da semana (aba aluno, B5:F6) vazios: os treinos aparecerão sem o dia.')
    observacoes = []
    obs_geral = resolver(wbf, wbv, asheet, 'A9')
    if not vazio(obs_geral):
        padrao = [sem_acento(x) for x in OBS_GERAL_PADRAO]
        observacoes = [l.strip() for l in str(obs_geral).splitlines() if l.strip() and sem_acento(l).strip() not in padrao]
        if observacoes:
            infos.append('OBS geral da planilha exibida no PDF como "Observação": ' + ' | '.join(observacoes))

    # cargas vs RM estimado (Brzycki) e bases por palavra-chave
    usadas_chave = {}
    for s in semanas:
        for num, rows in s['treinos'].items():
            for r in rows:
                if r['base'] and r['ex'].lower() != r['base'].lower():
                    usadas_chave.setdefault((r['ex'], r['base']), 0)
                if r['pct'] and r['obs'] and 'aquecimento' in r['obs'].lower():
                    continue
                rs = reps_superior(r['reps'])
                if r['pct'] and rs and 'amrap' not in str(r['reps']).lower():
                    ratio = r['pct'] / brzycki_frac(rs)
                    ref = r['kg'] / ratio if ratio else None
                    rep_txt = f'{fs(r["sets"])}×{r["reps"]}'
                    if ratio >= 1.005:
                        alertas.append(f'S{s["numero"]:02d} T{num} {r["ex"]} {rep_txt} @ {r["kg"]:g} kg = {r["pct"]*100:.0f}% do 1RM estimado, ACIMA do {rs:g}RM estimado (≈ {r["kg"]/ratio:.1f} kg). Séries no limite/além; confirme se é intencional.')
                    elif ratio >= 0.985:
                        alertas.append(f'S{s["numero"]:02d} T{num} {r["ex"]} {rep_txt} @ {r["kg"]:g} kg ≈ {rs:g}RM estimado ({r["pct"]*100:.0f}% do 1RM): séries tendendo ao limite' + (' (topo da faixa de reps)' if isinstance(r['reps'], str) else '') + '.')
    if usadas_chave:
        infos.append('%1RM por palavra-chave da planilha (exercício → 1RM usado): ' + '; '.join(f'{ex} → {base}' for ex, base in usadas_chave) + '. Variações usam o 1RM do exercício-base.')

    # VTT / VTR (macro)
    vt_excl = 0
    vtt = vtr = 0.0
    for s in semanas:
        for rows in s['treinos'].values():
            for r in rows:
                if isinstance(r['reps'], int) and isinstance(r['sets'], (int, float)):
                    vtr += r['sets'] * r['reps']
                    if r['kg']:
                        vtt += r['sets'] * r['reps'] * r['kg']
                else:
                    vt_excl += 1
    if vt_excl:
        alertas.append(f'VTT/VTR da aba macro só somam exercícios com séries e repetições NUMÉRICAS FIXAS; {vt_excl} linha(s) com faixa (ex.: "8 - 12", séries "3 - 4") ou AMRAP entram como zero. O volume real é maior que o da macro. (Soma dos exercícios fixos nas semanas prescritas: VTT {vtt:g} kg · VTR {vtr:g} reps.)')

    # PSE / AMRAP
    tem_amrap = any('amrap' in (str(r['reps']) + ' ' + str(r['obs'] or '')).lower()
                    for s in semanas for rows in s['treinos'].values() for r in rows)
    if tem_amrap:
        infos.append('AMRAP detectado: a legenda ganha o item AMRAP. VTT/VTR não calculam séries AMRAP.')
    if a.usa_pse:   # Miguel (5/out/2026): nenhum aluno usa PSE por enquanto; desligada por padrão e sem perguntar
        infos.append('PSE: ativado (--usa-pse). A legenda inclui PSE.')

    # evolução vs bloco anterior
    evol = []
    if rm_prev:
        for k, v in rm.items():
            if k in rm_prev:
                evol.append((k, rm_prev[k]['est'], v['est']))

    if anotar_auto:
        infos.insert(0, f'ANOTAR adicionado automaticamente (faixa de repetições ou AMRAP) em {len(anotar_auto)} linha(s); a legenda ganha o item ANOTAR: ' + '; '.join(anotar_auto))
    if gravar_auto:
        infos.insert(0, f'GRAVAR adicionado automaticamente (%1RM ≥ 80%) em {len(gravar_auto)} linha(s): ' + '; '.join(gravar_auto))
    if gravar_amrap:
        infos.insert(0, f'GRAVAR adicionado automaticamente (série AMRAP) em {len(gravar_amrap)} linha(s): ' + '; '.join(gravar_amrap))
    total_auto = sum(i[1] for i in infos if isinstance(i, tuple))
    infos = [i for i in infos if not isinstance(i, tuple)]
    if total_auto:
        infos.insert(0, f'Regra aquecimento/principal aplicada em {total_auto} linha(s) onde a etiqueta do PDF difere da planilha (exercícios repetidos no mesmo treino: 1ª = Aquecimento, última = Principal; sem repetição = sem etiqueta).')

    dados = {
        'aluno': {'nome': nome, 'objetivo': objetivo}, 'observacoes': observacoes,
        'bloco': {'numero': bloco, 'inicio': inicio.strftime('%d/%m/%Y') if inicio else None},
        'dias': dias, 'aqc': aqc, 'rm': rm_pdf, 'semanas': semanas, 'usa_pse': bool(a.usa_pse),
        'origem': {'arquivo': os.path.basename(a.planilha), 'aba_aluno': asheet, 'aba_bloco': bsheet},
    }
    with open(os.path.join(a.saida, 'dados.json'), 'w', encoding='utf-8') as f:
        json.dump(dados, f, ensure_ascii=False, indent=1, default=str)

    # --- relatório para Miguel ---------------------------------------------------
    L = [f'# Revisão para Miguel — {nome} · Bloco {bloco:02d}', '',
         '_Documento interno. Não vai para o aluno._', '',
         '## Dados observados',
         f'- Aluno: **{nome}** · Objetivo: {objetivo or "—"}',
         f'- Bloco {bloco:02d} · início: {inicio:%d/%m/%Y}' if inicio else f'- Bloco {bloco:02d} · início: não encontrado',
         f'- Semanas prescritas: {", ".join(f"{s["numero"]:02d}" for s in semanas)} · treinos/semana: {n_treinos} · dias: {", ".join(dias) or "—"}',
         '- 1RM estimado (aba prs): ' + (', '.join(f'{k} {v["est"]:.1f} kg ({v["reps"]}×{v["kg"]:g})' for k, v in rm.items()) or 'nenhum'),
         '- Anamnese/feedback: não fazem parte da planilha e **não entram no PDF do aluno** (decisão de Miguel); servem só para esta revisão.', '']
    if evol:
        L += ['## Evolução do 1RM estimado vs bloco anterior', '']
        for k, p, n in evol:
            L.append(f'- {k}: {p:.1f} → {n:.1f} kg ({n-p:+.1f} kg)')
        L.append('')
    L += ['## Alertas (decisão sua)', ''] + ([f'- ⚠ {t}' for t in alertas] or ['- nenhum']) + ['', '## Informações', ''] + ([f'- {t}' for t in infos] or ['- nenhuma']) + ['']
    with open(os.path.join(a.saida, 'relatorio_revisao.md'), 'w', encoding='utf-8') as f:
        f.write('\n'.join(L))
    print(f'OK  {nome} · bloco {bloco:02d} · semanas {[s["numero"] for s in semanas]} · {len(alertas)} alerta(s)')
    print('    ->', os.path.join(a.saida, 'dados.json'))
    print('    ->', os.path.join(a.saida, 'relatorio_revisao.md'))


if __name__ == '__main__':
    main()
