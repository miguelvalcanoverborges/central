# -*- coding: utf-8 -*-
"""Mensagens para o WhatsApp que acompanham os PDFs (Miguel, 11/out/2026). Regras fixas, sem IA: cada linha só entra se
o dado existe para aquele aluno; nada de prazo, promessa ou regra que não seja de Miguel. Mesmo texto da Central da
Consultoria (que importa este arquivo).

ENTREGA (acompanha a ficha, e o planejamento se ele foi gerado para o bloco):
  python mensagem.py entrega dados.json [--semanas 3,4] [--proxima dd/mm/aaaa] [--planejamento]
  dados.json = o mesmo que gerou a ficha (extrair_dados.py). --semanas = só as semanas enviadas.
RESULTADOS (acompanha o PDF de evolução, no fim do bloco):
  python mensagem.py resultados PLANILHA.xlsx --bloco N [--itens itens.json] [--proximo dd/mm/aaaa]
  Só aponta para o PDF e diz quando começa o próximo bloco: números e itens pessoais ficam no PDF, não na mensagem
  (Miguel, 11/out/2026). --itens é aceito por compatibilidade e não muda o texto.
"""
import argparse
import datetime as dt
import json
import os
import re
import sys

DIAS = ["SEGUNDA", "TERÇA", "QUARTA", "QUINTA", "SEXTA", "SÁBADO", "DOMINGO"]
CURTO = {"SEGUNDA": "seg", "TERÇA": "ter", "QUARTA": "qua", "QUINTA": "qui", "SEXTA": "sex", "SÁBADO": "sáb", "DOMINGO": "dom"}
LONGO = {"SEGUNDA": "segunda", "TERÇA": "terça", "QUARTA": "quarta", "QUINTA": "quinta", "SEXTA": "sexta",
         "SÁBADO": "sábado", "DOMINGO": "domingo"}


def _lista(xs: list) -> str:
    xs = [str(x) for x in xs if x]
    return "" if not xs else xs[0] if len(xs) == 1 else ", ".join(xs[:-1]) + " e " + xs[-1]


def _data(txt: str | None) -> dt.date | None:
    for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return dt.datetime.strptime(str(txt or "").strip(), fmt).date()
        except ValueError:
            pass
    return None


def _dm(d: dt.date | None) -> str:
    return f"{d:%d/%m}" if d else ""


def _primeiro(nome: str) -> str:
    return (nome or "").split()[0] if (nome or "").split() else ""


def _sem(n: int) -> str:
    return f"{int(n):02d}"


# ------------------------------------------------------------------ o que o aluno faz em cada linha da ficha
def _acoes(row: dict) -> list[tuple[str, str, str]]:
    """[(tipo, exercício, detalhe)] — tipo: 'amrap' (gravar e anotar), 'gravar' (com %1RM) ou 'anotar'."""
    obs = str(row.get("obs") or "").lower()
    reps = str(row.get("reps") or "").lower()
    ex = str(row.get("ex") or "").strip()
    if "aquecimento" in obs:
        return []
    if "amrap" in reps or "amrap" in obs:
        return [("amrap", ex, "")]
    out = []
    if "gravar" in obs:
        pct = row.get("pct")
        out.append(("gravar", ex, f"{round(pct * 100)}%" if pct else ""))
    if "anotar" in obs:
        out.append(("anotar", ex, ""))
    return out


def _frase(acoes: list[tuple[str, str, str]], com_detalhe=True) -> str:
    """Junta as ações de um treino numa frase: gravar …; AMRAP de … (gravar e anotar); anotar reps e carga em …"""
    ordem, vistos = [], set()
    for a in acoes:
        k = (a[0], a[1])
        if k not in vistos:
            vistos.add(k)
            ordem.append(a)
    grav = [f"{ex} ({d})" if (d and com_detalhe) else ex for t, ex, d in ordem if t == "gravar"]
    amr = [ex for t, ex, _ in ordem if t == "amrap"]
    ano = [ex for t, ex, _ in ordem if t == "anotar"]
    partes = []
    if grav:
        partes.append("gravar " + _lista(grav))
    if amr:
        partes.append("AMRAP de " + _lista(amr) + " (gravar e anotar)")
    if ano:
        partes.append("anotar reps e carga em " + _lista(ano))
    return "; ".join(partes)


def texto_entrega(D: dict, semanas: list[int] | None = None, proxima: str | None = None,
                  com_planejamento: bool = False) -> str:
    """D = dados.json da ficha. semanas = as escolhidas (vazio = todas). proxima = data da próxima revisão (ISO)."""
    todas = [int(s["numero"]) for s in D.get("semanas", [])]
    sel = sorted({int(s) for s in (semanas or []) if int(s) in todas}) or todas
    parcial = sel != todas
    bloco = int(D["bloco"]["numero"])
    ini = _data(D["bloco"].get("inicio"))
    nome = _primeiro(D["aluno"]["nome"])
    dia_de = {t: d for d, t in (D.get("dias") or {}).items()}          # 'T01' -> 'SEGUNDA'
    dias_ord = sorted(D.get("dias") or {}, key=lambda d: DIAS.index(d) if d in DIAS else 99)

    def ini_sem(n):
        return ini + dt.timedelta(days=7 * (n - 1)) if ini else None

    def data_treino(n, num):
        d0, dia = ini_sem(n), dia_de.get(f"T{num}")
        if not d0 or dia not in DIAS:
            return None
        return d0 + dt.timedelta(days=(DIAS.index(dia) - d0.weekday()) % 7)

    L = []
    # ---- abertura
    if not parcial:
        L.append(f"*{nome}, seu bloco {_sem(bloco)} está pronto.*")
    elif len(sel) == 1:
        L.append(f"*{nome}, segue a semana {_sem(sel[0])} do seu bloco {_sem(bloco)}.*")
    else:
        L.append(f"*{nome}, seguem as semanas {_lista([_sem(s) for s in sel])} do seu bloco {_sem(bloco)}.*")

    # ---- período e dias
    if ini:
        if not parcial:
            fim = ini_sem(todas[-1]) + dt.timedelta(days=6)
            n = len(todas)
            dia_ini = LONGO.get(DIAS[ini.weekday()], "")
            L.append(f"Começa {dia_ini}, {_dm(ini)}, e vai até {_dm(fim)} ({n} {'semana' if n == 1 else 'semanas'}).")
        elif sel == list(range(sel[0], sel[-1] + 1)):
            L.append(f"{'Semana' if len(sel) == 1 else 'Período'}: {_dm(ini_sem(sel[0]))} a {_dm(ini_sem(sel[-1]) + dt.timedelta(days=6))}.")
        else:
            L.append(_lista([f"semana {_sem(s)} de {_dm(ini_sem(s))} a {_dm(ini_sem(s) + dt.timedelta(days=6))}" for s in sel]).capitalize() + ".")
    if dias_ord:
        L.append(f"Treinos: {_lista([LONGO.get(d, d.lower()) for d in dias_ord])}.")

    # ---- arquivos
    L.append("")
    if com_planejamento and not parcial:
        L.append("Te mandei 2 PDFs: o seu *planejamento* (o que pensei a partir das suas respostas) e a *ficha* (o treino de cada semana).")
    else:
        L.append("Segue a ficha em PDF.")

    # ---- o que gravar e anotar, semana a semana
    por_sem = {}            # n -> [(num, [ações])]
    for s in D.get("semanas", []):
        n = int(s["numero"])
        if n not in sel:
            continue
        por_sem[n] = [(num, [a for r in rows for a in _acoes(r)]) for num, rows in sorted(s["treinos"].items())]
    # o que se repete em TODAS as semanas escolhidas (mesmo treino, mesma ação, mesmo exercício) aparece uma vez só
    fixo = set()
    if len(sel) > 1:
        conj = [{(num, t, ex) for num, acs in por_sem.get(n, []) for t, ex, _ in acs} for n in sel]
        fixo = set.intersection(*conj) if conj else set()
    linhas_fixas = []
    for num in sorted({k[0] for k in fixo}):
        acs = [(t, ex, "") for (n_, t, ex) in sorted(fixo) if n_ == num]
        dia = dia_de.get(f"T{num}")
        quando = f"Treino {num}" + (f" ({CURTO.get(dia, '')})" if dia in CURTO else "")
        linhas_fixas.append(f"• {quando}: {_frase(acs, com_detalhe=False)}")
    blocos_sem = []
    for n in sel:
        lin = []
        # em ordem de data (o bloco pode começar no meio da semana)
        treinos_n = sorted(por_sem.get(n, []), key=lambda x: (data_treino(n, x[0]) or dt.date.max, x[0]))
        for num, acs in treinos_n:
            resto = [a for a in acs if (num, a[0], a[1]) not in fixo]
            if not resto:
                continue
            dia, d = dia_de.get(f"T{num}"), data_treino(n, num)
            quando = (f"{CURTO.get(dia, '').capitalize()} {_dm(d)} · treino {num}" if d
                      else f"Treino {num}" + (f" ({CURTO.get(dia, '')})" if dia in CURTO else ""))
            lin.append(f"• {quando.strip()}: {_frase(resto)}")
        if lin:
            cab = f"*Semana {_sem(n)}" + (f" · {_dm(ini_sem(n))} a {_dm(ini_sem(n) + dt.timedelta(days=6))}*" if ini else "*")
            blocos_sem.append([cab] + lin)
    tem_amrap = any(t == "amrap" for n in sel for _, acs in por_sem.get(n, []) for t, _, _ in acs)
    if linhas_fixas or blocos_sem:
        L.append("")
        L.append("*O que eu preciso de você:*")
        if linhas_fixas:
            L.append("*Todas as semanas*" if not parcial else "*Nas duas semanas*" if len(sel) == 2
                     else f"*Nas semanas {_lista([_sem(s) for s in sel])}*")
            L.extend(linhas_fixas)
        for b in blocos_sem:
            L.append(b[0])
            L.extend(b[1:])
        if tem_amrap:
            L.append("O AMRAP é o teste que atualiza o seu 1RM: faça o máximo de repetições com boa técnica.")

    # ---- vídeos de demonstração e próxima revisão
    links = any(r.get("link") for n in sel for s in D.get("semanas", []) if int(s["numero"]) == n
                for rows in s["treinos"].values() for r in rows) or any(a.get("link") for a in D.get("aqc", []))
    fim_txt = []
    if links:
        fim_txt.append("Os exercícios com *Ver vídeo* na ficha têm o vídeo de demonstração.")
    if proxima and _data(proxima):
        fim_txt.append(f"Próxima revisão: {_data(proxima):%d/%m}.")
    if fim_txt:
        L.append("")
        L.extend(fim_txt)
    return re.sub(r"\n{3,}", "\n\n", "\n".join(x for x in L if x is not None)).strip()


# ------------------------------------------------------------------ resultados
def texto_resultados(nome: str, bloco: int, rm_evolucao: list[dict] | None = None, itens: list[dict] | None = None,
                     proximo_inicio: str | None = None) -> str:
    """Só aponta para o PDF (Miguel, 11/out/2026: o PDF é o principal; a mensagem não repete o que está nele, então
    1RM e itens pessoais ficam só no PDF). rm_evolucao e itens ficam na assinatura por compatibilidade, sem uso."""
    d = _data(proximo_inicio)
    L = [f"*{_primeiro(nome)}, fechamos o bloco {_sem(bloco)}.*", "",
         "Segue o PDF com a sua evolução. Abra para ver os seus números e o que mudou no bloco."]
    if d:
        L.append(f"O próximo bloco começa em {d:%d/%m}.")
    return "\n".join(L).strip()


# ------------------------------------------------------------------ 1RM por bloco direto da aba prs (uso fora da Central)
def rm_evolucao_planilha(path: str) -> list[dict]:
    sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
    import openpyxl
    import extrair_dados as X
    wb = openpyxl.load_workbook(path, data_only=True)
    if "prs" not in wb.sheetnames:
        return []
    wp, por = wb["prs"], {}
    for b, r0 in sorted(X.PRS_LINHA0.items()):
        for r in range(r0, r0 + 7):
            nm, reps, kg = wp.cell(r, 1).value, X.numero(wp.cell(r, 2).value), X.numero(wp.cell(r, 3).value)
            if nm and reps and kg:
                est = X.numero(wp.cell(r, 4).value) or kg / X.brzycki_frac(reps)
                por.setdefault(str(nm).strip(), []).append({"bloco": b, "est": round(est, 1), "reps": reps, "kg": kg})
    return [{"exercicio": n, "blocos": pts, "variacao": round((pts[-1]["est"] - pts[0]["est"]) / pts[0]["est"] * 100, 1)
             if len(pts) > 1 else None} for n, pts in por.items()]


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="tipo", required=True)
    a = sub.add_parser("entrega")
    a.add_argument("dados")
    a.add_argument("--semanas", default="")
    a.add_argument("--proxima")
    a.add_argument("--planejamento", action="store_true")
    r = sub.add_parser("resultados")
    r.add_argument("planilha")
    r.add_argument("--bloco", type=int, required=True)
    r.add_argument("--itens")
    r.add_argument("--proximo")
    r.add_argument("--nome")
    x = ap.parse_args()
    if x.tipo == "entrega":
        D = json.load(open(x.dados, encoding="utf-8"))
        sem = [int(s) for s in x.semanas.split(",") if s.strip().isdigit()]
        print(texto_entrega(D, sem, x.proxima, x.planejamento))
    else:
        nome = x.nome
        if not nome:
            import openpyxl
            wb = openpyxl.load_workbook(x.planilha, data_only=True)
            aba = next((n for n in wb.sheetnames if n.lower().startswith("aluno")), None)
            nome = str(wb[aba]["C2"].value or "") if aba else ""
        itens = json.load(open(x.itens, encoding="utf-8")) if x.itens else []
        print(texto_resultados(nome, x.bloco, rm_evolucao_planilha(x.planilha), itens, x.proximo))


if __name__ == "__main__":
    main()
