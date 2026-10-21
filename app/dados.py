# -*- coding: utf-8 -*-
"""Aba DADOS — leitura de TODAS as semanas da planilha matriz e cálculo de VTT, VTR e intensidade.

Definições idênticas às da aba MACRO da planilha:
  VTT (por semana)  = Σ séries × repetições × kg   (SUMPRODUCT das três colunas; células não numéricas contam 0)
  VTR (por semana)  = Σ séries × repetições
  Intensidade       = média simples do %1RM das linhas com %1RM > 0 (AVERAGEIF ">0")
Linhas com faixa ("8 - 12"), AMRAP ou tempo ("30s") somam ZERO em VTT/VTR, como na macro (regra de Miguel).

Só o planejado (Miguel não registra mais o treino realizado no app). Nada aqui altera a planilha nem a prescrição.
"""
import datetime as dt
import re
import sys
from pathlib import Path

import openpyxl

from . import armazenamento as A
from .config import SCRIPTS

sys.path.insert(0, str(SCRIPTS))
import extrair_dados as X  # noqa: E402  (funções de leitura do motor, sem alterá-lo)

LIMITE_AUMENTO = 25.0      # % de aumento de VTT, VTR ou intensidade semana a semana que gera alerta (Miguel: 10% → 25% em 8/out/2026)
DIAS_ORDEM = ["SEGUNDA", "TERÇA", "QUARTA", "QUINTA", "SEXTA", "SÁBADO", "DOMINGO"]
DIA_CURTO = {"SEGUNDA": "Segunda", "TERÇA": "Terça", "QUARTA": "Quarta", "QUINTA": "Quinta", "SEXTA": "Sexta",
             "SÁBADO": "Sábado", "DOMINGO": "Domingo"}
_cache: dict = {}


# ------------------------------------------------------------------ utilidades
def chave_ex(nome: str) -> str:
    return X.chave_ex(nome or "")


def base_de(nome: str) -> str | None:
    n = X.sem_acento(nome or "")
    for k, lab in X.BASES:
        if k in n:
            return lab
    return None


def _num_fixo(v):
    """Número como a SUMPRODUCT enxerga: só números de verdade (texto, faixa, AMRAP = None)."""
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", ".")
    return float(s) if re.fullmatch(r"\d+(\.\d+)?", s) else None


def _data(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    return None


def _fmt(v):
    v = X.reps_texto(v)   # "10/10" que o Excel virou data
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


# ------------------------------------------------------------------ leitura da planilha
VERSAO_CACHE = 4


def ler_planilha(path: Path) -> dict:
    """Prescrição planejada de todos os blocos/semanas.
    Guardada em cache na memória e em disco (<planilha>.dados.json), porque ler o .xlsx leva ~1 s por aluno."""
    st = path.stat()
    marca = [VERSAO_CACHE, st.st_size, int(st.st_mtime)]
    if str(path) in _cache and _cache[str(path)][0] == marca:
        return _cache[str(path)][1]
    disco = path.with_name(path.name + ".dados.json")
    salvo = A.ler_json(disco)
    if salvo and salvo.get("marca") == marca:
        out = salvo["dados"]
        out["rm"] = {int(k): v for k, v in out["rm"].items()}
        out["rm_ref"] = {int(k): v for k, v in out.get("rm_ref", {}).items()}
        _cache[str(path)] = (marca, out)
        return out
    out = _ler_xlsx(path)
    try:
        A.gravar_json(disco, {"marca": marca, "dados": out})
    except OSError:
        pass
    _cache[str(path)] = (marca, out)
    return out


def _ler_xlsx(path: Path) -> dict:
    wbv = openpyxl.load_workbook(path, data_only=True)
    wbf = openpyxl.load_workbook(path)
    nomes = wbv.sheetnames
    blocos = sorted(int(m.group(1)) for n in nomes if (m := re.fullmatch(r"\s*bloco\s*(\d+)\s*", n, flags=re.I)))
    out = {"semanas": [], "rm": {}, "rm_ref": {}, "treinos": {}, "avisos": []}
    data_b1 = _data(wbv["macro"].cell(6, 2).value) if "macro" in nomes else None

    for b in blocos:
        bsheet = next(n for n in nomes if re.fullmatch(rf"\s*bloco\s*0?{b}\s*", n, flags=re.I))
        asheet = f"aluno {b:02d}" if f"aluno {b:02d}" in nomes else "aluno 01"
        ws = wbv[bsheet]
        linhas = X.detectar_treinos(ws, wbf[bsheet])

        # dias da semana do bloco: {'T01': 'SEGUNDA'}
        dias = {}
        if asheet in nomes:
            for c in "BCDEF":
                d = X.resolver(wbf, wbv, asheet, c + "6") or X.resolver(wbf, wbv, "aluno 01", c + "6")
                t = X.resolver(wbf, wbv, asheet, c + "5") or X.resolver(wbf, wbv, "aluno 01", c + "5")
                if d and t and re.fullmatch(r"T\d+", str(d).strip().upper()):
                    dias[str(d).strip().upper()] = str(t).strip().upper()

        # 1RM de referência do bloco (aba prs) — o mesmo que a planilha usa no %1RM
        rm, rm_ref = {}, {}
        r0 = X.PRS_LINHA0.get(b)
        if r0 and "prs" in nomes:
            wp = wbv["prs"]
            for r in range(r0, r0 + 7):
                nm, reps, kg = wp.cell(r, 1).value, X.numero(wp.cell(r, 2).value), X.numero(wp.cell(r, 3).value)
                if nm and reps and kg:
                    cache = X.numero(wp.cell(r, 4).value)
                    rm[str(nm).strip()] = cache if cache else kg / X.brzycki_frac(reps)
                    rm_ref[str(nm).strip()] = {"reps": reps, "kg": kg, "est": rm[str(nm).strip()]}
        out["rm"][b] = rm
        out["rm_ref"][b] = rm_ref

        # data de início do bloco: macro (B6/F6/J6), senão bloco 01 + 28 dias × (b-1)
        mcol = X.MACRO_COL_INICIO.get(b)
        inicio = _data(wbv["macro"].cell(6, mcol).value) if (mcol and "macro" in nomes) else None
        if not inicio and data_b1:
            inicio = data_b1 + dt.timedelta(days=28 * (b - 1))

        for w, (ce, c0) in enumerate(X.colunas_semanas(ws)):   # layout antigo (C/I/O/U) ou novo (B/C, I/J, P/Q, W/X)
            treinos = {}
            for num, (a0, b0) in linhas.items():
                rows = []
                for r in range(a0, b0 + 1):
                    ex = ws.cell(r, ce).value
                    if X.vazio(ex):
                        continue
                    sets, reps, kg, pct = (ws.cell(r, c0 + i).value for i in range(4))
                    obs = ws.cell(r, c0 + 4).value
                    if X.vazio(sets) and X.vazio(reps):
                        continue
                    ex = re.sub(r"\s+", " ", str(ex)).strip()
                    base = base_de(ex)
                    pct_n = X.numero(pct)
                    kg_n = _num_fixo(kg)
                    if pct_n is None and base in rm and kg_n:
                        pct_n = kg_n / rm[base]
                    rows.append({"linha": r, "ex": ex, "base": base, "series": _fmt(sets), "reps": _fmt(reps),
                                 "kg": _fmt(kg), "sets_n": _num_fixo(sets), "reps_n": _num_fixo(reps), "kg_n": kg_n,
                                 "pct": pct_n if (pct_n and pct_n > 0) else None, "obs": _fmt(obs)})
                if rows:
                    treinos[num] = rows
            if not treinos:
                continue
            data = inicio + dt.timedelta(days=7 * w) if inicio else None
            out["semanas"].append({"bloco": b, "semana": w + 1, "global": (b - 1) * 4 + w + 1,
                                   "data": data.isoformat() if data else "", "dias": dias, "treinos": treinos})
            for num in treinos:
                out["treinos"][num] = dias.get(f"T{num}", out["treinos"].get(num, ""))
    return out


# ------------------------------------------------------------------ cálculo
FAIXAS = ["abaixo de 70%", "70 a 80%", "80 a 90%", "90% ou mais"]


def _faixa(pct: float) -> int:
    p = round(pct * 100, 1)
    return 0 if p < 70 else 1 if p < 80 else 2 if p < 90 else 3


def _metricas(row):
    """VTT/VTR de uma linha como a SUMPRODUCT da macro: faixa, AMRAP e tempo valem zero."""
    s, r, k = row["sets_n"], row["reps_n"], row["kg_n"]
    vtr = s * r if (s is not None and r is not None) else 0.0
    vtt = vtr * k if k is not None else 0.0
    zonas = [0, 0, 0, 0]
    if row["pct"] and s:
        zonas[_faixa(row["pct"])] += s
    return {"vtt": vtt, "vtr": vtr, "pct": row["pct"], "zonas": zonas}


def _agrega(itens):
    pcts = [i["pct"] for i in itens if i["pct"]]
    return {"vtt": round(sum(i["vtt"] for i in itens), 1), "vtr": round(sum(i["vtr"] for i in itens), 1),
            "int": round(sum(pcts) / len(pcts) * 100, 2) if pcts else None,
            "zonas": [sum(i["zonas"][k] for i in itens) for k in range(4)]}


def _var(atual, anterior):
    if atual is None or anterior in (None, 0):
        return None
    return round((atual - anterior) / anterior * 100, 1)


def _com_variacao(serie: list[dict]):
    """Acrescenta variação % semana a semana e o alerta de aumento > LIMITE_AUMENTO (25%; VTT, VTR ou intensidade média).
    Intensidade: aumento relativo da média do %1RM (ex.: 60% → 66,1% = +10,2%), como no VTT e no VTR."""
    ant = None
    for p in serie:
        if ant is None or p.get("vazio"):
            p.update(d_vtt=None, d_vtr=None, d_int=None, d_int_pp=None, alerta=False)
        else:
            p["d_vtt"] = _var(p["vtt"], ant["vtt"])
            p["d_vtr"] = _var(p["vtr"], ant["vtr"])
            p["d_int"] = _var(p["int"], ant["int"])
            p["d_int_pp"] = round(p["int"] - ant["int"], 2) if (p["int"] is not None and ant["int"] is not None) else None
            p["alerta"] = any(d is not None and d > LIMITE_AUMENTO for d in (p["d_vtt"], p["d_vtr"], p["d_int"]))
        if not p.get("vazio"):
            ant = p
    return serie


def analisar(planilha) -> dict:
    """Aba DADOS de uma planilha (Path) ou do histórico do aluno (o dict de historico())."""
    P = planilha if isinstance(planilha, dict) else ler_planilha(planilha)
    hoje = dt.date.today()
    semanas_out, por_ex = [], {}
    for s in P["semanas"]:
        itens = []
        for num in sorted(s["treinos"]):
            for row in s["treinos"][num]:
                m = _metricas(row)
                itens.append(m)
                d = por_ex.setdefault(chave_ex(row["ex"]), {"nome": row["ex"], "base": row["base"], "semanas": {}})
                d["semanas"].setdefault(s["global"], []).append(m)
        data = dt.date.fromisoformat(s["data"]) if s["data"] else None
        semanas_out.append({"global": s["global"], "rotulo": f"S{s['global']:02d}", "bloco": s["bloco"], "semana": s["semana"],
                            "ciclo": s.get("ciclo", 1), "rotulo_bloco": s.get("rotulo_bloco") or f"Bloco {s['bloco']:02d}",
                            "data": s["data"], "atual": bool(data and data <= hoje < data + dt.timedelta(days=7)),
                            "treinos": len(s["treinos"]), **_agrega(itens)})
    _com_variacao(semanas_out)

    exercicios = []
    for k, d in por_ex.items():
        pts = [{"g": x["global"], **_agrega(d["semanas"][x["global"]])} if x["global"] in d["semanas"]
               else {"g": x["global"], "vtt": 0, "vtr": 0, "int": None, "zonas": [0, 0, 0, 0], "vazio": True}
               for x in semanas_out]
        exercicios.append({"chave": k, "nome": d["nome"], "base": d["base"], "semanas": _com_variacao(pts)})
    ordem = {}
    for s in P["semanas"]:
        for t in sorted(s["treinos"]):
            for i, r in enumerate(s["treinos"][t]):
                ordem.setdefault(chave_ex(r["ex"]), (int(t), i))
    # principais (com 1RM de referência) primeiro, depois na ordem em que aparecem nos treinos
    exercicios.sort(key=lambda e: (e["base"] is None, ordem.get(e["chave"], (99, 99))))

    # evolução do 1RM estimado entre os blocos (aba prs); no histórico, bloco a bloco de todos os ciclos
    refs = P.get("rm_ciclos") or {(1, b): v for b, v in P["rm_ref"].items()}
    multi = len({c for c, _ in refs}) > 1
    nomes_rm = []
    for k in sorted(refs):
        for n in refs[k]:
            if n not in nomes_rm:
                nomes_rm.append(n)
    rm_evolucao = []
    for n in nomes_rm:
        pts = [{"bloco": b, "ciclo": c, "rotulo": rotulo_bloco(c, b, multi), "est": round(refs[(c, b)][n]["est"], 1),
                "reps": refs[(c, b)][n]["reps"], "kg": refs[(c, b)][n]["kg"]} for c, b in sorted(refs) if n in refs[(c, b)]]
        var = round((pts[-1]["est"] - pts[0]["est"]) / pts[0]["est"] * 100, 1) if len(pts) > 1 else None
        rm_evolucao.append({"exercicio": n, "blocos": pts, "variacao": var})

    return {"semanas": semanas_out, "exercicios": exercicios, "limite": LIMITE_AUMENTO, "faixas": FAIXAS,
            "rm_evolucao": rm_evolucao, "rm": {str(b): {k: round(v, 1) for k, v in r.items()} for b, r in P["rm"].items()}}


def planilha_atual(aluno: dict) -> Path | None:
    """A planilha como Miguel enviou (a original, sem o aquecimento da Central): VTT, VTR e %1RM contam só as séries
    que ele prescreveu e batem com a aba MACRO dele (Miguel, 10/out/2026). Envios antigos sem "original" usam o arquivo."""
    for p in aluno.get("planilhas", []):
        caminho = A.pasta(aluno["slug"]) / (p.get("original") or p["arquivo"])
        if not caminho.exists():
            caminho = A.pasta(aluno["slug"]) / p["arquivo"]
        if caminho.exists() and p.get("status") != "Erro na leitura":
            return caminho
    return None


# ------------------------------------------------------------------ histórico do aluno (todas as planilhas)
# Miguel (21/out/2026): a planilha matriz é reenviada ao longo do tempo e, a cada trimestre, ele começa uma planilha
# NOVA (bloco 01 de novo). O Resultados, a aba Dados e o feedback leem o histórico do aluno, não só a última planilha:
#   - cada semana é guardada pela DATA de início; a planilha mais nova vale para as datas que ela cobre, e as semanas
#     de fora (de planilhas anteriores) continuam valendo (uma planilha antiga reenviada não apaga o bloco seguinte);
#   - planilha cujo bloco 01 começa mais de 3 semanas depois do início do ciclo anterior abre um novo ciclo;
#   - a semana do ciclo é a contagem real das semanas (bloco de 3 semanas não deixa buraco na numeração).
NOVO_CICLO_DIAS = 21
_hist: dict = {}


def rotulo_bloco(ciclo: int, bloco: int, multi: bool) -> str:
    return f"Ciclo {ciclo} · Bloco {bloco:02d}" if multi else f"Bloco {bloco:02d}"


def _caminho(aluno: dict, p: dict) -> Path | None:
    caminho = A.pasta(aluno["slug"]) / (p.get("original") or p["arquivo"])
    if not caminho.exists():
        caminho = A.pasta(aluno["slug"]) / p["arquivo"]
    return caminho if caminho.exists() and p.get("status") != "Erro na leitura" else None


def historico(aluno: dict, sem_envio: str | None = None) -> dict:
    """O mesmo formato de ler_planilha() com as semanas de todas as planilhas do aluno, mais "ciclo", "envio" e
    "rotulo_bloco" em cada semana, "rm_ciclos" {(ciclo, bloco): 1RM} e "ciclo_envio" {envio: ciclo}.
    "rm"/"rm_ref" ficam com os blocos do ciclo mais recente (o da planilha atual). sem_envio: histórico sem esse envio."""
    envios = [(p, c) for p in reversed(aluno.get("planilhas", [])) if p["id"] != sem_envio and (c := _caminho(aluno, p))]
    if not envios:
        return {"semanas": [], "rm": {}, "rm_ref": {}, "rm_ciclos": {}, "treinos": {}, "avisos": [], "ciclo_envio": {}, "ciclos": 0}
    marca = tuple((p["id"], c.stat().st_size, int(c.stat().st_mtime)) for p, c in envios)
    chave = aluno["slug"]
    if chave in _hist and _hist[chave][0] == marca:
        return _hist[chave][1]
    lidos = [(p, ler_planilha(c)) for p, c in envios]          # da mais antiga para a mais recente
    if any(not s["data"] for _, P in lidos for s in P["semanas"]):
        # planilha sem datas na aba macro: não dá para alinhar as semanas, vale a mais recente (como antes)
        p, P = lidos[-1]
        out = {**P, "semanas": [{**s, "ciclo": 1, "envio": p["id"], "rotulo_bloco": f"Bloco {s['bloco']:02d}"} for s in P["semanas"]],
               "rm_ciclos": {(1, b): v for b, v in P["rm_ref"].items()}, "ciclo_envio": {x["id"]: 1 for x, _ in lidos}, "ciclos": 1}
        _hist[chave] = (marca, out)
        return out
    # ciclos: inícios de planilha (1ª semana) agrupados; mais de 3 semanas depois do início do ciclo = ciclo novo
    inicios = sorted({P["semanas"][0]["data"] for _, P in lidos if P["semanas"]})
    grupos = []
    for d in inicios:
        if not grupos or (dt.date.fromisoformat(d) - dt.date.fromisoformat(grupos[-1])).days > NOVO_CICLO_DIAS:
            grupos.append(d)

    def ciclo_de(d):
        return max(i for i, g in enumerate(grupos, 1) if g <= d)

    por_data, rm_ciclos, ciclo_envio, treinos = {}, {}, {}, {}
    for p, P in lidos:
        if not P["semanas"]:
            continue
        c = ciclo_de(P["semanas"][0]["data"])
        ciclo_envio[p["id"]] = c
        d0, d1 = P["semanas"][0]["data"], P["semanas"][-1]["data"]
        # a planilha mais nova vale para as datas que ela cobre e, no mesmo ciclo, para tudo até a última semana dela
        # (data de início corrigida na macro não deixa semana velha sobrando); o que vem depois continua valendo
        for d in [d for d, w in por_data.items() if d0 <= d <= d1 or (w["ciclo"] == c and d <= d1)]:
            del por_data[d]
        for s in P["semanas"]:
            por_data[s["data"]] = {**s, "ciclo": c, "envio": p["id"]}
        for b, v in P["rm_ref"].items():
            if v:                    # aba prs vazia nesta planilha não apaga o 1RM que outra já trazia
                rm_ciclos[(c, b)] = v
        treinos.update(P["treinos"])
    semanas = [por_data[d] for d in sorted(por_data)]
    multi = len({s["ciclo"] for s in semanas}) > 1
    for i, s in enumerate(semanas, 1):
        s["global"] = i
        s["rotulo_bloco"] = rotulo_bloco(s["ciclo"], s["bloco"], multi)
    ult = ciclo_envio.get(lidos[-1][0]["id"], len(grupos))
    P_ult = lidos[-1][1]
    out = {"semanas": semanas, "rm": P_ult["rm"], "rm_ref": P_ult["rm_ref"], "rm_ciclos": rm_ciclos, "treinos": treinos,
           "avisos": [], "ciclo_envio": ciclo_envio, "ciclos": len(grupos), "ciclo_atual": ult}
    _hist[chave] = (marca, out)
    return out


def analisar_aluno(aluno: dict) -> dict:
    return analisar(historico(aluno))


def nome_arquivo(nome: str, bloco: int, ciclo: int = 1) -> str:
    """'MariaSilva_Bloco02'; a partir do 2º ciclo 'MariaSilva_C2_Bloco02' (nunca substitui o bloco do ciclo anterior)."""
    return re.sub(r'[\\/:*?"<>|\s]+', "", nome) + (f"_C{ciclo}" if ciclo and ciclo > 1 else "") + f"_Bloco{int(bloco):02d}"

