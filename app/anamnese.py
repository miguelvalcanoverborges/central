# -*- coding: utf-8 -*-
"""Leitura da planilha de respostas da ANAMNESE (Google Forms), seguindo a skill consultoria-anamnese:
- várias pessoas e respostas repetidas na mesma planilha → agrupa por nome completo normalizado;
- a resposta mais recente é a base; as anteriores entram como histórico com data;
- perguntas encontradas pelo TEXTO do cabeçalho (o formulário mudou ao longo do tempo);
- em branco ou "--" = dado faltante;
- contatos (telefone, e-mail etc.) nunca são copiados.
"""
import csv
import datetime as dt
import io
import re
import unicodedata

import openpyxl

CONTATO = ("telefone", "celular", "whatsapp", "e-mail", "email", "endereco de e-mail", "instagram",
           "endereco", "cpf", "rg ")
VAZIOS = {"", "--", "-", "—", "n/a"}


def _sa(s) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", str(s)) if unicodedata.category(c) != "Mn").lower()


def _norm_nome(n: str) -> str:
    return re.sub(r"\s+", " ", _sa(n)).strip()


def _vazio(v) -> bool:
    return v is None or str(v).strip().lower() in VAZIOS


def _fmt(v) -> str:
    if isinstance(v, dt.datetime):
        return v.strftime("%d/%m/%Y %H:%M") if (v.hour or v.minute) else v.strftime("%d/%m/%Y")
    if isinstance(v, dt.date):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _data(v):
    if isinstance(v, dt.datetime):
        return v
    for f in ("%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S"):
        try:
            return dt.datetime.strptime(str(v).strip(), f)
        except (ValueError, TypeError):
            pass
    return None


def _tabelas(arquivo_bytes: bytes, nome_arquivo: str) -> list[list[list]]:
    """Todas as abas de respostas ("Respostas ao formulário 1", "… 2" …): o Forms cria uma aba nova quando o
    formulário é refeito, e as respostas antigas continuam valendo como histórico."""
    if nome_arquivo.lower().endswith(".csv"):
        txt = arquivo_bytes.decode("utf-8-sig", errors="replace")
        return [list(csv.reader(io.StringIO(txt)))]
    wb = openpyxl.load_workbook(io.BytesIO(arquivo_bytes), data_only=True, read_only=True)
    abas = [n for n in wb.sheetnames if "respostas" in _sa(n)] or [wb.sheetnames[0]]
    return [[list(r) for r in wb[n].iter_rows(values_only=True)] for n in abas]


def ler(arquivo_bytes: bytes, nome_arquivo: str) -> dict:
    """Retorna {nome_normalizado: {"nome": str, "respostas": [ {data, data_txt, pares:[(pergunta, resposta)]} ]}}
    com as respostas de cada pessoa da MAIS RECENTE para a mais antiga (todas as abas de respostas)."""
    pessoas: dict = {}
    for tabela in _tabelas(arquivo_bytes, nome_arquivo):
        _ler_tabela(tabela, pessoas)
    for p in pessoas.values():
        p["respostas"].sort(key=lambda x: x["data"], reverse=True)
    return pessoas


def _ler_tabela(tabela: list[list], pessoas: dict):
    linhas = [r for r in tabela if any(not _vazio(c) for c in r)]
    if len(linhas) < 2:
        return
    cab = [_fmt(c) if c is not None else "" for c in linhas[0]]
    # cabeçalhos duplicados ("5.4" usado duas vezes, dois campos de medicamento) ganham sufixo para não se sobrescreverem
    vistos, cab_u = {}, []
    for c in cab:
        vistos[c] = vistos.get(c, 0) + 1
        cab_u.append(c if vistos[c] == 1 else f"{c} ({vistos[c]})")
    i_data = next((i for i, c in enumerate(cab) if "carimbo" in _sa(c) or "timestamp" in _sa(c)), None)
    i_nome = next((i for i, c in enumerate(cab) if "nome" in _sa(c) and "completo" in _sa(c)), None)
    if i_nome is None:
        i_nome = next((i for i, c in enumerate(cab) if "nome" in _sa(c)), None)
    if i_nome is None:
        raise ValueError("Não achei a coluna de nome na planilha (cabeçalho com a palavra 'nome').")
    for r in linhas[1:]:
        r = list(r) + [None] * (len(cab) - len(r))
        if _vazio(r[i_nome]):
            continue
        nome = re.sub(r"\s+", " ", str(r[i_nome])).strip()
        data = _data(r[i_data]) if i_data is not None else None
        pares = []
        for i, (h, v) in enumerate(zip(cab_u, r)):
            if i in (i_data, i_nome) or _vazio(v) or not h:
                continue
            if any(k in _sa(h) for k in CONTATO):
                continue
            pares.append((h.strip(), _fmt(v)))
        p = pessoas.setdefault(_norm_nome(nome), {"nome": nome.title() if nome.isupper() or nome.islower() else nome,
                                                   "respostas": []})
        p["respostas"].append({"data": data.isoformat() if data else "",
                               "data_txt": data.strftime("%d/%m/%Y") if data else "sem data",
                               "pares": pares})


def resumo_pessoas(pessoas: dict) -> list[dict]:
    out = [{"chave": k, "nome": v["nome"], "respostas": len(v["respostas"]),
            "mais_recente": v["respostas"][0]["data_txt"], "_ord": v["respostas"][0]["data"]}
           for k, v in pessoas.items()]
    return sorted(out, key=lambda x: x["_ord"], reverse=True)


def como_texto(pessoa: dict) -> str:
    """Texto pergunta → resposta, para leitura humana e para a IA (resposta base + histórico datado)."""
    blocos = []
    for i, r in enumerate(pessoa["respostas"]):
        titulo = "RESPOSTA BASE (mais recente)" if i == 0 else "HISTÓRICO"
        blocos.append(f"### {titulo} — {r['data_txt']}")
        blocos += [f"- {h}: {v}" for h, v in r["pares"]]
        blocos.append("")
    return "\n".join(blocos)


def mudancas(pessoa: dict) -> list[str]:
    """Aponta perguntas cuja resposta mudou entre a resposta base e a anterior."""
    if len(pessoa["respostas"]) < 2:
        return []
    atual = dict(pessoa["respostas"][0]["pares"])
    ant = pessoa["respostas"][1]
    out = []
    for h, v in ant["pares"]:
        if h in atual and atual[h] != v:
            out.append(f"{h}: {ant['data_txt']} → \"{v}\" · agora → \"{atual[h]}\"")
    return out
