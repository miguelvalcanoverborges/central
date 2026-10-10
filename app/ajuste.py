# -*- coding: utf-8 -*-
"""Ajuste automático da planilha que Miguel envia (decisão de Miguel, 9/out/2026).

A prescrição e a periodização são de Miguel: a planilha chega totalmente planejada. Ao recebê-la, a Central grava uma
CÓPIA AJUSTADA (a original fica guardada, intacta) com as regras dele, e o PDF e a aba Dados saem dessa cópia:

1. Aquecimento antes da série principal dos exercícios principais (os da aba prs com 1RM no quadro do bloco):
   65% × 3, 70% × 2 e 80% × 1 (1 série cada; reps caindo, Miguel 10/out/2026), só os degraus abaixo da carga principal;
   principal acima de 90% ganha também 90% × 1. kg = %1RM × 1RM, no múltiplo de 1 kg mais próximo. Se Miguel já escreveu mais de uma linha não-AMRAP
   para o exercício naquele treino e semana (ex.: o próprio aquecimento), nada é inserido ali.
   A Central abre espaço descendo os exercícios de baixo; se faltar linha no treino, ficam os degraus mais próximos
   da carga e o aviso aparece em "Ajustes da Central".
2. O AMRAP fica exatamente como Miguel escreveu na planilha (ele indica quais séries são AMRAP; 9/out/2026).
3. GRAVAR (AMRAP e %1RM ≥ 80%, nunca no aquecimento), ANOTAR (faixa de repetições e AMRAP) e as etiquetas
   Aquecimento/Principal continuam automáticas no motor, na hora do PDF.

A cópia é escrita célula a célula no XML (`xlsx_celulas`): fórmulas, aba macro, gráficos e formatação ficam intactos;
o resultado das fórmulas de %1RM das linhas mexidas é atualizado (kg ÷ 1RM) para o motor e a aba Dados lerem certo.
"""
import math
import re
import sys
from pathlib import Path

import openpyxl

from . import xlsx_celulas
from .config import SCRIPTS

sys.path.insert(0, str(SCRIPTS))
import extrair_dados as X  # noqa: E402  (funções de leitura do motor, sem alterá-lo)

ARREDONDAR_KG = 1.0                 # Miguel (7/out/2026): carga calculada no múltiplo de 1 kg mais próximo
AQUECIMENTO = (.65, .70, .80)       # Miguel (7/out/2026): 1 série em cada degrau abaixo da carga principal
DEGRAU_EXTRA = .90                  # principal acima de 90% ganha também 90%
REPS_DEGRAU = {.65: 3, .70: 2, .80: 1, .90: 1}   # Miguel (10/out/2026): reps caindo, menos volume nas semanas intensas
L = xlsx_celulas.col_letras


class Erro(ValueError):
    """Mensagem pronta para mostrar a Miguel."""


# ------------------------------------------------------------------ utilidades
def _aba(nomes, tipo: str, n: int) -> str | None:
    return next((s for s in nomes if re.fullmatch(rf"\s*{tipo}\s*0*{n}\s*", s, flags=re.I)), None)


def _blocos(nomes) -> list[int]:
    return sorted(int(m.group(1)) for s in nomes if (m := re.fullmatch(r"\s*bloco\s*(\d+)\s*", s, flags=re.I)))


def _formula(wbf, aba: str, coord: str) -> bool:
    v = wbf[aba][coord].value
    return (isinstance(v, str) and v.startswith("=")) or type(v).__name__ in ("ArrayFormula", "DataTableFormula")


def _base(ex: str) -> str | None:
    n = X.sem_acento(ex or "")
    return next((lab for k, lab in X.BASES if k in n), None)


def _bruto(v):
    if v is None or isinstance(v, bool):
        return v
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, (int, float)):
        return v
    s = str(v).strip()
    return s or None


def _num(v):
    n = X.numero(v)
    return None if n is None or isinstance(n, bool) else float(n)


def arredondar(kg: float) -> float:
    q = math.floor(kg / ARREDONDAR_KG + 0.5) * ARREDONDAR_KG
    return int(q) if float(q).is_integer() else round(q, 2)


def _pct_txt(p: float) -> str:
    return f"{round(p * 100)}%"


def _link(h):
    t = (h.target or h.location or "") if h is not None else ""
    return t.strip() if str(t).strip().lower().startswith("http") else None


def _rm_bloco(wbv, n: int) -> dict[str, float]:
    """{'Agachamento': 112.5} — exercícios da aba prs do bloco n com 1RM (os principais, regra de Miguel)."""
    r0 = X.PRS_LINHA0.get(n)
    out = {}
    if not r0 or "prs" not in wbv.sheetnames:
        return out
    wp = wbv["prs"]
    for r in range(r0, r0 + 7):
        nome = wp.cell(r, 1).value
        if X.vazio(nome):
            continue
        reps, kg, cache = _num(wp.cell(r, 2).value), _num(wp.cell(r, 3).value), _num(wp.cell(r, 4).value)
        rm = cache if cache else (kg / X.brzycki_frac(reps) if reps and kg else None)
        if rm:
            out[str(nome).strip()] = rm
    return out


def _principal(nome: str, rms: dict[str, float]) -> str | None:
    """Nome da aba prs a que o exercício corresponde (mesma correspondência da fórmula de %1RM da planilha)."""
    por = {X.sem_acento(k): k for k in rms}
    b = _base(nome)
    if b and X.sem_acento(b) in por:
        return por[X.sem_acento(b)]
    n = X.sem_acento(nome or "")
    return next((v for k, v in por.items() if k and k in n), None)


def _layout(wbv, wbf, aba: str, colunas) -> list[tuple[str, list[int]]]:
    """[('01', [6, …, 15]), …]: linhas de exercício de cada treino (até a última com fórmula de %1RM ou exercício)."""
    ws = wbv[aba]
    heads = []
    for r in range(1, 220):
        m = re.fullmatch(r"\s*TREINO\s*(\d+)\s*", str(ws.cell(r, 1).value or ""), flags=re.I)
        if m:
            heads.append((f"{int(m.group(1)):02d}", r))
    out = []
    for i, (num, h) in enumerate(heads):
        r0 = h + 3
        limite = heads[i + 1][1] - 1 if i + 1 < len(heads) else r0 + 14
        usadas = [r for r in range(r0, limite + 1)
                  if any(not X.vazio(ws.cell(r, ce).value) or _formula(wbf, aba, f"{L(c0 + 3)}{r}")
                         for ce, c0 in colunas)]
        fim = usadas[-1] if usadas else min(limite - 2, r0 + 9)
        out.append((num, list(range(r0, fim + 1))))
    return out


def _eh_amrap(x: dict | None) -> bool:
    return bool(x) and "amrap" in str(x.get("reps") or "").lower()


def _eh_aquec(x: dict | None) -> bool:
    return bool(x) and "aquecimento" in str(x.get("obs") or "").lower()


# ------------------------------------------------------------------ ajuste de um bloco
def _ajustar_bloco(wbv, wbf, n: int, val: dict, cache: dict, ajustes: list, avisos: list) -> bool:
    bs = _aba(wbv.sheetnames, "bloco", n)
    ws = wbv[bs]
    colunas = X.colunas_semanas(ws)
    so_b = all(ce == 2 for ce, _ in colunas)
    rms = _rm_bloco(wbv, n)
    if not rms:
        return False
    nsem = len(colunas)

    def nome_linha(r):
        for ce, _ in colunas:
            v = ws.cell(r, ce).value
            if not X.vazio(v):
                return re.sub(r"\s+", " ", str(v)).strip()
        return None

    def ler(r) -> list:
        """A linha por semana: {ex, sets, reps, kg, pct, obs} (None onde a semana está vazia)."""
        out = []
        for ce, c0 in colunas:
            ex = ws.cell(r, ce).value
            sets, reps, kg, pct = (ws.cell(r, c0 + i).value for i in range(4))
            obs = _bruto(ws.cell(r, c0 + 4).value)
            link = _link(wbf[bs].cell(r, c0 + 4).hyperlink)
            if link and link not in str(obs or ""):
                obs = f"{obs} - {link}" if obs else link
            if X.vazio(sets) and X.vazio(reps) and X.vazio(kg) and X.vazio(obs):
                out.append({"ex": _bruto(ex), "vazia": True} if not X.vazio(ex) else None)
                continue
            out.append({"ex": _bruto(ex) if not X.vazio(ex) else None, "sets": _bruto(sets), "reps": _bruto(reps),
                        "kg": _bruto(kg), "pct": _num(pct), "obs": obs})
        return out

    # 1) itens de cada treino: grupo do mesmo principal (linhas seguidas) ou acessório (linha inteira)
    treinos = []
    for num, vagas in _layout(wbv, wbf, bs, colunas):
        itens = []
        for r in vagas:
            nome = nome_linha(r)
            if not nome:
                continue
            prin = _principal(nome, rms)
            linha = ler(r)
            if prin and itens and itens[-1]["tipo"] == "p" and itens[-1]["base"] == prin \
                    and X.sem_acento(itens[-1]["nome"]) == X.sem_acento(nome):
                itens[-1]["linhas"].append(linha)
                continue
            if prin:
                itens.append({"tipo": "p", "nome": nome, "base": prin, "linhas": [linha]})
            else:
                itens.append({"tipo": "a", "nome": nome, "linhas": [linha]})
        treinos.append({"num": num, "vagas": vagas, "itens": itens, "mudou": False})

    # séries por semana de cada principal: [semana] -> lista de linhas (só as preenchidas naquela semana)
    for t in treinos:
        for it in t["itens"]:
            if it["tipo"] != "p":
                continue
            it["sem"] = []
            for w in range(nsem):
                it["sem"].append([dict(lin[w]) for lin in it["linhas"] if lin[w] and not lin[w].get("vazia")])

    # 2) aquecimento antes da série principal (o AMRAP é como Miguel escreveu: ele indica na planilha)
    for t in treinos:
        for it in t["itens"]:
            if it["tipo"] != "p":
                continue
            rm = rms[it["base"]]
            it["aq"] = [[] for _ in range(nsem)]
            for w in range(nsem):
                series = [x for x in it["sem"][w] if not _eh_amrap(x)]
                if len(series) != 1 or _eh_aquec(series[0]):
                    continue                                  # Miguel já estruturou (ou não há série)
                pr = series[0]
                pct = pr.get("pct") or (_num(pr.get("kg")) / rm if _num(pr.get("kg")) else None)
                if not pct:
                    continue
                pct = round(pct, 2)                           # o %1RM que Miguel planejou (a carga em kg é arredondada)
                kg_pr = _num(pr.get("kg"))
                degraus = [a for a in AQUECIMENTO if a < pct - 1e-9] + ([DEGRAU_EXTRA] if pct > DEGRAU_EXTRA + 1e-9 else [])
                for a in degraus:
                    kg = arredondar(a * rm)
                    if kg_pr and kg >= kg_pr:
                        continue                              # degrau que não fica abaixo da carga principal
                    it["aq"][w].append({"ex": pr.get("ex"), "sets": 1, "reps": REPS_DEGRAU[a], "kg": kg, "pct": round(kg / rm, 4),
                                        "obs": None, "_degrau": a})
            if any(it["aq"]):
                t["mudou"] = True

    # 3) reescreve os treinos que mudaram, abrindo espaço (descendo os de baixo)
    mexeu = False
    for t in treinos:
        if not t["mudou"]:
            continue
        cap = len(t["vagas"])

        def altura(it):
            if it["tipo"] == "a":
                return len(it["linhas"])
            return max(len(it["aq"][w]) + len(it["sem"][w]) for w in range(nsem)) if any(it["sem"]) else len(it["linhas"])

        cortados = {}
        while sum(altura(it) for it in t["itens"]) > cap:
            cand = [it for it in t["itens"] if it["tipo"] == "p" and any(it["aq"])]
            if not cand:
                break
            it = max(cand, key=lambda i: max(len(a) for a in i["aq"]))
            for w in range(nsem):
                if it["aq"][w] and len(it["aq"][w]) == max(len(a) for a in it["aq"]):
                    it["aq"][w].pop(0)                        # tira o degrau mais leve
            cortados[it["nome"]] = True
        if sum(altura(it) for it in t["itens"]) > cap:
            avisos.append(f"Bloco {n:02d} · treino {t['num']}: não coube o ajuste (o treino tem {cap} linhas). "
                          "Ficou como você escreveu.")
            continue
        for nome in cortados:
            avisos.append(f"Bloco {n:02d} · treino {t['num']} · {nome}: nem todos os degraus de aquecimento couberam; "
                          "ficaram os mais próximos da carga.")

        linhas_novas = []                                     # [(nome_b, [semana] -> dict|None)]
        for it in t["itens"]:
            if it["tipo"] == "a" or not any(it["sem"]):
                for lin in it["linhas"]:
                    linhas_novas.append((it["nome"], lin))
                continue
            h = altura(it)
            bloco = [[None] * nsem for _ in range(h)]
            for w in range(nsem):
                col = it["aq"][w] + it["sem"][w]
                for j, x in enumerate(col):                   # alinhado embaixo (a principal fica na mesma linha)
                    bloco[h - len(col) + j][w] = x
            for lin in bloco:
                linhas_novas.append((it["nome"], lin))
            if any(it["aq"]):
                sem_aq = [w + 1 for w in range(nsem) if it["aq"][w]]
                graus = sorted({x["_degrau"] for w in range(nsem) for x in it["aq"][w]})
                ajustes.append(f"Bloco {n:02d} · treino {t['num']} · {it['nome']}: aquecimento "
                               + ", ".join(_pct_txt(g) for g in graus)
                               + (f" (semana {sem_aq[0]})" if len(sem_aq) == 1 else
                                  f" (semanas {', '.join(map(str, sem_aq[:-1]))} e {sem_aq[-1]})") + ".")

        for i, r in enumerate(t["vagas"]):
            nome_b, lin = linhas_novas[i] if i < len(linhas_novas) else (None, None)
            usada = bool(lin) and any(lin)
            if so_b:
                _por(wbf, bs, val, ws, f"B{r}", nome_b if usada else None)
            for w, (ce, c0) in enumerate(colunas):
                x = lin[w] if usada else None
                if not so_b:
                    _por(wbf, bs, val, ws, f"{L(ce)}{r}", (x or {}).get("ex") or (nome_b if x else None))
                if x and x.get("vazia"):
                    x = None
                _por(wbf, bs, val, ws, f"{L(c0)}{r}", x and x.get("sets"))
                _por(wbf, bs, val, ws, f"{L(c0 + 1)}{r}", x and x.get("reps"))
                _por(wbf, bs, val, ws, f"{L(c0 + 2)}{r}", x and x.get("kg"))
                _por(wbf, bs, val, ws, f"{L(c0 + 4)}{r}", x and x.get("obs"))
                cp = f"{L(c0 + 3)}{r}"
                pct = x.get("pct") if x else None
                if _formula(wbf, bs, cp):
                    cache.setdefault(bs, {})[cp] = pct if pct is not None else ""
                else:
                    _por(wbf, bs, val, ws, cp, pct)
        mexeu = True
    return mexeu


def _por(wbf, bs, val, ws, coord, v):
    """Escreve v (célula com fórmula nunca é sobrescrita; nada a fazer se já está igual)."""
    if _formula(wbf, bs, coord):
        return
    atual = _bruto(ws[coord].value)
    if X.vazio(atual) and X.vazio(v):
        return
    if atual == v:
        return
    val.setdefault(bs, {})[coord] = v


# ------------------------------------------------------------------ entrada
def ajustar(planilha: Path) -> tuple[bytes | None, list[str], list[str]]:
    """Devolve (bytes da planilha ajustada ou None se nada mudou, ajustes feitos, avisos)."""
    try:
        wbv = openpyxl.load_workbook(planilha, data_only=True)
        wbf = openpyxl.load_workbook(planilha)
    except Exception as ex:  # noqa: BLE001
        raise Erro(f"Não consegui abrir a planilha ({ex}).") from ex
    val, cache, ajustes, avisos = {}, {}, [], []
    mexeu = False
    for n in _blocos(wbv.sheetnames):
        if n not in X.PRS_LINHA0:
            continue
        mexeu |= _ajustar_bloco(wbv, wbf, n, val, cache, ajustes, avisos)
    if not mexeu or not (any(val.values()) or any(cache.values())):
        return None, ajustes, avisos
    novo, _ = xlsx_celulas.escrever(planilha.read_bytes(), {k: v for k, v in val.items() if v},
                                    {k: v for k, v in cache.items() if v})
    return novo, ajustes, avisos
