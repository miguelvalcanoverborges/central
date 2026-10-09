# -*- coding: utf-8 -*-
"""Recebimento da planilha matriz e ponte com o motor da skill consultoria-planilha (v1.11, sem alterações).

Fluxo automático ao receber uma planilha:
  1. guarda o arquivo e roda o extrator (bloco = último com prescrição);
  2. identifica o aluno pelo nome da planilha (ou cria um novo no painel);
  3. aplica as preferências já salvas do aluno (exercícios de "Seus números") e reextrai se preciso;
     (PSE fica sempre desligada na ficha: nenhum aluno de Miguel usa PSE por enquanto)
  4. atualiza plano (2x/3x/4x), bloco atual, data da última e da próxima atualização;
  5. registra a planilha no histórico do aluno e deixa a revisão pronta para gerar o PDF.
"""
import datetime as dt
import json
import os
import re
import shutil
import subprocess
from pathlib import Path

import openpyxl

from . import planos
from . import armazenamento as A
from .config import ENTRADA, SCRIPTS, env_motor, python_motor

SEM_JANELA = 0x08000000 if os.name == "nt" else 0   # CREATE_NO_WINDOW: nada de janela de terminal no Windows


# ------------------------------------------------------------------ motor
def _rodar(args: list[str], timeout=300) -> tuple[int, str]:
    r = subprocess.run([python_motor(), *args], capture_output=True, text=True, timeout=timeout,
                       encoding="utf-8", errors="replace", env=env_motor(), creationflags=SEM_JANELA)
    return r.returncode, (r.stdout + ("\n" + r.stderr if r.stderr.strip() else "")).strip()


def _extrair(planilha: Path, saida: Path, bloco=None, usa_pse=False, numeros=None) -> tuple[int, str]:
    saida.mkdir(parents=True, exist_ok=True)
    args = [str(SCRIPTS / "extrair_dados.py"), str(planilha), "--saida", str(saida)]
    if bloco:
        args += ["--bloco", str(bloco)]
    # PSE desligada na ficha (decisão de Miguel: nenhum aluno usa PSE por enquanto). Para religar, passar "--usa-pse".
    if numeros:
        args += ["--numeros", ",".join(numeros)]
    return _rodar(args)


def _gerar_pdf(dados_json: Path, pdf: Path, contexto: dict | None) -> tuple[int, str]:
    args = [str(SCRIPTS / "gerar_pdf.py"), str(dados_json), "--saida", str(pdf)]
    if contexto:
        ctx = dados_json.parent / "contexto.json"
        A.gravar_json(ctx, contexto)
        args += ["--contexto", str(ctx)]
    return _rodar(args)


def erro_amigavel(log: str) -> str:
    """Extrai do log do motor só a mensagem que interessa."""
    linhas = log.splitlines()
    for i, l in enumerate(linhas):
        if "ERRO" in l:
            msg = "\n".join(linhas[i:]).replace("ERRO DE LAYOUT:", "").replace("ERRO:", "").strip()
            return _traduzir(msg)
    return "Não foi possível ler a planilha. Detalhes técnicos abaixo."


def _traduzir(txt: str) -> str:
    # mensagens do motor escritas para o Claude → orientações diretas para você
    trocas = [
        (r"NÃO escolha um número por conta própria.*?corrigir a planilha\.", "Corrija na planilha e envie de novo."),
        (r"Não reescreva o texto de Miguel: peça a ele para encurtar a observação ou dizer o que tirar\.",
         "Encurte a observação geral na planilha (aba aluno, A9)."),
        (r"NÃO corte a prescrição: avise Miguel \(encurtar o texto na planilha ou autorizar ajuste do template\)\.",
         "Encurte esses textos na planilha e envie de novo."),
        (r"\s*Nunca ajuste o template por conta própria\.", ""),
        (r"Miguel decide quais exercícios mostrar\.", ""),
    ]
    for padrao, novo in trocas:
        txt = re.sub(padrao, novo, txt, flags=re.S)
    return (txt.replace("regerar com --usa-pse", "ative \"Aluno usa PSE\" em Ajustes do PDF")
               .replace("Escolha com --numeros no extrair_dados.py.", "Escolha quais exercícios mostrar na revisão.")
               .replace("--numeros", "a escolha de \"Seus números\"")
               .replace("--usa-pse", "\"Aluno usa PSE\"")
               .replace("--bloco", "a escolha do bloco")
               .replace("contexto.json", "PRs informados")
               .replace("limitado por Miguel", "limitado por você"))


# ------------------------------------------------------------------ leitura direta (identificação e datas)
def _dt(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    try:
        return dt.datetime.strptime(str(v).strip(), "%d/%m/%Y").date()
    except (ValueError, TypeError):
        return None


def _tem_prescricao(ws) -> bool:
    """Aba 'bloco NN' com pelo menos um exercício que tenha séries na mesma semana
    (layout antigo: exercício em B, séries em C, I, O ou U; layout novo: B/C, I/J, P/Q, W/X)."""
    linhas = [tuple(r) + (None,) * 30 for r in ws.iter_rows(min_row=1, max_row=80, max_col=30, values_only=True)]
    cols = [(2, c0) for c0 in (3, 9, 15, 21)]
    for i, row in enumerate(linhas[:-2]):
        if re.fullmatch(r"\s*TREINO\s*\d+\s*", str(row[0] or ""), flags=re.I):
            sets = [c + 1 for c, v in enumerate(linhas[i + 1][:60]) if str(v or "").strip().upper() == "SETS"]
            if len(sets) >= 4:
                cols = [(c0 - 1 if str(linhas[i + 2][c0 - 2] or "").strip().lower().startswith("exerc") else 2, c0)
                        for c0 in sets[:4]]
            break
    for row in linhas[3:]:
        for ce, c0 in cols:
            ex = row[ce - 1]
            if not ex or str(ex).strip().lower() in ("exercícios", "exercicios", "sets", ""):
                continue
            if row[c0 - 1] not in (None, "") and str(row[c0 - 1]).strip():
                return True
    return False


def info_planilha(planilha: Path) -> dict:
    """Nome do aluno, blocos existentes e datas da aba macro (B6/F6/J6), lidos sem o motor."""
    out = {"nome": "", "blocos": [], "datas_macro": {}}
    try:
        wb = openpyxl.load_workbook(planilha, data_only=True, read_only=True)
    except Exception:  # noqa: BLE001
        return out
    for n in wb.sheetnames:
        m = re.fullmatch(r"\s*bloco\s*(\d+)\s*", n, flags=re.I)
        if m and _tem_prescricao(wb[n]):
            out["blocos"].append(int(m.group(1)))
        if re.fullmatch(r"\s*aluno\s*\d+\s*", n, flags=re.I) and not out["nome"]:
            v = wb[n]["C2"].value
            if v and str(v).strip() and not str(v).startswith("="):
                out["nome"] = re.sub(r"\s+", " ", str(v)).strip()
    if "macro" in wb.sheetnames:
        ws = wb["macro"]
        for b, col in ((1, "B"), (2, "F"), (3, "J")):
            d = _dt(ws[f"{col}6"].value)
            if d:
                out["datas_macro"][b] = d.isoformat()
    out["blocos"].sort()
    wb.close()
    return out


def plano_de(D: dict) -> str:
    """Plano = nº de treinos por semana (planilhas-base Consultoria_2x/3x/4x)."""
    n = len({v for v in D.get("dias", {}).values() if re.fullmatch(r"T\d+", str(v))})
    if not n and D.get("semanas"):
        n = len(D["semanas"][0].get("treinos", {}))
    return f"{n}x por semana" if n else ""


def datas(D: dict, macro: dict) -> tuple[str, str]:
    """(início do bloco, próxima atualização) em ISO. Próxima = data do bloco seguinte na aba macro,
    ou início + 4 semanas (revisão do bloco a cada 4 semanas)."""
    n = int(D["bloco"]["numero"])
    inicio = _dt(D["bloco"].get("inicio")) or (dt.date.fromisoformat(macro[n]) if macro.get(n) else None)
    prox = None
    if macro.get(n + 1):
        cand = dt.date.fromisoformat(macro[n + 1])
        if not inicio or cand > inicio:
            prox = cand
    if not prox and inicio:
        prox = inicio + dt.timedelta(days=28)
    return (inicio.isoformat() if inicio else ""), (prox.isoformat() if prox else "")


def classificar(relatorio: str, pse_definido: bool) -> dict:
    """Separa o relatório do motor em 'precisa da sua decisão' e 'detalhes', em linguagem amigável."""
    def secao(titulo):
        m = re.search(rf"## {titulo}[^\n]*\n(.*?)(\n## |\Z)", relatorio, flags=re.S)
        return [l.strip()[2:].strip() for l in (m.group(1).splitlines() if m else []) if l.strip().startswith("- ")]
    alertas, avisos_fixos = [], []
    for a in secao("Alertas"):
        a = a.lstrip("⚠").strip()
        if a.startswith("PSE:"):
            continue   # PSE desligada para todos os alunos (decisão de Miguel)
        if a.startswith("VTT/VTR da aba macro"):
            avisos_fixos.append(_traduzir(a).replace("**", ""))   # aviso de sempre, não pede decisão
            continue
        alertas.append(_traduzir(a).replace("**", ""))
    infos = avisos_fixos + [_traduzir(i).replace("**", "") for i in secao("Informações")]
    return {"alertas": alertas, "infos": infos}


def tabela(D: dict) -> list[dict]:
    out = []
    for s in D.get("semanas", []):
        for t, rows in s.get("treinos", {}).items():
            for r in rows:
                pct = r.get("pct")
                out.append({"semana": s.get("numero"), "treino": t, "exercicio": r.get("ex") or "",
                            "series": r.get("sets"), "reps": r.get("reps"), "kg": r.get("kg"),
                            "pct": f"{round(pct * 100)}%" if isinstance(pct, (int, float)) else "",
                            "obs": r.get("obs") or "", "video": bool(r.get("link"))})
    return out


# ------------------------------------------------------------------ envio (uma planilha recebida)
def _trabalho(slug: str, envio: str) -> Path:
    return A.pasta(slug) / "trabalho" / envio


def _entrada_planilha(aluno: dict, envio: str) -> dict | None:
    return next((p for p in aluno.get("planilhas", []) if p["id"] == envio), None)


def receber(nome_arquivo: str, conteudo: bytes, slug_destino: str | None = None) -> dict:
    """Recebe a planilha matriz. Retorna {ok, slug, envio, erro?, novo_aluno}."""
    envio = dt.datetime.now().strftime("%Y%m%d-%H%M%S")
    tmp = ENTRADA / envio
    nome_arquivo = Path(nome_arquivo).name
    planilha_tmp = tmp / nome_arquivo
    A.gravar_seguro(planilha_tmp, conteudo)
    info = info_planilha(planilha_tmp)

    # ajuste automático (Miguel, 9/out/2026): aquecimento numa CÓPIA (AMRAP como Miguel escreveu); a original fica intacta
    from . import ajuste
    ajustes, avisos_aj, ajustada = [], [], None
    try:
        novo_aj, ajustes, avisos_aj = ajuste.ajustar(planilha_tmp)
        if novo_aj:
            ajustada = tmp / f"ajustada_{nome_arquivo}"
            A.gravar_seguro(ajustada, novo_aj)
    except Exception as ex:  # noqa: BLE001
        avisos_aj = [f"Não consegui ajustar a planilha ({ex}); ela foi lida como você mandou."]
    leitura = ajustada or planilha_tmp

    code, log = _extrair(leitura, tmp / "saida")
    D = json.loads((tmp / "saida" / "dados.json").read_text(encoding="utf-8")) if code == 0 else None
    nome = (D["aluno"]["nome"] if D else info["nome"]).strip()
    if not nome:
        shutil.rmtree(tmp, ignore_errors=True)
        return {"ok": False, "erro": erro_amigavel(log) if code else "Não encontrei o nome do aluno (aba aluno, C2).",
                "log": log}

    aluno = A.carregar(slug_destino) if slug_destino else A.encontrar_por_nome(nome)
    novo = aluno is None
    if novo:
        aluno = A.criar(nome)
    elif nome not in aluno.get("nomes_planilha", []):
        aluno.setdefault("nomes_planilha", []).append(nome)
    slug = aluno["slug"]

    # guarda a planilha no histórico do aluno (nunca apagada) e a leitura na pasta de trabalho
    original_pl = A.pasta(slug) / "planilhas" / f"{envio}_{nome_arquivo}"
    A.gravar_seguro(original_pl, conteudo)
    destino_pl = original_pl
    if ajustada:
        destino_pl = A.pasta(slug) / "planilhas" / f"{envio}_ajustada_{nome_arquivo}"
        A.gravar_seguro(destino_pl, ajustada.read_bytes())
    trab = _trabalho(slug, envio)
    trab.mkdir(parents=True, exist_ok=True)
    if D:
        for f in (tmp / "saida").iterdir():
            shutil.move(str(f), str(trab / f.name))
    shutil.rmtree(tmp, ignore_errors=True)

    estado = {"envio": envio, "planilha": str(destino_pl.relative_to(A.pasta(slug))), "nome_original": nome_arquivo,
              "bloco": None, "usa_pse": False, "numeros": None,
              "blocos_disponiveis": info["blocos"], "datas_macro": info["datas_macro"],
              "extraido_em": A.agora(), "resultado": None,
              "planilha_original": str(original_pl.relative_to(A.pasta(slug))),
              "ajustes": ajustes, "avisos_ajuste": avisos_aj}
    entrada = {"id": envio, "enviado_em": A.agora(), "arquivo": estado["planilha"], "nome_original": nome_arquivo,
               "original": estado["planilha_original"], "ajustes": len(ajustes),
               "bloco": None, "inicio_bloco": "", "status": "Em revisão", "pdf": "", "alertas": 0}
    aluno.setdefault("planilhas", []).insert(0, entrada)

    if not D:
        estado["erro"] = erro_amigavel(log)
        estado["log"] = log
        entrada["status"] = "Erro na leitura"
        A.gravar_json(trab / "estado.json", estado)
        A.registrar(aluno, f"Planilha recebida ({nome_arquivo}) — erro na leitura")
        A.salvar(aluno)
        return {"ok": True, "slug": slug, "envio": envio, "novo_aluno": novo}

    # preferências já salvas do aluno: "Seus números"
    precisa = False
    nums = aluno.get("numeros")
    if nums and len(D.get("rm", {})) > 4 and all(x in D["rm"] for x in nums):
        estado["numeros"] = nums
        precisa = True
    A.gravar_json(trab / "estado.json", estado)
    A.salvar(aluno)
    if precisa:
        reextrair(slug, envio, registrar_evento=False)
    else:
        _atualizar_aluno(slug, envio, D)
    try:   # prepara a aba DADOS já na chegada da planilha (fica instantânea depois)
        from . import dados
        dados.ler_planilha(destino_pl)
    except Exception:  # noqa: BLE001
        pass
    aluno = A.carregar(slug)
    A.registrar(aluno, f"Planilha recebida ({nome_arquivo}) — bloco {_entrada_planilha(aluno, envio)['bloco'] or '?'}"
                       + (" · aluno novo no painel" if novo else ""))
    A.salvar(aluno)
    return {"ok": True, "slug": slug, "envio": envio, "novo_aluno": novo}


def _atualizar_aluno(slug: str, envio: str, D: dict):
    """Atualiza plano, bloco e datas do aluno a partir da leitura (só se este for o envio mais recente)."""
    aluno = A.carregar(slug)
    estado = A.ler_json(_trabalho(slug, envio) / "estado.json", {})
    rel = (_trabalho(slug, envio) / "relatorio_revisao.md").read_text(encoding="utf-8")
    inicio, prox = datas(D, {int(k): v for k, v in estado.get("datas_macro", {}).items()})
    n = int(D["bloco"]["numero"])
    entrada = _entrada_planilha(aluno, envio)
    entrada.update(bloco=n, inicio_bloco=inicio,
                   alertas=len(classificar(rel, aluno.get("pse_definido", False))["alertas"]))
    if entrada["status"] == "Erro na leitura":
        entrada["status"] = "Em revisão"
    mais_recente = aluno["planilhas"][0]["id"] == envio
    if mais_recente:
        if not aluno.get("plano_manual"):
            aluno["plano"] = plano_de(D)
        aluno["objetivo"] = D["aluno"].get("objetivo") or aluno.get("objetivo", "")
        aluno["bloco_atual"] = n
        aluno["inicio_bloco"] = inicio
        aluno["ultima_atualizacao"] = dt.date.today().isoformat()
        if not aluno.get("proxima_manual"):
            aluno["proxima_atualizacao"] = prox
    A.salvar(aluno)


def reextrair(slug: str, envio: str, usa_pse=None, numeros="__manter__", bloco="__manter__",
              registrar_evento=True) -> dict:
    trab = _trabalho(slug, envio)
    estado = A.ler_json(trab / "estado.json")
    aluno = A.carregar(slug)
    if numeros != "__manter__":
        estado["numeros"] = numeros or None
        if numeros:
            aluno["numeros"] = numeros
    if bloco != "__manter__":
        estado["bloco"] = bloco or None
    A.salvar(aluno)
    planilha = A.pasta(slug) / estado["planilha"]
    saida = trab / "_nova"
    shutil.rmtree(saida, ignore_errors=True)
    code, log = _extrair(planilha, saida, estado.get("bloco"), estado.get("usa_pse"), estado.get("numeros"))
    if code != 0:
        shutil.rmtree(saida, ignore_errors=True)
        estado.update(erro=erro_amigavel(log), log=log)
        A.gravar_json(trab / "estado.json", estado)
        return {"ok": False, "erro": estado["erro"]}
    for f in saida.iterdir():
        shutil.move(str(f), str(trab / f.name))
    saida.rmdir()
    for velho in [*trab.glob("Entrega_*.pdf"), *trab.glob("*_Bloco*.pdf")]:
        velho.unlink()     # a prévia antiga não vale mais para a nova leitura (os PDFs entregues ficam em pdfs/)
    estado.pop("erro", None)
    estado.pop("log", None)
    estado.update(extraido_em=A.agora(), resultado=None)
    A.gravar_json(trab / "estado.json", estado)
    D = json.loads((trab / "dados.json").read_text(encoding="utf-8"))
    _atualizar_aluno(slug, envio, D)
    if registrar_evento:
        a = A.carregar(slug)
        A.registrar(a, f"Planilha {estado['nome_original']} lida de novo ")
        A.salvar(a)
    return {"ok": True}


def revisao(slug: str, envio: str) -> dict:
    """Tudo o que a tela de revisão precisa."""
    aluno = A.carregar(slug)
    trab = _trabalho(slug, envio)
    estado = A.ler_json(trab / "estado.json", {})
    entrada = _entrada_planilha(aluno, envio) or {}
    out = {"aluno": resumo(aluno), "envio": envio, "estado": estado, "entrada": entrada}
    if estado.get("erro") and not (trab / "dados.json").exists():
        return {**out, "erro": estado["erro"], "log": estado.get("log", "")}
    D = json.loads((trab / "dados.json").read_text(encoding="utf-8"))
    rel = (trab / "relatorio_revisao.md").read_text(encoding="utf-8")
    inicio, prox = datas(D, {int(k): v for k, v in estado.get("datas_macro", {}).items()})
    pdfs = sorted([*trab.glob("Entrega_*.pdf"), *trab.glob("*_Bloco*.pdf")], key=lambda p: p.stat().st_mtime)
    out.update({
        "nome_planilha": D["aluno"]["nome"], "objetivo": D["aluno"].get("objetivo") or "",
        "bloco": int(D["bloco"]["numero"]), "inicio": inicio, "proxima": prox, "plano": plano_de(D),
        "semanas": [s["numero"] for s in D.get("semanas", [])],
        "treinos_semana": len({v for v in D.get("dias", {}).values() if re.fullmatch(r"T\d+", str(v))}),
        "dias": D.get("dias", {}), "rm": {k: {"reps": v["reps"], "kg": v["kg"], "est": round(v["est"], 1)}
                                          for k, v in D.get("rm", {}).items()},
        "tabela": tabela(D), "relatorio": rel, **classificar(rel, aluno.get("pse_definido", False)),
        "pdf": pdfs[-1].name if pdfs else "", "erro": estado.get("erro"),
        "ajustes": estado.get("ajustes", []),
    })
    out["alertas"] = estado.get("avisos_ajuste", []) + out["alertas"]
    return out


def gerar(slug: str, envio: str) -> dict:
    trab = _trabalho(slug, envio)
    estado = A.ler_json(trab / "estado.json")
    aluno = A.carregar(slug)
    D = json.loads((trab / "dados.json").read_text(encoding="utf-8"))
    n = int(D["bloco"]["numero"])
    ctx = {}      # "PR real" retirado a pedido de Miguel (9/out/2026): "Seus números" mostra só o 1RM estimado da aba prs
    # nome do arquivo (pedido de Miguel, 7/out/2026): "MariaSilva_Bloco02.pdf" (nome sem espaços, como na planilha)
    nome = re.sub(r'[\\/:*?"<>|\s]+', "", D["aluno"]["nome"])
    pdf = trab / f"{nome}_Bloco{n:02d}.pdf"
    code, log = _gerar_pdf(trab / "dados.json", pdf, ctx or None)
    estado["resultado"] = {"codigo": code, "log": log, "em": A.agora()}
    entrada = _entrada_planilha(aluno, envio)
    if code == 0:
        # Miguel (9/out/2026): o PDF refeito do mesmo bloco SUBSTITUI o anterior — um arquivo por bloco,
        # pdfs/MariaSilva_Bloco02.pdf. (PDFs antigos em subpastas com data ficam onde estão.)
        final = A.pasta(slug) / "pdfs" / pdf.name
        A.gravar_seguro(final, pdf.read_bytes())
        rel_final = final.relative_to(A.pasta(slug)).as_posix()
        entrada.update(status="PDF gerado", pdf=rel_final)
        aluno["entregas"] = [e for e in aluno.get("entregas", []) if e.get("pdf") != rel_final]
        aluno["entregas"].insert(0, {"bloco": n, "data": A.agora(), "pdf": rel_final,
                                     "envio": envio, "planilha": estado["nome_original"]})
        A.registrar(aluno, f"PDF do bloco {n:02d} gerado e conferido contra a planilha")
        A.gravar_json(trab / "estado.json", estado)
        A.salvar(aluno)
        return {"ok": True, "pdf": rel_final}
    if pdf.exists():
        msg = "A conferência PDF × planilha encontrou divergências. Não entregue este PDF — veja os detalhes."
    else:
        msg = erro_amigavel(log)
    estado["resultado"]["mensagem"] = msg
    A.gravar_json(trab / "estado.json", estado)
    A.salvar(aluno)
    return {"ok": False, "erro": msg, "log": log}


def vincular(slug_origem: str, envio: str, slug_destino: str) -> dict:
    """'Não é este aluno': move o envio para outro aluno e lembra o nome da planilha para as próximas vezes."""
    if slug_origem == slug_destino:
        return {"ok": True, "slug": slug_destino}
    origem, destino = A.carregar(slug_origem), A.carregar(slug_destino)
    entrada = _entrada_planilha(origem, envio)
    trab_o, trab_d = _trabalho(slug_origem, envio), _trabalho(slug_destino, envio)
    estado = A.ler_json(trab_o / "estado.json")
    pl_o = A.pasta(slug_origem) / estado["planilha"]
    pl_d = A.pasta(slug_destino) / estado["planilha"]
    pl_d.parent.mkdir(parents=True, exist_ok=True)
    shutil.move(str(pl_o), str(pl_d))
    orig = estado.get("planilha_original")
    if orig and orig != estado["planilha"] and (A.pasta(slug_origem) / orig).exists():
        shutil.move(str(A.pasta(slug_origem) / orig), str(A.pasta(slug_destino) / orig))
    trab_d.parent.mkdir(parents=True, exist_ok=True)
    shutil.move(str(trab_o), str(trab_d))
    origem["planilhas"] = [p for p in origem["planilhas"] if p["id"] != envio]
    destino.setdefault("planilhas", []).insert(0, entrada)
    destino["planilhas"].sort(key=lambda p: p["id"], reverse=True)
    for nome in origem.get("nomes_planilha", []):
        if nome not in destino.setdefault("nomes_planilha", []):
            destino["nomes_planilha"].append(nome)
    A.registrar(destino, f"Planilha {estado['nome_original']} vinculada a este aluno")
    A.salvar(destino)
    if not origem["planilhas"] and not origem.get("entregas"):
        shutil.rmtree(A.pasta(slug_origem), ignore_errors=True)   # aluno criado só por engano
    else:
        A.salvar(origem)
    D = A.ler_json(trab_d / "dados.json")
    if D:
        _atualizar_aluno(slug_destino, envio, D)
    return {"ok": True, "slug": slug_destino}


# ------------------------------------------------------------------ resumo para o painel
def resumo(aluno: dict) -> dict:
    hoje = dt.date.today()
    prox = aluno.get("proxima_atualizacao")
    dias = (dt.date.fromisoformat(prox) - hoje).days if prox else None
    if aluno.get("status") != "Ativo" or dias is None:
        situacao = ""
    elif dias < 0:
        situacao = "atrasada"
    elif dias <= 7:
        situacao = "proxima"
    else:
        situacao = "em-dia"
    semana = None
    if aluno.get("inicio_bloco"):
        d = (hoje - dt.date.fromisoformat(aluno["inicio_bloco"])).days
        semana = d // 7 + 1 if d >= 0 else 0
    pendentes = [p for p in aluno.get("planilhas", []) if p["status"] in ("Em revisão", "Erro na leitura")]
    return {k: aluno.get(k) for k in ("slug", "nome", "status", "plano", "plano_manual", "objetivo", "usa_pse",
                                       "pse_definido", "notas", "bloco_atual", "inicio_bloco", "ultima_atualizacao",
                                       "proxima_atualizacao", "proxima_manual", "atualizado_em", "criado_em",
                                       "nomes_planilha")} | {
        "dias_para_proxima": dias, "situacao": situacao, "semana_do_bloco": semana,
        "ultimo_pdf": (aluno.get("entregas") or [{}])[0].get("pdf", ""),
        "revisao_pendente": pendentes[0]["id"] if pendentes else "",
        "total_planilhas": len(aluno.get("planilhas", [])),
        "contrato": planos.situacao(aluno),
    }
