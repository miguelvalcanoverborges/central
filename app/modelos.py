# -*- coding: utf-8 -*-
"""Modelos de bloco: um bloco já prescrito guardado para montar o próximo bloco de outro aluno.

"Salvar como modelo"   lê a aba `bloco NN` de uma planilha matriz (de um aluno ou do "Aluno Modelo") e guarda,
                       por treino e semana: exercício, séries, repetições, %1RM, OBS e link de vídeo.
                       O kg do modelo fica só como referência: ele é do outro aluno e nunca vai para a planilha nova.
"Novo bloco a partir   monta a planilha matriz do aluno: copia a planilha atual dele (ou a matriz em branco, no
 de modelo"            começo de um ciclo novo), escreve o modelo na aba do bloco e calcula
                       kg = %1RM do modelo × 1RM do aluno (Brzycki, como a aba prs), arredondado a 1 kg (Miguel).
                       Linha sem %1RM (acessórios) fica com kg em branco e vai para a lista "Para ajustar".
                       Exercício sem 1RM do aluno também fica em branco e vai para a lista.

Nada é inventado: o 1RM vem da aba prs do aluno ou da série de referência que Miguel digita (reps × kg), e entra
na aba prs da planilha nova. A planilha recebida nunca é alterada: o resultado é um arquivo novo em
dados/alunos/<aluno>/montados/. As fórmulas da planilha (%1RM, macro, aba aluno) ficam intactas.

Os dois layouts da aba bloco são aceitos (detectados pelo cabeçalho, `extrair_dados.colunas_semanas`):
antigo = exercício só em B, semanas em C/I/O/U; novo (planilhas-base de out/2026) = exercício de cada semana na
própria coluna: B/C, I/J, P/Q, W/X.

dados/modelos/<id>/modelo.json + origem.xlsx (cópia da planilha de onde o modelo saiu)
matrizes/Consultoria_2x|3x|4x.xlsx: planilhas matriz em branco de Miguel (vêm com o programa);
dados/modelos/_matriz/<n>x.xlsx (+ .json): a que Miguel trocar em Configurações (tem prioridade)
"""
import datetime as dt
import math
import re
import shutil
import sys
from pathlib import Path

import openpyxl

from . import armazenamento as A
from . import xlsx_celulas
from .config import DADOS, LIXEIRA, RAIZ, SCRIPTS

sys.path.insert(0, str(SCRIPTS))
import extrair_dados as X  # noqa: E402  (funções de leitura do motor, sem alterá-lo)

MODELOS = DADOS / "modelos"
MATRIZ = MODELOS / "_matriz"
MATRIZES_PADRAO = RAIZ / "matrizes"
ENTRADA_MODELOS = DADOS / "_entrada"
ARREDONDAR_KG = 1.0          # decisão de Miguel (7/out/2026): carga calculada vai para o múltiplo de 1 kg mais próximo


class Erro(ValueError):
    """Mensagem pronta para mostrar a Miguel."""


# ------------------------------------------------------------------ utilidades
def _aba(nomes, tipo: str, n: int) -> str | None:
    return next((s for s in nomes if re.fullmatch(rf"\s*{tipo}\s*0*{n}\s*", s, flags=re.I)), None)


def _blocos(nomes) -> list[int]:
    return sorted(int(m.group(1)) for s in nomes if (m := re.fullmatch(r"\s*bloco\s*(\d+)\s*", s, flags=re.I)))


def _formula(wbf, aba: str, coord: str) -> bool:
    v = wbf[aba][coord].value
    return (isinstance(v, str) and v.startswith("=")) or type(v).__name__ in ("ArrayFormula", "DataTableFormula")


def _base(ex: str) -> str | None:
    n = X.sem_acento(ex or "")
    return next((lab for k, lab in X.BASES if k in n), None)


def _bruto(v):
    """Valor como Miguel escreveu (10.0 → 10; texto sem espaços sobrando)."""
    if v is None or isinstance(v, bool):
        return v
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, (int, float)):
        return v
    s = str(v).strip()
    return s or None


def _num(v):
    n = X.numero(v)
    return None if n is None or isinstance(n, bool) else float(n)


def rm_brzycki(reps, kg) -> float:
    return float(kg) / X.brzycki_frac(float(reps))


def arredondar(kg: float) -> float:
    q = math.floor(kg / ARREDONDAR_KG + 0.5) * ARREDONDAR_KG
    return int(q) if float(q).is_integer() else round(q, 2)


def _semanas_txt(ns: list[int]) -> str:
    ns = sorted(set(ns))
    if len(ns) == 1:
        return f"semana {ns[0]}"
    if ns == list(range(ns[0], ns[-1] + 1)):
        return f"semanas {ns[0]} a {ns[-1]}"
    return "semanas " + ", ".join(map(str, ns[:-1])) + f" e {ns[-1]}"


def _link(h):
    t = (h.target or h.location or "") if h is not None else ""
    return t.strip() if str(t).strip().lower().startswith("http") else None


def _abrir(p: Path):
    try:
        return openpyxl.load_workbook(p, data_only=True), openpyxl.load_workbook(p)
    except Exception as ex:  # noqa: BLE001
        raise Erro(f"Não consegui abrir a planilha ({ex}).") from ex


def _prs_quadro(wbv, n: int) -> dict[str, dict]:
    """{'Agachamento': {'linha': 18, 'reps': 5, 'kg': 100, 'rm': 112.5}} do quadro do bloco n na aba prs."""
    r0 = X.PRS_LINHA0.get(n)
    out = {}
    if not r0 or "prs" not in wbv.sheetnames:
        return out
    wp = wbv["prs"]
    for r in range(r0, r0 + 7):
        nome = wp.cell(r, 1).value
        if X.vazio(nome):
            continue
        reps, kg, cache = _num(wp.cell(r, 2).value), _num(wp.cell(r, 3).value), _num(wp.cell(r, 4).value)
        rm = cache if cache else (rm_brzycki(reps, kg) if reps and kg else None)
        out[str(nome).strip()] = {"linha": r, "reps": _bruto(reps) if reps else None, "kg": _bruto(kg) if kg else None,
                                  "rm": round(rm, 2) if rm else None}
    return out


def _dias(wbf, wbv, aba: str | None) -> dict[str, str]:
    """{'B': 'T01', 'D': 'T02'} — marcadores da linha 6 da aba aluno (T01, T02… e outros, como JJ)."""
    out = {}
    if not aba:
        return out
    for c in "BCDEFGH":
        v = X.resolver(wbf, wbv, aba, f"{c}6")
        if not X.vazio(v):
            out[c] = str(v).strip()
    return out


# ------------------------------------------------------------------ leitura de um bloco (para o modelo)
def blocos_prescritos(planilha: Path) -> list[int]:
    wbv, wbf = _abrir(planilha)
    out = []
    for n in _blocos(wbv.sheetnames):
        if _ler_treinos(wbv, wbf, n):
            out.append(n)
    return out


def _ler_treinos(wbv, wbf, n: int) -> list[dict]:
    bs = _aba(wbv.sheetnames, "bloco", n)
    if not bs:
        return []
    ws, wf = wbv[bs], wbf[bs]
    rm = {k: v["rm"] for k, v in _prs_quadro(wbv, n).items() if v["rm"]}

    # links de vídeo presos na aba aluno (espelho do bloco): {(semana, linha do bloco): url}
    links_al = {}
    asheet = _aba(wbv.sheetnames, "aluno", n)
    if asheet:
        af, sem = wbf[asheet], None
        for rr in range(1, min(af.max_row, 400) + 1):
            a0 = af.cell(rr, 1).value
            m = re.fullmatch(r"\s*SEMANA\s*(\d+)\s*", str(a0 or ""), flags=re.I)
            if m:
                sem = int(m.group(1))
                continue
            mm = re.search(r"'?!\$?[A-Z]+\$?(\d+)", str(a0 or ""))
            u = _link(af.cell(rr, 6).hyperlink)
            if u and sem and mm:
                links_al[(sem, int(mm.group(1)))] = u

    colunas = X.colunas_semanas(ws)
    treinos = []
    for num, (r0, r1) in X.detectar_treinos(ws, wf).items():
        linhas = []
        for r in range(r0, r1 + 1):
            semanas = []
            for w, (ce, c0) in enumerate(colunas):
                ex = ws.cell(r, ce).value
                sets, reps, kg, pct = (ws.cell(r, c0 + i).value for i in range(4))
                obs = ws.cell(r, c0 + 4).value
                if X.vazio(ex) or (X.vazio(sets) and X.vazio(reps)):
                    semanas.append(None)
                    continue
                ex = re.sub(r"\s+", " ", str(ex)).strip()
                base = _base(ex)
                kg_n, pct_n = _num(kg), _num(pct)
                if not (pct_n and pct_n > 0):
                    pct_n = (kg_n / rm[base]) if (kg_n and base in rm) else None
                semanas.append({"ex": ex, "base": base, "sets": _bruto(sets), "reps": _bruto(reps),
                                "pct": round(pct_n, 4) if pct_n else None, "kg_modelo": _bruto(kg), "obs": _bruto(obs),
                                "link": _link(wf.cell(r, c0 + 4).hyperlink) or links_al.get((w + 1, r))})
            if any(semanas):
                prim = next(x for x in semanas if x)
                linhas.append({"ex": prim["ex"], "base": prim["base"], "semanas": semanas})
        if linhas:
            treinos.append({"numero": num, "linhas": linhas})
    return treinos


def _sem_ex(lin: dict, s: dict) -> tuple[str, str | None]:
    """(exercício, base) da semana — modelos salvos antes do layout novo só têm o da linha."""
    return s.get("ex") or lin["ex"], (s.get("base") if s.get("ex") else lin["base"])


# ------------------------------------------------------------------ modelos guardados
def _resumo(m: dict) -> dict:
    linhas = [l for t in m["treinos"] for l in t["linhas"]]
    sem = sorted({w + 1 for l in linhas for w, s in enumerate(l["semanas"]) if s})
    bases = sorted({_sem_ex(l, s)[1] for l in linhas for s in l["semanas"] if s and s["pct"] and _sem_ex(l, s)[1]},
                   key=lambda b: [lab for _, lab in X.BASES].index(b))
    return {"id": m["id"], "nome": m["nome"], "criado_em": m["criado_em"], "origem": m["origem"],
            "treinos": len(m["treinos"]), "exercicios": len(linhas), "semanas": sem, "bases": bases}


def listar() -> list[dict]:
    out = []
    for p in sorted(MODELOS.glob("*/modelo.json")):
        m = A.ler_json(p)
        if m:
            out.append(_resumo(m))
    return sorted(out, key=lambda m: m["nome"].lower())


def carregar(mid: str) -> dict:
    if not re.fullmatch(r"[\w-]+", mid or ""):
        raise Erro("Modelo não encontrado.")
    m = A.ler_json(MODELOS / mid / "modelo.json")
    if not m:
        raise Erro("Modelo não encontrado.")
    return m


def salvar(planilha: Path, bloco: int, nome: str, origem: dict) -> dict:
    nome = re.sub(r"\s+", " ", nome or "").strip()
    if not nome:
        raise Erro("Dê um nome ao modelo.")
    if any(m["nome"].lower() == nome.lower() for m in listar()):
        raise Erro(f"Já existe um modelo chamado “{nome}”.")
    wbv, wbf = _abrir(planilha)
    treinos = _ler_treinos(wbv, wbf, int(bloco))
    if not treinos:
        raise Erro(f"O bloco {int(bloco):02d} desta planilha não tem prescrição.")
    mid = dt.datetime.now().strftime("%Y%m%d-%H%M%S") + "-" + A.slugify(nome)[:40]
    pasta = MODELOS / mid
    A.gravar_seguro(pasta / "origem.xlsx", planilha.read_bytes())
    m = {"id": mid, "nome": nome, "criado_em": A.agora(), "origem": {**origem, "bloco": int(bloco)},
         "treinos": treinos}
    A.gravar_json(pasta / "modelo.json", m)
    return _resumo(m)


def excluir(mid: str) -> None:
    carregar(mid)
    destino = LIXEIRA / "modelos" / f"{mid}__{dt.datetime.now():%Y%m%d-%H%M%S}"
    destino.parent.mkdir(parents=True, exist_ok=True)
    shutil.move(str(MODELOS / mid), str(destino))


def detalhe(mid: str) -> dict:
    m = carregar(mid)
    return {**_resumo(m), "treinos_detalhe": [
        {"numero": t["numero"], "linhas": [{"ex": l["ex"], "semanas": [
            s and {"ex": _sem_ex(l, s)[0], "sets": s["sets"], "reps": s["reps"], "pct": s["pct"],
                   "obs": X.separar_link(str(s["obs"]))[0] if s["obs"] is not None else None,
                   "video": bool(s["link"] or (s["obs"] is not None and X.separar_link(str(s["obs"]))[1]))}
            for s in l["semanas"]]} for l in t["linhas"]]} for t in m["treinos"]]}


# ------------------------------------------------------------------ planilhas matriz em branco (2x, 3x, 4x)
def _n_treinos(wbv) -> int:
    bs = _aba(wbv.sheetnames, "bloco", 1)
    ws = wbv[bs]
    return sum(1 for r in range(1, 220)
               if re.fullmatch(r"\s*TREINO\s*\d+\s*", str(ws.cell(r, 1).value or ""), flags=re.I))


def matrizes() -> dict[int, dict]:
    """{3: {'arquivo': Path, 'nome_original': ..., 'trocada': bool}} — a trocada em Configurações vale mais."""
    out = {}
    for p in sorted(MATRIZES_PADRAO.glob("Consultoria_*x.xlsx")):
        m = re.search(r"_(\d+)x\.xlsx$", p.name)
        if m:
            out[int(m.group(1))] = {"arquivo": p, "nome_original": p.name, "trocada": False}
    for p in sorted(MATRIZ.glob("*x.xlsx")):
        m = re.fullmatch(r"(\d+)x\.xlsx", p.name)
        info = A.ler_json(p.with_suffix(".json")) or {}
        if m:
            out[int(m.group(1))] = {"arquivo": p, "nome_original": info.get("nome_original", p.name), "trocada": True,
                                    "enviada_em": info.get("enviada_em", "")}
    return out


def matriz() -> list[dict]:
    """Resumo para a interface: [{'treinos': 3, 'nome_original': 'Consultoria_3x.xlsx', 'trocada': False}]."""
    return [{"treinos": n, "nome_original": v["nome_original"], "trocada": v["trocada"]}
            for n, v in sorted(matrizes().items())]


def salvar_matriz(nome_arquivo: str, conteudo: bytes) -> dict:
    tmp = DADOS / "_entrada" / f"matriz-{dt.datetime.now():%Y%m%d-%H%M%S}.xlsx"
    A.gravar_seguro(tmp, conteudo)
    try:
        wbv, wbf = _abrir(tmp)
        nomes = wbv.sheetnames
        if not _aba(nomes, "bloco", 1) or "prs" not in nomes or not _aba(nomes, "aluno", 1):
            raise Erro("Esta não parece a planilha matriz: faltam as abas aluno 01, bloco 01 ou prs.")
        cheios = [n for n in _blocos(nomes) if _ler_treinos(wbv, wbf, n)]
        if cheios:
            raise Erro("Esta planilha já tem prescrição (bloco " + ", ".join(f"{n:02d}" for n in cheios)
                       + "). Envie a planilha matriz em branco.")
        n = _n_treinos(wbv)
        if not n:
            raise Erro("Não encontrei os treinos (TREINO 01…) na aba bloco 01.")
        destino = MATRIZ / f"{n}x.xlsx"
        if destino.exists():   # a anterior fica guardada, nunca apagada
            antiga = MATRIZ / "anteriores" / f"{dt.datetime.now():%Y%m%d-%H%M%S}_{n}x.xlsx"
            antiga.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(destino, antiga)
        A.gravar_seguro(destino, conteudo)
        info = {"nome_original": Path(nome_arquivo).name, "enviada_em": A.agora(), "treinos": n}
        A.gravar_json(destino.with_suffix(".json"), info)
        return info
    finally:
        tmp.unlink(missing_ok=True)


# ------------------------------------------------------------------ novo bloco a partir de modelo
def _layout(wbv, wbf, aba: str) -> list[tuple[str, list[int]]]:
    """[('01', [6, 7, …, 15]), …]: linhas de exercício de cada treino na aba do bloco de destino.
    Vão da 3ª linha abaixo de "TREINO NN" até a última linha do treino com fórmula de %1RM ou exercício escrito."""
    ws = wbv[aba]
    colunas = X.colunas_semanas(ws)
    heads = []
    for r in range(1, 220):
        m = re.fullmatch(r"\s*TREINO\s*(\d+)\s*", str(ws.cell(r, 1).value or ""), flags=re.I)
        if m:
            heads.append((f"{int(m.group(1)):02d}", r))
    out = []
    for i, (num, h) in enumerate(heads):
        r0 = h + 3
        limite = heads[i + 1][1] - 1 if i + 1 < len(heads) else r0 + 14
        usadas = [r for r in range(r0, limite + 1)
                  if any(not X.vazio(ws.cell(r, ce).value)
                         or _formula(wbf, aba, f"{xlsx_celulas.col_letras(c0 + 3)}{r}") for ce, c0 in colunas)]
        fim = usadas[-1] if usadas else min(limite - 2, r0 + 9)
        vagas = list(range(r0, fim + 1))
        out.append((num, vagas))
    return out


def opcoes(aluno: dict, planilha: Path | None) -> dict:
    """O que a janela "Novo bloco a partir de modelo" precisa: destinos possíveis e o 1RM de cada um."""
    destinos, rm = [], {}
    anterior = {}     # último 1RM conhecido de cada exercício (só para mostrar como referência, nunca preenchido)
    if planilha:
        wbv, wbf = _abrir(planilha)
        prescritos = set()
        for n in _blocos(wbv.sheetnames):
            if _ler_treinos(wbv, wbf, n):
                prescritos.add(n)
            q = _prs_quadro(wbv, n)
            rm[str(n)] = {k: v for k, v in q.items() if v["reps"] and v["kg"]}
            for k, v in q.items():
                if v["reps"] and v["kg"]:
                    anterior[k] = {**v, "bloco": n}
        for n in _blocos(wbv.sheetnames):
            if n in X.PRS_LINHA0:
                destinos.append({"id": str(n), "rotulo": f"Bloco {n:02d}", "prescrito": n in prescritos})
    mz = matriz()
    destinos.append({"id": "novo", "rotulo": "Ciclo novo · bloco 01", "prescrito": False,
                     "indisponivel": not mz, "treinos": [m["treinos"] for m in mz]})
    rm["novo"] = {}
    atual = aluno.get("bloco_atual") or 0
    padrao = next((d["id"] for d in destinos if d["id"] != "novo" and int(d["id"]) > atual and not d["prescrito"]),
                  "novo")
    return {"destinos": destinos, "padrao": padrao, "rm": rm, "anterior": anterior, "matriz": mz,
            "modelos": listar(), "base": listar_base(), "principais": principais_blocos(planilha) if planilha else {},
            "ultimo_escrito": _ultimo_escrito(planilha) if planilha else None}


def montar(aluno: dict, planilha_atual: Path | None, modelo_id: str, destino: str, rms: dict) -> dict:
    """Monta a planilha do aluno com o modelo. rms = {'Agachamento': {'reps': 5, 'kg': 100}} (série de referência).
    Devolve {arquivo (relativo à pasta do aluno), bloco, pendencias, resumo}."""
    m = carregar(modelo_id)
    if destino == "novo":
        k = len(m["treinos"])
        mz = matrizes().get(k)
        if not mz:
            raise Erro(f"Falta a planilha matriz em branco de {k} treinos. Escolha em Configurações → Modelos de bloco.")
        origem, n = mz["arquivo"], 1
    else:
        if not planilha_atual:
            raise Erro("Este aluno ainda não tem planilha matriz lida.")
        origem, n = planilha_atual, int(destino)
    wbv, wbf = _abrir(origem)
    bs = _aba(wbv.sheetnames, "bloco", n)
    if not bs:
        raise Erro(f"A planilha não tem a aba do bloco {n:02d}.")
    layout = _layout(wbv, wbf, bs)
    if len(m["treinos"]) > len(layout):
        raise Erro(f"O modelo tem {len(m['treinos'])} treinos e a aba “{bs}” tem {len(layout)}. "
                   "Use a planilha matriz com o mesmo número de treinos.")
    for t, (num, vagas) in zip(m["treinos"], layout):
        if len(t["linhas"]) > len(vagas):
            raise Erro(f"O treino {t['numero']} do modelo tem {len(t['linhas'])} exercícios e o treino {num} da "
                       f"planilha tem {len(vagas)} linhas.")

    # 1RM do aluno: série de referência digitada (ou já na aba prs) → Brzycki, igual à planilha
    rm, ref = {}, {}
    for base, v in (rms or {}).items():
        reps, kg = _num((v or {}).get("reps")), _num((v or {}).get("kg"))
        if reps is None and kg is None:
            continue
        if not reps or not kg or reps < 1 or reps > 20 or kg <= 0:
            raise Erro(f"{base}: confira as repetições (1 a 20) e o kg da série de referência.")
        rm[base] = rm_brzycki(reps, kg)
        ref[base] = (_bruto(reps), _bruto(kg))

    val: dict[str, dict] = {bs: {}, "prs": {}}
    cache: dict[str, dict] = {bs: {}, "prs": {}}
    L = xlsx_celulas.col_letras

    def por(aba, coord, v, so_se_mudar=True):
        """Escreve v; fórmula fica (só o resultado guardado é atualizado no %1RM e no 1RM da prs)."""
        if _formula(wbf, aba, coord):
            return "formula"
        atual = wbv[aba][coord].value
        if so_se_mudar and (X.vazio(atual) and X.vazio(v)):
            return "igual"
        val.setdefault(aba, {})[coord] = v
        return "ok"

    colunas = X.colunas_semanas(wbv[bs])
    so_b = all(ce == 2 for ce, _ in colunas)        # layout antigo: um nome de exercício por linha (coluna B)
    sem_rm: dict[str, dict] = {}
    acessorios: list[str] = []
    nomes_diferentes: list[str] = []
    cargas = exercicios = 0
    for ti, (num, vagas) in enumerate(layout):
        linhas = m["treinos"][ti]["linhas"] if ti < len(m["treinos"]) else []
        for si, r in enumerate(vagas):
            lin = linhas[si] if si < len(linhas) else None
            if lin:
                exercicios += 1
            if so_b:
                por(bs, f"B{r}", lin["ex"] if lin else None)
                if lin and len({_sem_ex(lin, x)[0] for x in lin["semanas"] if x}) > 1:
                    nomes_diferentes.append(lin["ex"])
            for w, (ce, c0) in enumerate(colunas):
                s = lin["semanas"][w] if lin else None
                cs, cr, ck, cp, co = (f"{L(c0 + i)}{r}" for i in range(5))
                if not s:
                    for c in ((cs, cr, ck, co) if so_b else (f"{L(ce)}{r}", cs, cr, ck, co)):
                        por(bs, c, None)
                    if _formula(wbf, bs, cp):
                        cache[bs][cp] = ""
                    else:
                        por(bs, cp, None)
                    continue
                ex, base = _sem_ex(lin, s)
                if not so_b:
                    por(bs, f"{L(ce)}{r}", ex, False)
                kg = None
                if s["pct"] and base in rm:
                    kg = arredondar(s["pct"] * rm[base])
                    cargas += 1
                elif s["pct"]:
                    sem_rm.setdefault(base or ex, {}).setdefault(ex, []).append(w + 1)
                elif not X.vazio(s["kg_modelo"]) and ex not in acessorios:
                    acessorios.append(ex)
                obs = s["obs"] if not X.vazio(s["obs"]) else ""
                if s["link"] and s["link"] not in str(obs):
                    obs = f"{obs} - {s['link']}" if obs else s["link"]
                por(bs, cs, s["sets"], False)
                por(bs, cr, s["reps"], False)
                por(bs, ck, kg)
                por(bs, co, obs or None)
                pct_real = round(kg / rm[base], 4) if kg else ""
                if _formula(wbf, bs, cp):
                    cache[bs][cp] = pct_real          # o que a fórmula da planilha vai dar (kg ÷ 1RM)
                else:
                    por(bs, cp, pct_real or None)

    # aba prs do bloco: a série de referência que gerou o 1RM
    pend = []
    quadro = _prs_quadro(wbv, n)
    por_nome = {X.sem_acento(k): v for k, v in quadro.items()}
    for base, (reps, kg) in ref.items():
        q = por_nome.get(X.sem_acento(base))
        if not q:
            pend.append(f"A aba prs não tem a linha de {base} no bloco {n:02d}: o 1RM digitado foi usado só no cálculo.")
            continue
        r = q["linha"]
        por("prs", f"B{r}", reps, False)
        por("prs", f"C{r}", kg, False)
        if _formula(wbf, "prs", f"D{r}"):
            cache["prs"][f"D{r}"] = round(rm[base], 4)
        else:
            por("prs", f"D{r}", round(rm[base], 4), False)

    # ciclo novo: nome, objetivo e dias de treino vêm da planilha atual do aluno (dados dele, não do modelo)
    asheet = _aba(wbv.sheetnames, "aluno", n)
    if destino == "novo" and asheet:
        val.setdefault(asheet, {})
        nome = (aluno.get("nomes_planilha") or [aluno["nome"]])[0]
        por(asheet, "C2", nome, False)
        if aluno.get("objetivo"):
            por(asheet, "C3", aluno["objetivo"], False)
        if planilha_atual:
            pv, pf = _abrir(planilha_atual)
            ult = _aba(pv.sheetnames, "aluno", aluno.get("bloco_atual") or 1) or _aba(pv.sheetnames, "aluno", 1)
            for c, marca in _dias(pf, pv, ult).items():
                por(asheet, f"{c}6", marca, False)
            dias = _dias(pf, pv, ult)
        else:
            dias = {}
    else:
        dias = _dias(wbf, wbv, asheet)
    marcados = sorted({d for d in dias.values() if re.fullmatch(r"T\d+", d.upper())})
    if marcados and len(marcados) != len(m["treinos"]):
        pend.append(f"O modelo tem {len(m['treinos'])} treinos e os dias da planilha marcam {len(marcados)} "
                    f"({', '.join(marcados)}).")
    elif not marcados:
        pend.append("Marque os dias de treino na aba aluno (T01, T02…).")

    for base, exs in sem_rm.items():
        pend.append(f"Sem 1RM de {base}: kg em branco em "
                    + ", ".join(f"{ex} ({_semanas_txt(ws_)})" for ex, ws_ in exs.items()) + ".")
    if acessorios:
        pend.append("Carga para preencher (sem %1RM no modelo): " + ", ".join(acessorios) + ".")
    if nomes_diferentes:
        pend.append("O modelo troca o exercício entre semanas e esta planilha só tem um nome por linha (coluna B): "
                    + ", ".join(nomes_diferentes) + ". Confira.")

    bruto = origem.read_bytes()
    novo, mantidas = xlsx_celulas.escrever(bruto, {k: v for k, v in val.items() if v},
                                           {k: v for k, v in cache.items() if v})
    if mantidas:
        pend.append("Células com fórmula mantidas (o modelo não entrou nelas): " + ", ".join(mantidas[:8])
                    + (" …" if len(mantidas) > 8 else "") + ".")

    nome_arq = re.sub(r"[^\w-]+", "_", A.slugify(aluno["nome"]).replace("-", "_")).strip("_")
    rel = Path("montados") / f"{dt.datetime.now():%Y-%m-%d_%H%M%S}_{nome_arq}_Bloco{n:02d}.xlsx"
    A.gravar_seguro(A.pasta(aluno["slug"]) / rel, novo)
    return {"arquivo": rel.as_posix(), "bloco": n, "modelo": m["nome"], "ciclo_novo": destino == "novo",
            "pendencias": pend, "resumo": {"exercicios": exercicios, "cargas": cargas}}


# ------------------------------------------------------------------ modelos-base de periodização (só principais)
# Modelos de Miguel (atualizados em 8/out/2026; Ondulatória e Linear excluídas por ele). Valem SOMENTE para os exercícios principais (os de 1RM na aba prs: agachamento,
# supino, terra, desenvolvimento, remada, arranco, arremesso). Acessórios ficam exatamente como estão.
# "2×3 + 1×AMRAP" vai em DUAS linhas do mesmo exercício: a de cima com as séries fixas, a de baixo com 1 × AMRAP
# (regra de Miguel), no mesmo %1RM.
# AMRAP só 1 vez por semana para cada exercício principal (Miguel, 8/out/2026): no primeiro treino da semana em que ele
# aparece e o mais longe possível do AMRAP de outro principal; nos outros treinos, só as séries principais.
MODELOS_BASE = {
    "acumulacao": ("Acumulação linear", [[(4, 8, .60)], [(4, 8, .65)], [(5, 5, .70)], [(3, 3, .75), (1, "AMRAP", .75)]]),
    "acumulacao_reversa": ("Acumulação reversa", [[(4, 6, .75)], [(5, 5, .70)], [(3, 10, .65)], [(4, 8, .60)]]),
    "transformacao": ("Transformação", [[(4, 5, .70)], [(4, 4, .75)], [(4, 3, .80)], [(4, 2, .80)]]),
    "realizacao": ("Realização", [[(4, 3, .80)], [(3, 3, .85)], [(1, 2, .90)], [(1, 1, .95)]]),
}


def _linha_txt(linhas) -> str:
    return " + ".join(f"{s}×{r}" for s, r, _ in linhas) + f" {round(linhas[0][2] * 100)}%"


def listar_base() -> list[dict]:
    return [{"id": f"base:{k}", "nome": nome, "base": True, "semanas": [1, 2, 3, 4],
             "esquema": " · ".join(_linha_txt(w) for w in sem)} for k, (nome, sem) in MODELOS_BASE.items()]


def _nome_linha(ws, r: int, colunas) -> str | None:
    for ce, _ in colunas:
        v = ws.cell(r, ce).value
        if not X.vazio(v):
            return re.sub(r"\s+", " ", str(v)).strip()
    return None


def _principal(nome: str, nomes_prs: dict[str, str]) -> str | None:
    """Nome da aba prs a que o exercício corresponde (regra de Miguel: principais = os da aba prs), ou None.
    Mesma correspondência da fórmula de %1RM da planilha (agachamento, supino, terra…)."""
    b = _base(nome)
    if b and X.sem_acento(b) in nomes_prs:
        return nomes_prs[X.sem_acento(b)]
    n = X.sem_acento(nome or "")
    return next((v for k, v in nomes_prs.items() if k and k in n), None)


def _nomes_prs(wbv, n: int) -> dict[str, str]:
    """Exercícios da aba prs do bloco n que estão PREENCHIDOS (com 1RM estimado). Regra de Miguel (7/out/2026): os que
    estão na aba prs sem 1RM (ex.: Desenvolvimento, Remada vazios) não contam como principais."""
    return {X.sem_acento(k): k for k, v in _prs_quadro(wbv, n).items() if v["rm"]}


def _tem_exercicios(wbv, wbf, n: int) -> bool:
    bs = _aba(wbv.sheetnames, "bloco", n)
    if not bs:
        return False
    ws = wbv[bs]
    colunas = X.colunas_semanas(ws)
    return any(_nome_linha(ws, r, colunas) for _, vagas in _layout(wbv, wbf, bs) for r in vagas)


def _nomes_prs_para(wbv, n: int) -> dict[str, str]:
    """1RM preenchido na aba prs do bloco n; se ainda não foi preenchido, os do bloco anterior (o próximo bloco usa os
    mesmos principais até Miguel preencher o 1RM novo)."""
    return _nomes_prs(wbv, n) or (_nomes_prs(wbv, n - 1) if n > 1 else {})


def _bases_escritas(wbv, wbf, n: int, nprs: dict) -> list[str]:
    bs = _aba(wbv.sheetnames, "bloco", n)
    ws = wbv[bs]
    colunas = X.colunas_semanas(ws)
    achados = []
    for _, vagas in _layout(wbv, wbf, bs):
        for r in vagas:
            for ce, _ in colunas:
                v = ws.cell(r, ce).value
                p = _principal(str(v), nprs) if not X.vazio(v) else None
                if p and p not in achados:
                    achados.append(p)
    ordem = list(nprs.values())
    return sorted(achados, key=ordem.index)


def principais_blocos(planilha: Path) -> dict[str, list[str]]:
    """{'2': ['Agachamento', 'Supino'], 'novo': [...]} — principais de cada destino (para pedir o 1RM).
    Bloco já escrito: os exercícios dele com 1RM preenchido na prs. Bloco ainda vazio (ou ciclo novo): os do bloco
    anterior, que serão copiados para ele."""
    wbv, wbf = _abrir(planilha)
    out, ultimo = {}, None
    for n in _blocos(wbv.sheetnames):
        if _tem_exercicios(wbv, wbf, n):
            out[str(n)] = _bases_escritas(wbv, wbf, n, _nomes_prs_para(wbv, n))
            ultimo = n
        elif ultimo:
            out[str(n)] = _bases_escritas(wbv, wbf, ultimo, _nomes_prs(wbv, ultimo))
    if ultimo:
        out["novo"] = _bases_escritas(wbv, wbf, ultimo, _nomes_prs(wbv, ultimo))
    return out


def _copiar_bloco(aluno: dict, origem: Path, n_ori: int, base: Path, n_dst: int, nprs: dict,
                  ciclo_novo: bool) -> tuple[bytes, list[str]]:
    """Copia os exercícios do bloco n_ori (planilha do aluno) para o bloco n_dst de `base`, na mesma posição.
    Principais: só o nome (o modelo preenche o resto; a linha do AMRAP fica livre). Acessórios: como Miguel prescreveu
    no bloco anterior (séries, reps, kg, OBS e vídeo), para ele ajustar."""
    ov, of = _abrir(origem)
    dv, df = _abrir(base)
    bo, bd = _aba(ov.sheetnames, "bloco", n_ori), _aba(dv.sheetnames, "bloco", n_dst)
    if not bd:
        raise Erro(f"A planilha não tem a aba do bloco {n_dst:02d}.")
    wo, wd = ov[bo], dv[bd]
    co, cd = X.colunas_semanas(wo), X.colunas_semanas(wd)
    so_b = all(ce == 2 for ce, _ in cd)
    lo, ld = _layout(ov, of, bo), _layout(dv, df, bd)
    com = [(num, vagas) for num, vagas in lo if any(_nome_linha(wo, r, co) for r in vagas)]
    if len(com) > len(ld):
        raise Erro(f"O bloco {n_ori:02d} tem {len(com)} treinos e a aba “{bd}” tem {len(ld)}.")
    L = xlsx_celulas.col_letras
    val, pend, acessorios = {bd: {}}, [], []

    def por(coord, v):
        if not _formula(df, bd, coord) and not X.vazio(v):
            val[bd][coord] = v

    for (num, vagas), (_, dvagas) in zip(com, ld):
        ant = None
        for i, r in enumerate(vagas):
            nome = _nome_linha(wo, r, co)
            prin = bool(nome and _principal(nome, nprs))
            if not nome or (prin and ant and X.sem_acento(nome) == X.sem_acento(ant)):
                ant = nome
                continue                       # linha vazia, ou a linha do AMRAP (o modelo refaz)
            ant = nome
            if i >= len(dvagas):
                pend.append(f"{nome} (treino {num}) não coube no bloco {n_dst:02d}.")
                continue
            rd = dvagas[i]
            if so_b:
                por(f"B{rd}", nome)
            for w, ((ceo, c0o), (ced, c0d)) in enumerate(zip(co, cd)):
                ex_w = wo.cell(r, ceo).value
                ex_w = re.sub(r"\s+", " ", str(ex_w)).strip() if not X.vazio(ex_w) else nome
                if not so_b:
                    por(f"{L(ced)}{rd}", ex_w)
                if prin:
                    continue
                sets, reps, kg = (wo.cell(r, c0o + k).value for k in range(3))
                if X.vazio(sets) and X.vazio(reps):
                    continue
                obs = _bruto(wo.cell(r, c0o + 4).value)
                link = _link(of[bo].cell(r, c0o + 4).hyperlink)
                if link and link not in str(obs or ""):
                    obs = f"{obs} - {link}" if obs else link
                for k, v in enumerate((_bruto(sets), _bruto(reps), _bruto(kg))):
                    por(f"{L(c0d + k)}{rd}", v)
                por(f"{L(c0d + 4)}{rd}", obs)
                if nome not in acessorios:
                    acessorios.append(nome)
    if acessorios:
        pend.append(f"Acessórios copiados do bloco {n_ori:02d}, para você ajustar: " + ", ".join(acessorios) + ".")

    if ciclo_novo:                              # dados do aluno (dele, não do modelo) na matriz em branco
        asheet = _aba(dv.sheetnames, "aluno", 1)
        if asheet:
            val[asheet] = {}
            nome = (aluno.get("nomes_planilha") or [aluno["nome"]])[0]
            val[asheet]["C2"] = nome
            if aluno.get("objetivo"):
                val[asheet]["C3"] = aluno["objetivo"]
            ult = _aba(ov.sheetnames, "aluno", n_ori) or _aba(ov.sheetnames, "aluno", 1)
            for c, marca in _dias(of, ov, ult).items():
                if not _formula(df, asheet, f"{c}6"):
                    val[asheet][f"{c}6"] = marca
    novo, _ = xlsx_celulas.escrever(base.read_bytes(), {k: v for k, v in val.items() if v}, {})
    return novo, pend


def montar_proximo(aluno: dict, planilha_atual: Path | None, esquema: str, destino: str, rms: dict) -> dict:
    """Planilha do próximo bloco com o modelo-base. Se o bloco de destino ainda está vazio (ou é ciclo novo), os
    exercícios do último bloco escrito são copiados para ele na mesma posição; depois o modelo entra nos principais."""
    if not planilha_atual:
        raise Erro("Este aluno ainda não tem planilha matriz lida.")
    wbv, wbf = _abrir(planilha_atual)
    escritos = [n for n in _blocos(wbv.sheetnames) if _tem_exercicios(wbv, wbf, n)]
    if not escritos:
        raise Erro("A planilha não tem exercícios escritos.")
    if destino != "novo" and int(destino) in escritos:
        return montar_base(aluno, planilha_atual, esquema, destino, rms)
    ciclo_novo = destino == "novo"
    if ciclo_novo:
        n_ori = escritos[-1]
        ws_o = wbv[_aba(wbv.sheetnames, "bloco", n_ori)]
        k = sum(1 for _, vagas in _layout(wbv, wbf, _aba(wbv.sheetnames, "bloco", n_ori))
                if any(_nome_linha(ws_o, r, X.colunas_semanas(ws_o)) for r in vagas))
        mz = matrizes().get(k)
        if not mz:
            raise Erro(f"Falta a planilha matriz em branco de {k} treinos.")
        base, n_dst = mz["arquivo"], 1
    else:
        n_dst = int(destino)
        anteriores = [n for n in escritos if n < n_dst]
        if not anteriores:
            raise Erro(f"Não há bloco escrito antes do bloco {n_dst:02d} para copiar os exercícios.")
        n_ori, base = anteriores[-1], planilha_atual
    nprs = _nomes_prs(wbv, n_ori)
    if not nprs:
        raise Erro(f"Nenhum exercício com 1RM preenchido na aba prs do bloco {n_ori:02d}.")
    bruto, pend = _copiar_bloco(aluno, planilha_atual, n_ori, base, n_dst, nprs, ciclo_novo)
    tmp = ENTRADA_MODELOS / f"proximo-{dt.datetime.now():%Y%m%d-%H%M%S-%f}.xlsx"
    A.gravar_seguro(tmp, bruto)
    try:
        r = montar_base(aluno, tmp, esquema, str(n_dst), rms, nprs_fixo=nprs)
    finally:
        tmp.unlink(missing_ok=True)
    r["pendencias"] = pend + r["pendencias"]
    r["ciclo_novo"] = ciclo_novo
    return r


def montar_base(aluno: dict, planilha_atual: Path | None, esquema: str, destino: str, rms: dict,
                nprs_fixo: dict | None = None) -> dict:
    """Encaixa o modelo-base nos exercícios principais do bloco `destino` — os da aba prs (regra de Miguel), na
    posição em que ele os escreveu. Séries, reps, kg (= %1RM × 1RM, a 1 kg) e %1RM; AMRAP na linha de baixo.
    O resto (acessórios) não é tocado."""
    if esquema not in MODELOS_BASE:
        raise Erro("Modelo não encontrado.")
    nome_modelo, semanas_modelo = MODELOS_BASE[esquema]
    if not planilha_atual:
        raise Erro("Este aluno ainda não tem planilha matriz lida.")
    if not str(destino).isdigit():
        raise Erro("Escolha um bloco da planilha do aluno: o modelo de periodização entra nos exercícios já escritos.")
    n = int(destino)
    wbv, wbf = _abrir(planilha_atual)
    bs = _aba(wbv.sheetnames, "bloco", n)
    if not bs:
        raise Erro(f"A planilha não tem a aba do bloco {n:02d}.")
    ws = wbv[bs]
    colunas = X.colunas_semanas(ws)
    so_b = all(ce == 2 for ce, _ in colunas)
    L = xlsx_celulas.col_letras
    nprs = nprs_fixo or _nomes_prs_para(wbv, n)

    rm, ref = {}, {}
    for base, v in (rms or {}).items():
        reps, kg = _num((v or {}).get("reps")), _num((v or {}).get("kg"))
        if reps is None and kg is None:
            continue
        if not reps or not kg or reps < 1 or reps > 20 or kg <= 0:
            raise Erro(f"{base}: confira as repetições (1 a 20) e o kg da série de referência.")
        rm[base] = rm_brzycki(reps, kg)
        ref[base] = (_bruto(reps), _bruto(kg))

    val: dict[str, dict] = {bs: {}, "prs": {}}
    cache: dict[str, dict] = {bs: {}, "prs": {}}

    def por(aba, coord, v):
        if _formula(wbf, aba, coord):
            return
        if X.vazio(wbv[aba][coord].value) and X.vazio(v):
            return
        val.setdefault(aba, {})[coord] = v

    def pct_cel(coord, v):
        if _formula(wbf, bs, coord):
            cache[bs][coord] = v if v is not None else ""
        else:
            por(bs, coord, v)

    def ler_semanas(r) -> list[dict]:
        """Valores de uma linha, por semana, para movê-la (acessórios): nome, séries, reps, kg, OBS (+ vídeo)."""
        out = []
        for ce, c0 in colunas:
            obs = _bruto(ws.cell(r, c0 + 4).value)
            link = _link(wbf[bs].cell(r, c0 + 4).hyperlink)
            if link and link not in str(obs or ""):
                obs = f"{obs} - {link}" if obs else link
            out.append({"ex": _bruto(ws.cell(r, ce).value), "sets": _bruto(ws.cell(r, c0).value),
                        "reps": _bruto(ws.cell(r, c0 + 1).value), "kg": _bruto(ws.cell(r, c0 + 2).value),
                        "obs": obs, "pct": ws.cell(r, c0 + 3).value})
        return out

    def escrever_linha(r: int, nome_b, semanas: list | None):
        if so_b:
            por(bs, f"B{r}", nome_b)
        for w, (ce, c0) in enumerate(colunas):
            x = semanas[w] if semanas else None
            if not so_b:
                por(bs, f"{L(ce)}{r}", x and x["ex"])
            por(bs, f"{L(c0)}{r}", x and x["sets"])
            por(bs, f"{L(c0 + 1)}{r}", x and x["reps"])
            por(bs, f"{L(c0 + 2)}{r}", x and x["kg"])
            por(bs, f"{L(c0 + 4)}{r}", x and x["obs"])
            pct_cel(f"{L(c0 + 3)}{r}", x.get("pct") if x else None)

    # Aquecimento antes da série principal (regra de Miguel, 7/out/2026): 1 × 3 a 65%, 70% e 80% do 1RM, só os degraus
    # abaixo da carga principal (principal a 60% ou 65%: sem aquecimento). Ex.: principal 85% → 65, 70, 80 → principal.
    # Principal acima de 90% (realização, semana 4 a 95%) ganha mais um degrau de 90% (Miguel, 7/out/2026).
    AQUEC = (.65, .70, .80)

    def aquecimentos(pct: float) -> list[float]:
        return [a for a in AQUEC if a < pct - 1e-9] + ([.90] if pct > .90 + 1e-9 else [])

    def serie(base, ex_w, w, sets, reps, pct):
        nonlocal cargas
        kg = arredondar(pct * rm[base]) if base in rm else None
        if kg is not None:
            cargas += 1
        else:
            sem_rm.setdefault(base, set()).add(w + 1)
        return {"ex": ex_w, "sets": sets, "reps": reps, "kg": kg, "obs": None,
                "pct": round(kg / rm[base], 4) if kg is not None else None}

    W_MAX = max(len(aquecimentos(sem[0][2])) for sem in semanas_modelo)
    TEM_AMRAP = any(len(sem) > 1 for sem in semanas_modelo)
    pend, sem_rm = [], {}
    principais_n = cargas = 0
    # 1) o que há em cada treino, na ordem: principal (as linhas seguidas do mesmo principal viram um item só:
    #    aquecimento, série principal e AMRAP de antes) e acessórios (linha inteira, como Miguel escreveu)
    treinos = []
    for num, vagas in _layout(wbv, wbf, bs):
        itens = []
        for r in vagas:
            nome = _nome_linha(ws, r, colunas)
            if not nome:
                continue
            if _principal(nome, nprs):
                if itens and itens[-1]["tipo"] == "p" and X.sem_acento(itens[-1]["nome"]) == X.sem_acento(nome):
                    continue
                ex_sem = [re.sub(r"\s+", " ", str(ws.cell(r, ce).value)).strip() if not X.vazio(ws.cell(r, ce).value)
                          else nome for ce, _ in colunas]
                itens.append({"tipo": "p", "nome": nome, "ex": ex_sem, "aq": W_MAX, "base": _principal(nome, nprs),
                              "amrap": False})
            else:
                itens.append({"tipo": "a", "nome": nome, "r": r, "semanas": ler_semanas(r)})
        treinos.append((num, vagas, itens))

    # 1b) AMRAP só 1 vez por semana para cada principal (Miguel, 8/out/2026): no primeiro treino da semana em que ele
    #     aparece; se esse treino já tem o AMRAP de outro principal (o primeiro na ordem da planilha fica lá), vai para o
    #     treino em que ele aparece mais longe dos AMRAPs já marcados (empate: o mais cedo).
    if TEM_AMRAP:
        onde: dict[str, list[int]] = {}
        for t, (_, _, itens) in enumerate(treinos):
            for it in itens:
                if it["tipo"] == "p" and t not in onde.setdefault(it["base"], []):
                    onde[it["base"]].append(t)
        marcados: dict[int, str] = {}
        for base_p, cands in onde.items():
            if cands[0] not in marcados:
                t_escolhido = cands[0]
            else:
                t_escolhido = max(cands, key=lambda t: (min(abs(t - m) for m in marcados), -t))
            marcados.setdefault(t_escolhido, base_p)
            it = next(i for i in treinos[t_escolhido][2] if i["tipo"] == "p" and i["base"] == base_p)
            it["amrap"] = True

    for num, vagas, itens in treinos:
        if not any(it["tipo"] == "p" for it in itens):
            continue
        # 2) linhas: cada principal = aquecimentos + principal (+ AMRAP, se é o treino dele); acessório = 1.
        #    Abre espaço descendo os de baixo.
        cap = len(vagas)
        base_n = 0
        cabem = []
        for it in itens:
            precisa = 1 + (1 if it["tipo"] == "p" and it["amrap"] else 0)
            if base_n + precisa > cap:
                pend.append(f"Treino {num}: {it['nome']} não coube (o treino tem {cap} linhas).")
                continue
            base_n += precisa
            cabem.append(it)
        livre = cap - base_n
        for it in cabem:
            if it["tipo"] == "p":
                it["aq"] = min(W_MAX, livre)
                livre -= it["aq"]
                if it["aq"] < W_MAX:
                    pend.append(f"Treino {num}: só {it['aq']} de {W_MAX} linhas de aquecimento de {it['nome']} couberam.")
        # 3) escreve na ordem, linha por linha; o que sobrar no fim do treino fica vazio
        pos = 0
        for it in cabem:
            if it["tipo"] == "a":
                r = vagas[pos]
                pos += 1
                if r != it["r"]:                      # acessório desceu: leva a linha inteira
                    escrever_linha(r, it["nome"], it["semanas"])
                continue
            principais_n += 1
            k = it["aq"]
            linhas_p = [[None] * 4 for _ in range(k + 1 + (1 if it["amrap"] else 0))]
            for w, sem in enumerate(semanas_modelo):
                ex_w = it["ex"][w]
                base = _principal(ex_w, nprs) or _principal(it["nome"], nprs)
                sets, reps, pct = sem[0]
                aq = aquecimentos(pct)[-k:] if k else []          # se faltou linha, ficam os mais próximos da carga
                for j, a in enumerate(aq):
                    linhas_p[k - len(aq) + j][w] = serie(base, ex_w, w, 1, 3, a)
                linhas_p[k][w] = serie(base, ex_w, w, sets, reps, pct)
                if len(sem) > 1 and it["amrap"]:
                    linhas_p[k + 1][w] = serie(base, ex_w, w, *sem[1])
            for semanas_l in linhas_p:
                r = vagas[pos]
                pos += 1
                usada = any(semanas_l)
                escrever_linha(r, it["nome"] if usada else None, semanas_l if usada else None)
        for r in vagas[pos:]:
            if _nome_linha(ws, r, colunas) or any(not X.vazio(ws.cell(r, c0).value) for _, c0 in colunas):
                escrever_linha(r, None, None)
    if not principais_n:
        raise Erro(f"Nenhum exercício com 1RM preenchido na aba prs (bloco {n:02d}) está escrito no bloco. "
                   "Preencha a aba prs e os exercícios e envie a planilha de novo.")

    # aba prs do bloco: a série de referência que gerou o 1RM
    quadro = _prs_quadro(wbv, n)
    por_nome = {X.sem_acento(k): v for k, v in quadro.items()}
    for base, (reps, kg) in ref.items():
        q = por_nome.get(X.sem_acento(base))
        if not q:
            pend.append(f"A aba prs não tem a linha de {base} no bloco {n:02d}: o 1RM digitado foi usado só no cálculo.")
            continue
        r = q["linha"]
        if q["reps"] != reps or q["kg"] != kg:
            por("prs", f"B{r}", reps)
            por("prs", f"C{r}", kg)
        if _formula(wbf, "prs", f"D{r}"):
            cache["prs"][f"D{r}"] = round(rm[base], 4)
        else:
            por("prs", f"D{r}", round(rm[base], 4))

    for base, ws_ in sem_rm.items():
        pend.append(f"Sem 1RM de {base}: kg em branco nas {_semanas_txt(list(ws_))}.")

    novo, mantidas = xlsx_celulas.escrever(planilha_atual.read_bytes(), {k: v for k, v in val.items() if v},
                                           {k: v for k, v in cache.items() if v})
    if mantidas:
        pend.append("Células com fórmula mantidas (o modelo não entrou nelas): " + ", ".join(mantidas[:8])
                    + (" …" if len(mantidas) > 8 else "") + ".")
    nome_arq = re.sub(r"[^\w-]+", "_", A.slugify(aluno["nome"]).replace("-", "_")).strip("_")
    rel = Path("montados") / f"{dt.datetime.now():%Y-%m-%d_%H%M%S}_{nome_arq}_Bloco{n:02d}.xlsx"
    A.gravar_seguro(A.pasta(aluno["slug"]) / rel, novo)
    return {"arquivo": rel.as_posix(), "bloco": n, "modelo": nome_modelo, "ciclo_novo": False,
            "pendencias": pend, "resumo": {"exercicios": principais_n, "cargas": cargas}}


def _ultimo_escrito(planilha: Path) -> int | None:
    wbv, wbf = _abrir(planilha)
    escritos = [n for n in _blocos(wbv.sheetnames) if _tem_exercicios(wbv, wbf, n)]
    return escritos[-1] if escritos else None
