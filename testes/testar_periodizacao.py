# -*- coding: utf-8 -*-
"""Teste dos modelos-base de periodização (acumulação linear, acumulação reversa, transformação, realização), só com aluno FICTÍCIO,
numa cópia temporária da Central (os dados reais não são tocados).
Uso: python testes/testar_periodizacao.py      (código de saída ≠ 0 se algo falhar)

Regras de Miguel conferidas:
- séries, reps e %1RM do modelo valem SÓ para os exercícios principais; acessórios ficam como estavam;
- "2×3 + 1×AMRAP" em duas linhas do mesmo exercício (a de baixo com 1 × AMRAP, mesmo %1RM);
- kg = %1RM × 1RM (Brzycki da série de referência), múltiplo de 1 kg mais próximo; sem 1RM → kg em branco + aviso;
- aquecimento 1 × 3 a 65, 70 e 80% antes da principal (só os abaixo dela); a Central abre espaço descendo os
  exercícios de baixo; se o treino não tem linhas, aviso; a planilha recebida nunca é alterada;
- AMRAP só 1 vez por semana para cada principal: no primeiro treino em que ele aparece e o mais longe possível do
  AMRAP de outro principal; nos outros treinos, só a série principal;
- aplicar outro modelo por cima tira o AMRAP antigo das semanas que não têm mais AMRAP;
- a planilha montada sai no PDF pelo motor, conferida, com GRAVAR no AMRAP e ≥ 80%.
"""
import json
import os
import math
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
TMP = Path(tempfile.mkdtemp(prefix="central-periodizacao-"))
for parte in ("app", "motor", "icone", "matrizes"):
    shutil.copytree(RAIZ / parte, TMP / parte)
sys.path.insert(0, str(TMP))

from app import armazenamento as A  # noqa: E402
from app import modelos as M  # noqa: E402
from app import xlsx_celulas as XL  # noqa: E402

falhas = []


def confere(cond, msg):
    print(("  ok  " if cond else "  FALHOU  ") + msg)
    if not cond:
        falhas.append(msg)


L = XL.col_letras
BRZ = lambda reps, kg: kg / (1.0278 - 0.0278 * reps)  # noqa: E731
COLS = [(2, 3), (9, 10), (16, 17), (23, 24)]          # layout novo: (exercício, SETS) de cada semana
arred = lambda x: int(math.floor(x + 0.5))  # noqa: E731

# ---------------------------------------------------------------- planilha fictícia (bloco 02 com os exercícios escritos)
RM_REF = {"Agachamento": (5, 100), "Supino": (5, 70), "Terra": (3, 140)}     # Remada e Desenvolvimento sem 1RM na prs
val = {"prs": {"B18": 5, "C18": 100, "B19": 5, "C19": 70, "B20": 3, "C20": 140},   # aba prs do bloco 02 (fictícia)
       "aluno 01": {"C2": "Aluno Periodização (fictício)", "C3": "Ficar mais forte (fictício)",
                    "B6": "T01", "D6": "T02", "F6": "T03"},
       "bloco 02": {
           # treino 01: nomes só na coluna da semana 1 (B); linha 7 livre para o AMRAP
           "B6": "Agachamento livre", "B8": "Supino reto", "B10": "Cadeira extensora",
           "C10": 3, "D10": "10 - 12", "J10": 3, "K10": "10 - 12", "Q10": 3, "R10": "10 - 12", "X10": 3, "Y10": "10 - 12",
           "B11": "Prancha", "C11": 3, "D11": "30s",
           # treino 02: nomes em todas as semanas; Desenvolvimento sem linha livre embaixo
           **{f"{L(ce)}21": "Levantamento terra" for ce, _ in COLS},
           **{f"{L(ce)}23": "Remada curvada" for ce, _ in COLS},
           **{f"{L(ce)}25": "Desenvolvimento" for ce, _ in COLS},
           **{f"{L(ce)}26": "Elevação lateral" for ce, _ in COLS},
           "C26": 3, "D26": 12, "E26": 8,
           # treino 03: agachamento e supino de novo (AMRAP só 1 vez por semana para cada principal)
           "B36": "Agachamento livre", "B39": "Supino reto",
       }}
matriz = TMP / "matrizes" / "Consultoria_3x.xlsx"
origem = TMP / "ficticio.xlsx"
novo, _ = XL.escrever(matriz.read_bytes(), val, {})
origem.write_bytes(novo)
bytes_origem = origem.read_bytes()

aluno = {"slug": "aluno-periodizacao", "nome": "Aluno Periodização (fictício)", "bloco_atual": 1}
A.pasta(aluno["slug"]).mkdir(parents=True, exist_ok=True)
rms = {k: {"reps": r, "kg": kg} for k, (r, kg) in RM_REF.items()}
RM = {k: BRZ(*v) for k, v in RM_REF.items()}


def celulas(p: Path):
    return openpyxl.load_workbook(p, data_only=True)["bloco 02"], openpyxl.load_workbook(p)["bloco 02"]


def linha(ws, r, w):
    ce, c0 = COLS[w]
    return tuple(ws.cell(r, c).value for c in (ce, c0, c0 + 1, c0 + 2, c0 + 3))


print("1. Opções da janela")
pr = M.principais_blocos(origem)["2"]
confere(pr == ["Agachamento", "Supino", "Terra"],
        f"principais = exercícios da aba prs PREENCHIDOS (com 1RM) escritos no bloco 02: {pr}")
confere([m["id"] for m in M.listar_base()] == ["base:acumulacao", "base:acumulacao_reversa", "base:transformacao",
                                              "base:realizacao"], "4 modelos-base listados (sem a Linear)")

AQ = (.65, .70, .80)          # aquecimento de Miguel: 1 × 3 a 65, 70 e 80% (só os abaixo da carga principal)


def aq_de(p):                 # + degrau de 90% quando a principal passa de 90% (realização, semana 4)
    return [a for a in AQ if a < p - 1e-9] + ([.90] if p > .90 + 1e-9 else [])

TREINOS = {"01": list(range(6, 16)), "02": list(range(21, 31)), "03": list(range(36, 46))}
wo_v = openpyxl.load_workbook(origem, data_only=True)["bloco 02"]
ITENS = {"01": [("p", "Agachamento livre", "Agachamento"), ("p", "Supino reto", "Supino"), ("a", 10), ("a", 11)],
         "02": [("p", "Levantamento terra", "Terra"), ("a", 23), ("a", 25), ("a", 26)],
         "03": [("p", "Agachamento livre", "Agachamento"), ("p", "Supino reto", "Supino")]}
# onde fica o AMRAP (regra de Miguel): agachamento no T01 (primeiro); supino também aparece primeiro no T01, que já tem
# o AMRAP do agachamento → vai para o T03 (onde mais longe); terra só no T02.
AMRAP_EM = {("01", "Agachamento"), ("02", "Terra"), ("03", "Supino")}


def ler(ws, r):
    out = []
    for ce, c0 in COLS:
        t = tuple(ws.cell(r, c).value for c in (ce, c0, c0 + 1, c0 + 2))
        out.append(None if all(v in (None, "") for v in t) else t)
    return out


def esperado(sem, num):
    """Linhas que o treino deve ter, na ordem (mesma regra, escrita de outro jeito para conferir)."""
    cap = len(TREINOS[num])
    W = max(len(aq_de(s_[0][2])) for s_ in sem)
    A_ = 1 if any(len(s_) > 1 for s_ in sem) else 0
    tem = lambda it: A_ if (num, it[2]) in AMRAP_EM else 0  # noqa: E731
    livre = cap - sum(1 + tem(it) if it[0] == "p" else 1 for it in ITENS[num])
    linhas, avisos = [], 0
    for it in ITENS[num]:
        if it[0] == "a":
            linhas.append(ler(wo_v, it[1]))
            continue
        _, ex, b = it
        k = min(W, livre)
        livre -= k
        avisos += k < W
        bloco = [[None] * 4 for _ in range(k + 1 + tem(it))]
        for w, s_ in enumerate(sem):
            aq = aq_de(s_[0][2])[-k:] if k else []
            for j_, a in enumerate(aq):
                bloco[k - len(aq) + j_][w] = (ex, 1, 3, arred(a * RM[b]))
            bloco[k][w] = (ex, s_[0][0], s_[0][1], arred(s_[0][2] * RM[b]))
            if len(s_) > 1 and tem(it):
                bloco[k + 1][w] = (ex, s_[1][0], s_[1][1], arred(s_[1][2] * RM[b]))
        linhas += bloco
    return linhas + [[None] * 4] * (cap - len(linhas)), avisos


def _so_valores(linha_):
    return [None if t is None or all(v in (None, "") for v in t[1:]) else t[1:] for t in linha_]


def confere_treinos(arquivo, sem, rotulo, aba="bloco 02", so_valores=False):
    ws = openpyxl.load_workbook(arquivo, data_only=True)[aba]
    ok = True
    for num, rows in TREINOS.items():
        esp, _ = esperado(sem, num)
        for r, e in zip(rows, esp):
            got = ler(ws, r)
            if so_valores:               # bloco copiado: o nome vai para todas as semanas; confere séries/reps/kg
                got, e = _so_valores(got), _so_valores(e)
            if got != e:
                ok = False
                print(f"     treino {num} linha {r}: esperado {e} veio {got}")
    confere(ok, rotulo)


for esquema, (nome, sem) in M.MODELOS_BASE.items():
    print(f"2. {nome}")
    r = M.montar_base(aluno, origem, esquema, "2", rms)
    saida = A.pasta(aluno["slug"]) / r["arquivo"]
    confere_treinos(saida, sem, "aquecimento 65/70/80% só abaixo da principal, principal e AMRAP com o modelo; "
                                "acessórios descem quando falta espaço, na mesma ordem")
    avisos = sum(esperado(sem, n_)[1] for n_ in TREINOS)
    confere(sum("linhas de aquecimento" in p for p in r["pendencias"]) == avisos,
            f"aviso quando o aquecimento não cabe inteiro ({avisos})")
    wv, wf = celulas(saida)
    confere("=IFERROR" in str(wf["F6"].value) and "=IFERROR" in str(wf["M6"].value), "fórmulas de %1RM mantidas")
    pq = openpyxl.load_workbook(saida, data_only=True)["prs"]
    confere(pq["B18"].value == 5 and pq["C18"].value == 100 and abs(pq["D18"].value - RM["Agachamento"]) < 1e-3,
            "aba prs do bloco 02 com a série de referência e o 1RM")
confere(origem.read_bytes() == bytes_origem, "planilha recebida não foi alterada")
sem60 = M.MODELOS_BASE["acumulacao"][1]
confere(not aq_de(sem60[0][0][2]), "principal a 60%: sem aquecimento")

print("3. Outro modelo por cima (acumulação linear, com AMRAP → reversa → linear de novo)")
r1 = M.montar_base(aluno, origem, "acumulacao", "2", rms)
p1 = A.pasta(aluno["slug"]) / r1["arquivo"]
r1b = M.montar_base(aluno, p1, "acumulacao_reversa", "2", rms)
p1b = A.pasta(aluno["slug"]) / r1b["arquivo"]
confere_treinos(p1b, M.MODELOS_BASE["acumulacao_reversa"][1], "reversa por cima da linear: o AMRAP antigo some")
r2 = M.montar_base(aluno, p1b, "acumulacao", "2", rms)
confere_treinos(A.pasta(aluno["slug"]) / r2["arquivo"], M.MODELOS_BASE["acumulacao"][1],
                "aplicado por cima: aquecimentos e AMRAP antigos dão lugar aos novos, sem sobra")
w4 = openpyxl.load_workbook(A.pasta(aluno["slug"]) / r2["arquivo"], data_only=True)["bloco 02"]
confere(len([1 for num, rows in TREINOS.items() for r in rows if str(w4.cell(r, 25).value).upper() == "AMRAP"]) == 3,
        "semana 4: 3 AMRAPs (agachamento T01, terra T02, supino T03), um por exercício principal")

print("4. Motor: leitura e PDF conferido")
saida = A.pasta(aluno["slug"]) / r2["arquivo"]
motor = TMP / "motor" / "consultoria-planilha" / "scripts"
pasta = TMP / "pdf"
x = subprocess.run([sys.executable, str(motor / "extrair_dados.py"), str(saida), "--saida", str(pasta), "--bloco", "2"],
                   capture_output=True, text=True)
confere(x.returncode == 0, "extrair_dados leu o bloco 02 montado" + ("" if x.returncode == 0 else f": {x.stderr[-400:]}"))
if x.returncode == 0:
    d = json.loads((pasta / "dados.json").read_text(encoding="utf-8"))
    def linhas_de(o):
        if isinstance(o, dict):
            if "reps" in o and "obs" in o:
                yield o
            for v in o.values():
                yield from linhas_de(v)
        elif isinstance(o, list):
            for v in o:
                yield from linhas_de(v)
    linhas = list(linhas_de(d["semanas"]))
    amrap = [r for r in linhas if "amrap" in str(r.get("reps", "")).lower()]
    confere(amrap and all("gravar" in str(r.get("obs") or "").lower() for r in amrap), f"{len(amrap)} linhas AMRAP com GRAVAR")
    g = subprocess.run([sys.executable, str(motor / "gerar_pdf.py"), str(pasta / "dados.json"), "--saida",
                        str(pasta / "ficha.pdf")], capture_output=True, text=True)
    confere(g.returncode == 0 and "0 divergência" in g.stdout, "PDF gerado e conferido sem divergência"
            + ("" if g.returncode == 0 else f": {(g.stdout + g.stderr)[-400:]}"))

print("5. Próximo bloco (bloco 03 vazio) e ciclo novo, copiando os exercícios do bloco 02")
pb = M.principais_blocos(origem)
confere(pb.get("3") == ["Agachamento", "Supino", "Terra"] and pb.get("novo") == ["Agachamento", "Supino", "Terra"],
        f"principais do próximo bloco e do ciclo novo vêm do bloco 02: {pb.get('3')} / {pb.get('novo')}")
for destino, aba, linha_prs in (("3", "bloco 03", 30), ("novo", "bloco 01", 6)):
    rr = M.montar_proximo(aluno, origem, "acumulacao", destino, rms)
    sa = A.pasta(aluno["slug"]) / rr["arquivo"]
    wb = openpyxl.load_workbook(sa, data_only=True)
    ws = wb[aba]
    confere_treinos(sa, M.MODELOS_BASE["acumulacao"][1], f"{destino}: exercícios do bloco 02 copiados e Acumulação linear com aquecimento",
                    aba=aba, so_valores=True)
    confere(any("Acessórios copiados do bloco 02" in p for p in rr["pendencias"]), f"{destino}: aviso dos acessórios copiados")
    pq = wb["prs"]
    confere(pq[f"B{linha_prs}"].value == 5 and pq[f"C{linha_prs}"].value == 100, f"{destino}: 1RM na aba prs do bloco de destino")
    if destino == "novo":
        confere(wb["aluno 01"]["C2"].value == "Aluno Periodização (fictício)" and rr["ciclo_novo"],
                "ciclo novo: matriz em branco com o nome do aluno")
confere(origem.read_bytes() == bytes_origem, "planilha recebida continua sem alteração")

if os.environ.get("MANTER"):
    print("pasta do teste:", TMP)
else:
    shutil.rmtree(TMP, ignore_errors=True)
print("\nRESULTADO:", "PASSOU" if not falhas else f"FALHOU ({len(falhas)})")
sys.exit(1 if falhas else 0)
