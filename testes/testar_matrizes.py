# -*- coding: utf-8 -*-
"""Teste das planilhas matriz (Miguel, 22/out/2026: versão leve + layout "Leve" com a aba início).
Só aluno FICTÍCIO, numa cópia temporária da Central. Uso: python testes/testar_matrizes.py

Confere, na 2x, 3x e 4x:
- %1RM com uma busca simples (sem as 7 REGEXMATCH aninhadas) e sem a escala de cor das colunas de exercício;
- aba "início" na frente; as abas que a Central lê continuam com o mesmo nome;
- a Central lê a planilha preenchida: exercícios, séries, reps e kg do bloco 01 no dados.json.
"""
import json
import re
import shutil
import sys
import tempfile
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
TMP = Path(tempfile.mkdtemp(prefix="central-matrizes-"))
for parte in ("app", "motor"):
    shutil.copytree(RAIZ / parte, TMP / parte)
sys.path.insert(0, str(TMP))

from app import armazenamento as A  # noqa: E402
from app import planilha  # noqa: E402

falhas = []


def confere(cond, msg):
    print(("ok   " if cond else "FALHA ") + msg)
    if not cond:
        falhas.append(msg)


EXS = [("Agachamento", 4, 6, 70), ("Supino inclinado", 3, 8, 40), ("Remada curvada", 3, 10, 40)]

for n in ("2x", "3x", "4x"):
    arq = RAIZ / "matrizes" / f"Consultoria_{n}.xlsx"
    wb = openpyxl.load_workbook(arq)
    confere(wb.sheetnames[0] == "início", f"{n}: aba início na frente")
    confere({"aluno 01", "bloco 01", "bloco 02", "bloco 03", "prs", "macro"} <= set(wb.sheetnames), f"{n}: abas de sempre")
    formulas = [c.value for ws in wb.worksheets if ws.title.startswith("bloco") for row in ws.iter_rows() for c in row
                if isinstance(c.value, str) and c.value.startswith("=")]
    confere(not any("REGEXMATCH" in f for f in formulas), f"{n}: sem REGEXMATCH")
    confere(sum("INDEX(prs" in f for f in formulas) == 120 * int(n[0]), f"{n}: %1RM em todas as linhas de exercício")
    confere(all(not wb[s].conditional_formatting for s in wb.sheetnames if s.startswith("bloco")), f"{n}: sem escala de cor")
    # preenche um aluno fictício e manda para a Central
    wb["aluno 01"]["C2"] = f"Teste Matriz {n}"
    ws = wb["bloco 01"]
    for k, (ex, s, r, kg) in enumerate(EXS):
        ws.cell(6 + k, 2).value, ws.cell(6 + k, 3).value, ws.cell(6 + k, 4).value, ws.cell(6 + k, 5).value = ex, s, r, kg
    wb["prs"]["B6"], wb["prs"]["C6"] = 5, 100
    teste = TMP / f"Consultoria_{n} Teste.xlsx"
    wb.save(teste)
    res = planilha.receber(teste.name, teste.read_bytes())
    D = json.loads((A.pasta(res["slug"]) / "trabalho" / res["envio"] / "dados.json").read_text(encoding="utf-8"))
    linhas = [(x["ex"], x["sets"], x["reps"], x["kg"]) for x in D["semanas"][0]["treinos"]["01"]]
    confere([tuple(e) for e in EXS] == [(a, b, c, d) for a, b, c, d in linhas], f"{n}: a Central lê o treino do bloco 01")
    confere(re.sub(r"\s", "", D["aluno"]["nome"]) == f"TesteMatriz{n}", f"{n}: a Central lê o nome do aluno")

shutil.rmtree(TMP, ignore_errors=True)
print("\nRESULTADO:", "PASSOU" if not falhas else f"FALHOU ({len(falhas)})")
sys.exit(1 if falhas else 0)
