# -*- coding: utf-8 -*-
"""Teste do ajuste automático da planilha (aquecimento; AMRAP intocado), só com aluno FICTÍCIO, numa cópia
temporária da Central (os dados reais não são tocados). Uso: python testes/testar_ajuste.py

A planilha fictícia é feita na planilha matriz em branco de Miguel (testes/matriz_3x.xlsx, layout novo).
Confere:
- aquecimento 1 × 3 a 65/70/80% (+90% acima de 90%) antes da principal, só os degraus abaixo da carga;
  kg = %1RM × 1RM no múltiplo de 1 kg; acessórios descem e continuam iguais;
- AMRAP fica como Miguel escreveu (ele indica na planilha): escrito em dois treinos da mesma semana, fica nos dois;
- a original fica intacta; a ajustada guarda gráficos e fórmulas; o PDF sai conferido (GRAVAR nunca no aquecimento);
- a aba Dados lê a ajustada (o aquecimento entra no VTT).
"""
import math
import os
import shutil
import sys
import tempfile
import zipfile
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
TMP = Path(tempfile.mkdtemp(prefix="central-ajuste-"))
for parte in ("app", "motor", "icone"):
    shutil.copytree(RAIZ / parte, TMP / parte)
sys.path.insert(0, str(TMP))

from app import armazenamento as A  # noqa: E402
from app import dados as DADOS  # noqa: E402
from app import planilha as P  # noqa: E402
from app import xlsx_celulas as XL  # noqa: E402

falhas = []


def confere(cond, msg):
    print(("  ok  " if cond else "  FALHOU  ") + msg)
    if not cond:
        falhas.append(msg)


L = XL.col_letras
BRZ = lambda reps, kg: kg / (1.0278 - 0.0278 * reps)  # noqa: E731
arred = lambda x: int(math.floor(x + 0.5))  # noqa: E731
COLS = [(2, 3), (9, 10), (16, 17), (23, 24)]      # (exercício, SETS) de cada semana, layout novo
RMREF = {"Agachamento": (5, 100, 6), "Supino": (5, 70, 7), "Terra": (3, 140, 8)}   # reps, kg, linha da prs (bloco 01)
RM = {k: BRZ(r, kg) for k, (r, kg, _) in RMREF.items()}

# prescrição de Miguel (fictícia): semanas 1-4, principal com %1RM; AMRAP escrito na semana 4 em DOIS treinos
PRINC = [(4, 8, .60), (4, 8, .65), (5, 5, .70), (3, 3, .75)]
val = {"aluno 01": {"C2": "Aluno Ajuste (fictício)", "C3": "Ficar mais forte (fictício)",
                    "B6": "T01", "D6": "T02", "F6": "T03"},
       "prs": {}, "bloco 01": {}}
cache = {"bloco 01": {}, "prs": {}}
for ex, (reps, kg, r) in RMREF.items():
    val["prs"][f"B{r}"], val["prs"][f"C{r}"] = reps, kg
    cache["prs"][f"D{r}"] = round(RM[ex], 4)


def principal(r, nome, base, amrap):
    for w, ((ce, c0), (s_, rp, pc)) in enumerate(zip(COLS, PRINC)):
        kg = arred(pc * RM[base])
        val["bloco 01"].update({f"{L(ce)}{r}": nome, f"{L(c0)}{r}": s_, f"{L(c0 + 1)}{r}": rp, f"{L(c0 + 2)}{r}": kg})
        cache["bloco 01"][f"{L(c0 + 3)}{r}"] = round(kg / RM[base], 4)
        if w == 3 and amrap:
            val["bloco 01"].update({f"{L(ce)}{r + 1}": nome, f"{L(c0)}{r + 1}": 1, f"{L(c0 + 1)}{r + 1}": "AMRAP",
                                    f"{L(c0 + 2)}{r + 1}": kg})
            cache["bloco 01"][f"{L(c0 + 3)}{r + 1}"] = round(kg / RM[base], 4)


def acessorio(r, nome, sets, reps):
    for ce, c0 in COLS:
        val["bloco 01"].update({f"{L(ce)}{r}": nome, f"{L(c0)}{r}": sets, f"{L(c0 + 1)}{r}": reps})


# treino 01 (linhas 6-15): agachamento (AMRAP), supino (AMRAP), cadeira extensora
principal(6, "Agachamento livre", "Agachamento", True)
principal(8, "Supino reto", "Supino", True)
acessorio(10, "Cadeira extensora", 3, "10 - 12")
# treino 02 (21-30): terra (AMRAP), remada (acessório: sem 1RM na prs)
principal(21, "Levantamento terra", "Terra", True)
acessorio(23, "Remada curvada", 3, 10)
# treino 03 (36-45): agachamento e supino de novo, os dois com AMRAP escrito
principal(36, "Agachamento livre", "Agachamento", True)
principal(38, "Supino reto", "Supino", True)

origem = TMP / "Consultoria_3x Aluno Ajuste.xlsx"
novo, _ = XL.escrever((RAIZ / "testes" / "matriz_3x.xlsx").read_bytes(), val, cache)
origem.write_bytes(novo)
bytes_origem = origem.read_bytes()

print("1. Planilha enviada à Central")
r = P.receber(origem.name, bytes_origem)
confere(r["ok"], "planilha recebida")
slug, envio = r["slug"], r["envio"]
aluno = A.carregar(slug)
ent = aluno["planilhas"][0]
confere(ent["arquivo"] != ent["original"] and "ajustada" in ent["arquivo"], "cópia ajustada guardada ao lado da original")
confere((A.pasta(slug) / ent["original"]).read_bytes() == bytes_origem, "original intacta")
rv = P.revisao(slug, envio)
print("   ajustes:", *rv["ajustes"], sep="\n     - ")
confere(not any("AMRAP" in a for a in rv["ajustes"]), "nenhum ajuste de AMRAP (é de Miguel)")
confere(not rv["alertas"] or all("ajust" not in a.lower() for a in rv["alertas"]), "nenhum aviso de ajuste que não coube")

print("2. Conteúdo da planilha ajustada")
aj = A.pasta(slug) / ent["arquivo"]
ws = openpyxl.load_workbook(aj, data_only=True)["bloco 01"]


def treino(r0, n=10):
    out = []
    for r in range(r0, r0 + n):
        out.append([None if ws.cell(r, c0).value in (None, "") else
                    (ws.cell(r, ce).value, ws.cell(r, c0).value, ws.cell(r, c0 + 1).value, ws.cell(r, c0 + 2).value)
                    for ce, c0 in COLS])
    return out


def esperado(base, nome, amrap):
    """Linhas de um principal: aquecimento (alinhado embaixo), principal e, se for o treino dele, o AMRAP."""
    sem = []
    for w, (s_, rp, pc) in enumerate(PRINC):
        aq = [a for a in (.65, .70, .80) if a < pc - 1e-9]
        col = [(nome, 1, 3, arred(a * RM[base])) for a in aq] + [(nome, s_, rp, arred(pc * RM[base]))]
        if w == 3 and amrap:
            col.append((nome, 1, "AMRAP", arred(pc * RM[base])))
        sem.append(col)
    h = max(len(c) for c in sem)
    return [[(sem[w][len(sem[w]) - h + i] if i >= h - len(sem[w]) else None) for w in range(4)] for i in range(h)]


def acc(nome, sets, reps):
    return [[(nome, sets, reps, None)] * 4]


t1 = esperado("Agachamento", "Agachamento livre", True) + esperado("Supino", "Supino reto", True) \
    + acc("Cadeira extensora", 3, "10 - 12")
t2 = esperado("Terra", "Levantamento terra", True) + acc("Remada curvada", 3, 10)
t3 = esperado("Agachamento", "Agachamento livre", True) + esperado("Supino", "Supino reto", True)
for nome_t, r0, esp in (("01", 6, t1), ("02", 21, t2), ("03", 36, t3)):
    got = treino(r0)
    esp = esp + [[None] * 4] * (10 - len(esp))
    erros = [f"linha {r0 + i}: {g} ≠ {e}" for i, (g, e) in enumerate(zip(got, esp)) if g != e]
    confere(not erros, f"treino {nome_t}: aquecimento, principal, AMRAP e acessórios" + ("" if not erros else " — " + "; ".join(erros[:3])))
wf = openpyxl.load_workbook(aj)["bloco 01"]
confere(str(wf["F6"].value).startswith("=IFERROR") and str(wf["AA45"].value).startswith("=IFERROR"), "fórmulas de %1RM mantidas")
confere(any("chart" in i.filename for i in zipfile.ZipFile(aj).infolist()), "gráficos da planilha preservados")

print("3. PDF e aba Dados a partir da ajustada")
g = P.gerar(slug, envio, [])
confere(g["ok"], "PDF gerado e conferido" + ("" if g["ok"] else f": {g.get('erro')}"))
import json  # noqa: E402
D = json.loads((A.pasta(slug) / "trabalho" / envio / "dados.json").read_text(encoding="utf-8"))
linhas = [x for s in D["semanas"] for rows in s["treinos"].values() for x in rows]
aq_gravar = [x for x in linhas if "aquecimento" in str(x.get("obs") or "").lower() and "gravar" in str(x.get("obs") or "").lower()]
confere(not aq_gravar, "aquecimento nunca leva GRAVAR")
s4 = next(s for s in D["semanas"] if s["numero"] == 4)
n_amrap = sum(1 for rows in s4["treinos"].values() for x in rows if "amrap" in str(x.get("reps")).lower())
confere(n_amrap == 5, f"semana 4 no PDF: os 5 AMRAPs que Miguel escreveu — veio {n_amrap}")
R = DADOS.analisar(DADOS.planilha_atual(A.carregar(slug)))
s1 = next(s for s in R["semanas"] if s["global"] == 4)
vtt = 0
for base, nome in (("Agachamento", 2), ("Supino", 2), ("Terra", 1)):
    s_, rp, pc = PRINC[3]
    vtt += nome * (s_ * rp * arred(pc * RM[base]) + 3 * arred(.65 * RM[base]) + 3 * arred(.70 * RM[base]))
confere(abs(s1["vtt"] - vtt) < 0.5, f"aba Dados: VTT da semana 4 com o aquecimento = {s1['vtt']} (esperado {vtt})")

print("4. Planilha que já vem com o aquecimento escrito por Miguel")
val2 = {k: dict(v) for k, v in val.items()}
cache2 = {k: dict(v) for k, v in cache.items()}
for ce, c0 in COLS:      # treino 01 linha 5? não: Miguel escreve 2 linhas não-AMRAP no supino do T01 (linha 9 vira série extra)
    val2["bloco 01"].update({f"{L(ce)}9": "Supino reto", f"{L(c0)}9": 2, f"{L(c0 + 1)}9": 5, f"{L(c0 + 2)}9": 50})
o2 = TMP / "Consultoria_3x Aluno Ajuste 2.xlsx"
o2.write_bytes(XL.escrever((RAIZ / "testes" / "matriz_3x.xlsx").read_bytes(), val2, cache2)[0])
r2 = P.receber(o2.name, o2.read_bytes(), slug)
rv2 = P.revisao(slug, r2["envio"])
confere(not any("Supino reto: aquecimento" in a and "treino 01" in a for a in rv2["ajustes"]),
        "exercício com mais de uma série escrita por Miguel no treino não ganha aquecimento automático")

if not falhas and not os.environ.get("MANTER"):
    shutil.rmtree(TMP, ignore_errors=True)
else:
    print("pasta do teste:", TMP)
print("\nRESULTADO:", "PASSOU" if not falhas else f"{len(falhas)} FALHA(S)")
sys.exit(1 if falhas else 0)
