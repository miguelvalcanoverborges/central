# -*- coding: utf-8 -*-
"""Planilha FICTÍCIA de 12 semanas (3 blocos x 4 semanas) para testar a aba DADOS. Não é aluno real.
Uso: python testes/montar_planilha_12s.py saida.xlsx [nome] [data_inicio AAAA-MM-DD]"""
import datetime as dt
import sys
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
BASE = RAIZ / "motor" / "consultoria-planilha" / "exemplos" / "exemplo.xlsx"
saida = sys.argv[1]
nome = sys.argv[2] if len(sys.argv) > 2 else "Aluno Doze Semanas"
ini = dt.datetime.fromisoformat(sys.argv[3]) if len(sys.argv) > 3 else dt.datetime(2026, 8, 3)

wb = openpyxl.load_workbook(BASE)
wb["aluno 01"]["C2"] = nome
wb["aluno 01"]["C3"] = "Ficar mais forte (exemplo fictício)"

# 1RM de referência por bloco (aba prs): (reps, kg)
PRS = {1: {"Agachamento": (5, 100), "Supino": (5, 70), "Terra": (3, 130), "Arranco": (3, 50)},
       2: {"Agachamento": (5, 105), "Supino": (5, 72.5), "Terra": (3, 137.5), "Arranco": (3, 52.5)},
       3: {"Agachamento": (3, 115), "Supino": (3, 80), "Terra": (2, 150), "Arranco": (2, 57.5)}}
LINHA_PRS = {1: 6, 2: 18, 3: 30}
pr = wb["prs"]
RM = {}
for b, d in PRS.items():
    RM[b] = {}
    for r in range(LINHA_PRS[b], LINHA_PRS[b] + 7):
        n = pr[f"A{r}"].value
        if n in d:
            reps, kg = d[n]
            est = kg / (1.0278 - 0.0278 * reps)
            RM[b][n] = est
            pr[f"B{r}"], pr[f"C{r}"], pr[f"D{r}"] = reps, kg, round(est, 4)
    pr[f"B{LINHA_PRS[b] - 3}"] = ini + dt.timedelta(days=28 * (b - 1))

mc = wb["macro"]
for b, col in ((1, "B"), (2, "F"), (3, "J")):
    mc[f"{col}6"] = ini + dt.timedelta(days=28 * (b - 1))

# (exercício, base, séries[s1..s4], reps[s1..s4], %1RM alvo[s1..s4] ou kg fixo, obs)
INT = {1: [.62, .65, .68, .60], 2: [.72, .75, .78, .70], 3: [.80, .84, .88, .70]}
VOL = {1: [4, 5, 6, 3], 2: [5, 5, 6, 3], 3: [4, 4, 3, 2]}      # S03 do bloco 1 sobe >10%
REP = {1: 6, 2: 5, 3: 3}


def kg_de(base, pct, b):
    return round(RM[b][base] * pct / 2.5) * 2.5


def treinos(b, w):
    i, v, rp = INT[b][w], VOL[b][w], REP[b]
    return {
        6: [("Agachamento", "Agachamento", 3, 5, kg_de("Agachamento", .5, b), "aquecimento"),
            ("Agachamento", "Agachamento", v, rp, kg_de("Agachamento", i, b), "principal"),
            ("Supino", "Supino", v, rp + 1, kg_de("Supino", i - .02, b), None),
            ("Remada curvada", None, 3, 10, 40 + 2.5 * w, "biset"),
            ("Prancha lateral", None, 3, "30s", None, "biset")],
        21: [("Arranco", "Arranco", v + 1, 2 if b == 3 else 3, kg_de("Arranco", i - .05, b), "gravar"),
             ("Levantamento terra", "Terra", v - 1, rp - 1, kg_de("Terra", i, b), None),
             ("Desenvolvimento com halteres", None, 3, "8 - 12", 14, None),
             ("Afundo", None, 3, 10, 20 + 2.5 * (b - 1), None)],
        36: [("Agachamento frontal", "Agachamento", max(3, v - 1), rp, kg_de("Agachamento", i - .12, b), None),
             ("Supino inclinado", "Supino", 3, 8, kg_de("Supino", i - .1, b), None),
             ("Barra fixa", None, 3, "AMRAP", None, "amrap"),
             ("Elevação pélvica", None, 3, 12, 60 + 5 * w, None)],
    }


for b in (1, 2, 3):
    ws = wb[f"bloco {b:02d}"]
    for w, c0 in enumerate([3, 9, 15, 21]):
        for r0 in (3, 18, 33):
            ws.cell(r0, c0).value = ini + dt.timedelta(days=28 * (b - 1) + 7 * w)
        for r0, exs in treinos(b, w).items():
            for k, (ex, base, s, rp, kg, obs) in enumerate(exs):
                r = r0 + k
                ws.cell(r, 2).value = ex
                ws.cell(r, c0).value, ws.cell(r, c0 + 1).value = s, rp
                ws.cell(r, c0 + 2).value = kg
                ws.cell(r, c0 + 3).value = round(kg / RM[b][base], 4) if base and kg else None
                ws.cell(r, c0 + 4).value = obs
wb.save(saida)
print("ok", saida)
