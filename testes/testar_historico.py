# -*- coding: utf-8 -*-
"""Teste do histórico do aluno (Miguel, 21/out/2026: a planilha matriz é reenviada e, a cada trimestre, começa uma
planilha nova). Só alunos FICTÍCIOS, numa cópia temporária da Central. Uso: python testes/testar_historico.py

Confere:
- planilha nova do trimestre seguinte vira o ciclo 2: semanas 1 a 16 sem buraco, 1RM dos dois ciclos, PDFs com C2;
- bloco de 3 semanas: numeração sem pulo e aviso da data da aba macro;
- planilha antiga reenviada: aviso, o aluno não volta de bloco e o histórico mantém o bloco seguinte;
- semana já entregue alterada: aviso;
- feedback importado antes da planilha do bloco seguinte cai na semana certa pela data;
- bloco sem 1RM na aba prs e nome parecido: avisos.
"""
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
TMP = Path(tempfile.mkdtemp(prefix="central-historico-"))
for parte in ("app", "motor", "testes"):
    shutil.copytree(RAIZ / parte, TMP / parte)
sys.path.insert(0, str(TMP))

from app import armazenamento as A  # noqa: E402
from app import dados, individual, planilha  # noqa: E402

PL = TMP / "_planilhas"
PL.mkdir()
falhas = []


def confere(cond, msg):
    print(("ok   " if cond else "FALHA ") + msg)
    if not cond:
        falhas.append(msg)


def montar(arq, nome, ini, blocos=(1, 2, 3), semanas=None, sem_prs=(), kg_extra=None):
    """Planilha fictícia de 12 semanas (testes/montar_planilha_12s.py) só com os blocos pedidos."""
    arq = PL / arq
    subprocess.run([sys.executable, str(TMP / "testes" / "montar_planilha_12s.py"), str(arq), nome, ini],
                   check=True, capture_output=True)
    wb = openpyxl.load_workbook(arq)
    for b in (1, 2, 3):
        ws = wb[f"bloco {b:02d}"]
        n = 0 if b not in blocos else (semanas or {}).get(b, 4)
        for w, c0 in enumerate([3, 9, 15, 21]):
            for r in [*range(6, 16), *range(21, 31), *range(36, 46)]:
                if w >= n:
                    for i in range(5):
                        if type(ws.cell(r, c0 + i)).__name__ != "MergedCell":
                            ws.cell(r, c0 + i).value = None
                elif kg_extra and (b, w + 1) in kg_extra and isinstance(ws.cell(r, c0 + 2).value, (int, float)):
                    ws.cell(r, c0 + 2).value += kg_extra[(b, w + 1)]
        if b not in blocos or b in sem_prs:
            for r in range({1: 6, 2: 18, 3: 30}[b], {1: 6, 2: 18, 3: 30}[b] + 7):
                for c in "BCD":
                    wb["prs"][f"{c}{r}"] = None
    wb.save(arq)
    return arq


def enviar(arq):
    r = planilha.receber(arq.name, arq.read_bytes())
    est = A.ler_json(A.pasta(r["slug"]) / "trabalho" / r["envio"] / "estado.json")
    return r["slug"], r["envio"], est.get("conferencia", [])


def tem(avisos, trecho):
    return any(trecho in a for a in avisos)


# ---------------------------------------------------------------- ciclo 2 em planilha nova
slug, e1, _ = enviar(montar("carla_t1.xlsx", "Carla Teste", "2025-01-06"))
slug, e2, av = enviar(montar("carla_t2.xlsx", "Carla Teste", "2025-03-31", blocos=(1,)))
a = A.carregar(slug)
P = dados.historico(a)
confere([s["global"] for s in P["semanas"]] == list(range(1, 17)), "ciclo 2: semanas 1 a 16 sem buraco")
confere(P["ciclos"] == 2 and P["ciclo_envio"] == {e1: 1, e2: 2}, "ciclo 2: cada planilha no seu ciclo")
confere(P["semanas"][12]["rotulo_bloco"] == "Ciclo 2 · Bloco 01", "ciclo 2: rótulo do bloco com o ciclo")
R = dados.analisar(P)
confere(len(R["rm_evolucao"][0]["blocos"]) == 4, "ciclo 2: 1RM dos 3 blocos do ciclo 1 + o do ciclo 2")
confere(tem(av, "novo ciclo"), "ciclo 2: aviso de novo ciclo na revisão")
confere(dados.nome_arquivo("Carla Teste", 1, 2) == "CarlaTeste_C2_Bloco01", "ciclo 2: arquivo com C2 no nome")
confere(a["bloco_atual"] == 1, "ciclo 2: bloco atual = bloco 01 da planilha nova")

# ---------------------------------------------------------------- bloco de 3 semanas
slug, _, av = enviar(montar("bruno.xlsx", "Bruno Teste", "2025-01-06", blocos=(1, 2), semanas={1: 3}))
P = dados.historico(A.carregar(slug))
confere([s["global"] for s in P["semanas"]] == list(range(1, 8)), "bloco de 3 semanas: semanas 1 a 7 sem pulo")
confere(tem(av, "corrija a data na aba macro"), "bloco de 3 semanas: aviso da data na aba macro")

# ---------------------------------------------------------------- planilha antiga reenviada
slug, _, _ = enviar(montar("diego_nova.xlsx", "Diego Teste", "2025-01-06", blocos=(1, 2)))
slug, _, av = enviar(montar("diego_antiga.xlsx", "Diego Teste", "2025-01-06", blocos=(1,)))
a = A.carregar(slug)
confere(tem(av, "Parece uma versão antiga"), "planilha antiga: aviso")
confere(a["bloco_atual"] == 2, "planilha antiga: o aluno continua no bloco 02")
confere(len(dados.historico(a)["semanas"]) == 8, "planilha antiga: o histórico mantém as 8 semanas")
confere(len(dados.analisar_aluno(a)["rm_evolucao"][0]["blocos"]) == 2, "planilha antiga: o 1RM do bloco 02 continua")

# ---------------------------------------------------------------- semana já entregue alterada
slug, e1, _ = enviar(montar("elisa_v1.xlsx", "Elisa Teste", "2025-01-06", blocos=(1,)))
a = A.carregar(slug)
a["entregas"] = [{"bloco": 1, "ciclo": 1, "envio": e1, "pdf": "pdfs/ElisaTeste_Bloco01.pdf", "data": A.agora()}]
A.salvar(a)
slug, _, av = enviar(montar("elisa_v2.xlsx", "Elisa Teste", "2025-01-06", blocos=(1, 2), kg_extra={(1, 2): 10}))
confere(tem(av, "bloco 01 semana 02"), "semana entregue alterada: aviso com a semana")
slug, _, av = enviar(montar("elisa_v3.xlsx", "Elisa Teste", "2025-01-06", blocos=(1, 2, 3), kg_extra={(1, 2): 10}))
confere(not tem(av, "já entregues"), "semana entregue igual à da planilha anterior: sem aviso")

# ---------------------------------------------------------------- feedback importado antes da planilha do bloco 02
slug, _, _ = enviar(montar("hugo_v1.xlsx", "Hugo Teste", "2025-01-06", blocos=(1,)))
fb = [{"id": f"f{i}", "data": d, "global": 4, "pares": [["Adesão aos treinos", "Fiz todos os treinos"]]}
      for i, d in enumerate(["2025-02-09T20:00:00", "2025-02-16T20:00:00"])]   # gravados presos na semana 4
A.gravar_json(A.pasta(slug) / "feedbacks.json", fb)
slug, _, _ = enviar(montar("hugo_v2.xlsx", "Hugo Teste", "2025-01-06", blocos=(1, 2)))
sem = dados.analisar_aluno(A.carregar(slug))["semanas"]
lt = {x["global"]: x["notas"] for x in individual.linha_do_tempo(slug, sem)}
confere(bool(lt.get(5)) and bool(lt.get(6)) and not lt.get(4), "feedback: semanas 5 e 6 pela data, não presos na 4")

# ---------------------------------------------------------------- bloco sem 1RM e nome parecido
_, _, av = enviar(montar("gabi.xlsx", "Gabriela Teste", "2025-01-06", blocos=(1, 2), sem_prs=(2,)))
confere(tem(av, "sem 1RM na aba prs"), "bloco sem 1RM: aviso")
_, _, av = enviar(montar("gabi2.xlsx", "Gabriela Teste Souza", "2025-01-06", blocos=(1,)))
confere(tem(av, "O nome parece com o de Gabriela Teste"), "nome parecido: aviso")
_, _, av = enviar(montar("outro.xlsx", "Rafael Teste", "2025-01-06", blocos=(1,)))
confere(not tem(av, "O nome parece"), "primeiro nome diferente: sem aviso de nome")

shutil.rmtree(TMP, ignore_errors=True)
print("\nRESULTADO:", "PASSOU" if not falhas else f"FALHOU ({len(falhas)})")
sys.exit(1 if falhas else 0)
